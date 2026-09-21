# ==============================================================================
# Shared by BOTH partner-workbook generators - build_partner_dc_packages.py
# (the push tier) and refresh_partner_workbooks_daily.py (the daily tier) -
# so the rule deciding which clusters are off a partner's to-do list exists
# in exactly one place. Created 2026-09-21 (Jack's choice: one shared module,
# not a second copy), after the daily tier was found never to have received
# the 2026-09-19 exclusion fix at all: its 17:27-17:38 rebuild that day put
# 201 clusters the PARTNER had reported inaccessible, plus 64 dropped as
# excess capacity, back onto all 19 partners' "Available to Collect" sheets.
# Same shape as that day's other duplicated-logic misses (the oversampling
# cap, the dominant-ward pick): a fix landing in one copy, not its sibling.
# Import from here; never re-copy these into either script.
#
# Two things live here:
#   1. The exclusion rule - two optional, additive overlay files, "missing
#      file = no exclusions, not an error", matching the R-side loaders in
#      scripts/shared/frame_status.R:
#        - resampling/output/cluster_accessibility_overlay.csv: clusters a
#          partner reported inaccessible at CLUSTER grain (the ward itself
#          may be fine) - relocated IDP sites, unsafe coordinates. Built by
#          resampling/scripts/build_cluster_accessibility_overlay.py.
#        - resampling/output/target_correction_dropped_clusters.csv: Task 5
#          (2026-09-14) excess-capacity drops, append-only, built by
#          resampling/scripts/compute_target_correction_drops.R.
#   2. Two read-back helpers used by both generators' self-checks: the
#      placemark names in a KML file, and the ids a written workbook's
#      Sampling Points sheet still marks outstanding. Lifted verbatim from
#      build_partner_dc_packages.py's Section 9 (2026-09-19), so the push
#      tier's reconciliation and the daily tier's map check read files the
#      same way.
# ==============================================================================
import csv
import os
import re
from xml.sax.saxutils import unescape

import openpyxl

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
CLUSTER_ACCESSIBILITY_OVERLAY_CSV = PROJECT_DIR + r"\resampling\output\cluster_accessibility_overlay.csv"
TARGET_CORRECTION_DROPPED_CLUSTERS_CSV = PROJECT_DIR + r"\resampling\output\target_correction_dropped_clusters.csv"


def _load_cluster_id_set(path):
    if not os.path.exists(path):
        return set()
    with open(path, encoding="utf-8") as f:
        return {r["cluster_id"] for r in csv.DictReader(f)}


CLUSTER_ACCESSIBILITY_OVERLAY_EXCLUDED = _load_cluster_id_set(CLUSTER_ACCESSIBILITY_OVERLAY_CSV)
TARGET_CORRECTION_DROPPED_CLUSTERS = _load_cluster_id_set(TARGET_CORRECTION_DROPPED_CLUSTERS_CSV)


def cluster_overlay_excluded(cluster_id):
    """True if this cluster is off every partner to-do list by overlay:
    partner-reported inaccessible at cluster grain, or dropped as excess
    capacity. Ward accessibility and the Non-IDP <4-accessible-household
    threshold are separate checks, applied by each caller."""
    return cluster_id in CLUSTER_ACCESSIBILITY_OVERLAY_EXCLUDED or cluster_id in TARGET_CORRECTION_DROPPED_CLUSTERS


def exclusions_summary():
    return (f"Loaded {len(CLUSTER_ACCESSIBILITY_OVERLAY_EXCLUDED)} cluster-accessibility-overlay exclusion(s), "
            f"{len(TARGET_CORRECTION_DROPPED_CLUSTERS)} target-correction drop(s).")


# ---- read-back helpers (verbatim from build_partner_dc_packages.py Section 9) ----
_PLACEMARK_NAME_RE = re.compile(r"<Placemark\b[^>]*>\s*<name>(.*?)</name>", re.DOTALL)


def extract_kml_names(path):
    if not os.path.exists(path):
        return set()
    with open(path, encoding="utf-8") as f:
        content = f.read()
    return {unescape(n) for n in _PLACEMARK_NAME_RE.findall(content)}


def extract_workbook_active_ids(xlsx_path):
    """(non_idp_survey_ids, idp_cluster_ids) that a written workbook's
    'Sampling Points' sheet marks Not started/Partial - still outstanding."""
    if not os.path.exists(xlsx_path):
        return set(), set()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if "Sampling Points" not in wb.sheetnames:
            return set(), set()
        ws = wb["Sampling Points"]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            return set(), set()
        idx = {h: i for i, h in enumerate(header) if h is not None}
        needed = ("Point Type", "Cluster ID", "Survey ID", "Collection Status")
        if not all(h in idx for h in needed):
            return set(), set()
        non_idp_ids, idp_ids = set(), set()
        for row in rows_iter:
            if row[idx["Collection Status"]] not in ("Not started", "Partial"):
                continue
            point_type = row[idx["Point Type"]] or ""
            if point_type.startswith("Non-IDP household"):
                sid = row[idx["Survey ID"]]
                if sid:
                    non_idp_ids.add(sid)
            elif point_type.startswith("IDP cluster"):
                cid = row[idx["Cluster ID"]]
                if cid:
                    idp_ids.add(cid)
        return non_idp_ids, idp_ids
    finally:
        wb.close()


def kml_active_ids_for_partner(partner_root):
    """(non_idp_survey_ids, idp_cluster_ids) across a partner's core KML files
    (non_idp_households_{primary,reserve}.kml, idp_clusters_primary.kml).
    Skips MSNA_Light (its own isolated KML/sheet pair, never subject to this
    logic) and any in-place _archive/ backup folder (historical copies no
    field team loads - a false-positive source if walked)."""
    kml_non_idp_ids, kml_idp_ids = set(), set()
    for dirpath, _dirnames, files in os.walk(partner_root):
        norm_dirpath = dirpath.replace("\\", "/")
        if "/MSNA_Light" in norm_dirpath or "/_archive" in norm_dirpath:
            continue
        for fn in files:
            if fn in ("non_idp_households_primary.kml", "non_idp_households_reserve.kml"):
                kml_non_idp_ids |= extract_kml_names(os.path.join(dirpath, fn))
            elif fn == "idp_clusters_primary.kml":
                kml_idp_ids |= extract_kml_names(os.path.join(dirpath, fn))
    return kml_non_idp_ids, kml_idp_ids
