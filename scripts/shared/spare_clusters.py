"""
spare_clusters.py - spare (buffer) clusters, written 4 Oct 2026 for Jack's
weekend top-up ("2 spare clusters per stratum ... marked as spares, to use only
if a cluster can't be completed", recorded so unused spares can be left out of
the weights).

A spare cluster is an ordinary cluster in FULL/WORKING (so a submission at it
matches normally) that is ALSO listed in the register, which sits next to the
frame:
    1_sampling/output/data/data_collection/buffer_cluster_register.csv
    columns: cluster_id, strata_id, pop_type, partner, buffer_rank,
             drawn_batch, drawn_at, source_note

A spare with 0 achieved interviews is UNUSED:
  - excluded from every count (to-do lists, Sampling Points, Cluster Summary,
    05's capacity and projections, the dashboard's progress) and from weights;
  - shown to partners only in its own places: <pop>/KML/spare_clusters.kml
    (placemark names prefixed "SPARE - ") and a "Spare Clusters" sheet.
A spare with >= 1 achieved interview is USED and is an ordinary cluster
everywhere. "Achieved" is the canonical rule (completed, matched, not a
confirmed/contested deletion), counted per matched_cluster_id.

No register file = no spares: every caller behaves exactly as before.
The same names are constants in validity_checks' spare_cluster_integrity
module and in 2_monitoring's dashboard; change them together.
"""
import csv
import os

REGISTER_NAME = "buffer_cluster_register.csv"
REGISTER_COLUMNS = ["cluster_id", "strata_id", "pop_type", "partner", "buffer_rank", "drawn_batch", "drawn_at", "source_note"]
SPARE_KML = "spare_clusters.kml"
SPARE_SHEET = "Spare Clusters"
SPARE_PREFIX = "SPARE - "
SPARE_INSTRUCTION = ("SPARE cluster - use ONLY if a cluster in this LGA can't be completed, "
                     "and use spares in order of Spare Rank (1 first).")


def register_path(sampling_dir):
    return os.path.join(sampling_dir, "output", "data", "data_collection", REGISTER_NAME)


def load_register(sampling_dir):
    """{cluster_id: register row}; {} when there is no register. Stops on a malformed or duplicated register,
    since a silently half-read register would put spares on ordinary to-do lists."""
    path = register_path(sampling_dir)
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        missing = [c for c in REGISTER_COLUMNS if c not in (reader.fieldnames or [])]
        if missing:
            raise SystemExit(f"STOP: {path} lacks column(s) {missing}")
        rows = list(reader)
    ids = [r["cluster_id"] for r in rows]
    if len(ids) != len(set(ids)):
        raise SystemExit(f"STOP: {path} lists a cluster_id more than once")
    return {r["cluster_id"]: r for r in rows}


def unused_spare_ids(register, achieved_by_cluster):
    """Register clusters with 0 achieved interviews. achieved_by_cluster: {matched_cluster_id: count}."""
    return {cid for cid in register if achieved_by_cluster.get(cid, 0) <= 0}


def spare_rank(register, cluster_id):
    return register.get(cluster_id, {}).get("buffer_rank", "")


README_NOTE = (" 'Spare Clusters' = spare clusters for this partner's LGAs: NOT part of your target or to-do list. "
               "Use one ONLY if a cluster in the same LGA can't be completed, in order of Spare Rank (1 first); "
               "once you collect at a spare it becomes an ordinary cluster in the next update.")


def add_spare_sheet(wb, spare_rows, base_columns):
    """Append the "Spare Clusters" sheet (header in row 1, then one row per spare point/site, Spare Rank first).
    Shared by build_partner_dc_packages.py and refresh_partner_workbooks_daily.py so both tiers write the same."""
    import openpyxl
    columns = ["Spare Rank"] + [c for c in base_columns if c != "Spare Rank"]
    ws = wb.create_sheet(SPARE_SHEET)
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="595959")
    for row in sorted(spare_rows, key=lambda r: (r.get("State", ""), r.get("LGA", ""), str(r.get("Spare Rank", "")),
                                                 r.get("Cluster ID", ""), r.get("Survey ID", ""))):
        ws.append([row.get(c, "") for c in columns])
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(28, len(col) + 4))
    return ws


def summary(register, unused):
    if not register:
        return "Spare clusters: none registered."
    return (f"Spare clusters: {len(register)} registered, {len(unused)} unused (shown only as spares), "
            f"{len(register) - len(unused)} used (ordinary clusters now).")
