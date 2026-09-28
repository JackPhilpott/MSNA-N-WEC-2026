# ==============================================================================
# R3 (staged 2026-09-26, for Jack's click on the register page): release the Conpad-accessibility change to the LIVE partner package folder.
# Sibling of handover_kankia_and_leftovers_2026-09-25.py (R2) with the same lessons built in: extended-length paths (the backup root is ~198
# chars deep, so every backup path is over 260), no rmtree needed (files are removed or moved one by one), content-level comparison for
# .xlsx/.docx (Microsoft 365 rewrites metadata parts in synced folders, so md5s of office files drift), a per-file manifest BEFORE anything
# changes, a verified backup, dated log lines, and a stop-on-first-failure rule.
# DRY-RUN BY DEFAULT. --execute needs --jack-go "<the recorded go: who, when, which register decision>" and only touches what is listed in
# resampling/output/staged_packages_2026-09-26_conpad_accessibility/LIVE_CHANGE_LIST_FACT_package.csv (FACT package) and the workbooks
# that WORKBOOK_COMPARISON_vs_live.csv marks as changed.
#
# SCOPE (the two register buttons):
#   --scope A     Group A (Dandume, Faskari, Funtua, Malumfashi, Matazu: the Conpad accessibility change) + workbooks
#   --scope A_B   Group A + Group B (10 LGAs whose KMLs predate the 25 Sep chain: completions only) + workbooks
# Workbooks: replace the live workbook of every partner whose staged workbook differs from live (11 partners), never the 6 identical ones.
#
# Actions per file (from the change list): REPLACE = back up, then copy the staged file over the live one;
#   REMOVE = back up, then delete the live file;  ARCHIVE = back up, then MOVE the live factsheet into
#   Cluster_guide/_archived_dropped_clusters_2026-09-27/ (never deleted); IDENTICAL = must still be unchanged, not touched.
# NOT touched: LGA boundary KMLs, LGA map PNGs, master_log folders, accessibility reports, PDFs, any other partner's package, the KoBo tool.
# No partner note is sent by this script.
#
# Phases: baseline (staging only) | preflight (read-only) | run (backup, packages, workbooks, verify) | verify (read-only)
# Usage: python handover_conpad_accessibility_r3_2026-09-27.py <phase> --scope A|A_B [--execute --jack-go "..."]
# Rehearsal on a scratch replica: --live-root / --backup-root / --log-dir / --trace-dir override the defaults.
# ==============================================================================
import argparse
import csv
import hashlib
import os
import shutil
import stat
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path

BS = chr(92)
WS = Path(r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA")
SAMPLING = WS / "4. Data" / "MSNA N-WEC 2026" / "1_sampling"
OUT = SAMPLING / "resampling" / "output"
STAGE = OUT / "staged_packages_2026-09-26_conpad_accessibility"
CHANGES = STAGE / "LIVE_CHANGE_LIST_FACT_package.csv"
WB_CMP = STAGE / "WORKBOOK_COMPARISON_vs_live.csv"
MANIFEST = STAGE / "MANIFEST_staged_files.csv"
BASELINE = STAGE / "LIVE_BASELINE_2026-09-26.csv"
WORKING = SAMPLING / "output" / "data" / "data_collection" / "NGA_MSNA_2026_stage2_sampling_frame_v13_WORKING.csv"
TRACE_DEFAULT = OUT / "handover_traceability_2026-09-27_r3"
ARCHIVE_NAME = "_archived_dropped_clusters_2026-09-27"
META = ("docProps/", "customXml", "[trash]/", "[Content_Types].xml")
WB_PARTNERS = ["CARE", "COOPI", "CRS", "DRC", "FACT", "FHI 360", "INTERSOS", "IRC", "LHI", "MDM", "Malteser", "NRC", "PLAN", "Save the Children", "Solidarités",
               "Street Child of Nigeria", "ZOA"]
sys.path.insert(0, str(SAMPLING / "scripts" / "shared"))


def lp(p):
    s = os.path.abspath(str(p))
    pre = BS * 2 + "?" + BS
    return Path(s if s.startswith(pre) else pre + s)


def plain(p):
    s = str(p)
    pre = BS * 2 + "?" + BS
    return s[4:] if s.startswith(pre) else s


def md5(path):
    h = hashlib.md5()
    with open(lp(path), "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def office_content_sig(path):
    """Signature of an xlsx/docx that ignores Microsoft 365 metadata parts."""
    z = zipfile.ZipFile(lp(path))
    return hashlib.md5(repr(sorted((i.filename, i.CRC) for i in z.infolist() if not i.filename.startswith(META))).encode()).hexdigest()


def xlsx_values_sig(path):
    import openpyxl
    h = hashlib.md5()
    wb = openpyxl.load_workbook(lp(path), read_only=True, data_only=True)
    for n in wb.sheetnames:
        h.update(n.encode())
        for row in wb[n].iter_rows(values_only=True):
            if any(v is not None for v in row):
                h.update(repr(row).encode())
    return h.hexdigest()


def sig(path):
    """(md5, content_sig, values_sig) - what identifies 'this file, unchanged' robustly."""
    p = str(path)
    m = md5(path)
    if p.endswith(".docx"):
        return m, office_content_sig(path), ""
    if p.endswith(".xlsx"):
        return m, office_content_sig(path), xlsx_values_sig(path)
    return m, "", ""


def same_as_baseline(cur, base):
    """cur/base = (md5, content_sig, values_sig). Unchanged if md5 equal, or the content parts equal, or (xlsx) every cell value equal."""
    if cur[0] == base[0]:
        return True
    if cur[1] and cur[1] == base[1]:
        return True
    return bool(cur[2]) and cur[2] == base[2]


def read_csv(p):
    with open(lp(p), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def change_rows(scope):
    groups = ("A",) if scope == "A" else ("A", "B")
    return [r for r in read_csv(CHANGES) if r["group"][0] in groups]


def changed_workbook_partners():
    ch = set()
    for r in read_csv(WB_CMP):
        if r.get("sheet") in ("Strata Summary", "Sampling Points", "Available to Collect", "Cluster Summary"):
            if any(int(r.get(k) or 0) for k in ("keys_added", "keys_removed", "rows_with_changed_cells")):
                ch.add(r["partner"])
    return [p for p in WB_PARTNERS if p in ch]


def wb_rel(p):
    return f"{p}/{p}_sampling_points_summary.xlsx"


class Run:
    def __init__(self, a):
        self.a = a
        self.live, self.backup_root, self.trace = Path(a.live_root), Path(a.backup_root), Path(a.trace_dir)
        self.log_path = Path(a.log_dir) / f"R3_execution_log_{a.scope}.csv"
        self.ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.fails = []
        self.rows = change_rows(a.scope)
        self.wb_partners = changed_workbook_partners()

    def check(self, ok, msg):
        print(("  PASS  " if ok else "  FAIL  ") + msg)
        if not ok:
            self.fails.append(msg)
        return ok

    def log(self, phase, action, src="", dst="", h="", ok=True):
        if not self.a.execute:
            return
        new = not lp(self.log_path).exists()
        lp(self.log_path).parent.mkdir(parents=True, exist_ok=True)
        with open(lp(self.log_path), "a", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            if new:
                w.writerow(["timestamp", "phase", "action", "src", "dst", "md5", "ok", "scope", "go"])
            w.writerow([datetime.now().isoformat(timespec="seconds"), phase, action, src, dst, h, ok, self.a.scope, self.a.jack_go])

    def stop(self):
        print("\nSTOPPED, nothing further was done:")
        for m in self.fails:
            print("  -", m)
        sys.exit(1)

    def live_path(self, rel):
        return self.live / rel.replace("/", os.sep)

    # ---------------- baseline (staging only) ----------------
    def baseline(self):
        print("== baseline: record what every live file to be touched looks like NOW (writes only inside staging) ==")
        rows = []
        for r in read_csv(CHANGES):
            if r["action"] in ("REPLACE", "REMOVE", "ARCHIVE", "IDENTICAL"):
                p = self.live_path(r["path_under_package_FACT"])
                if p.is_file() or lp(p).is_file():
                    m, c, v = sig(p)
                    rows.append({"kind": "pkg", "path": r["path_under_package_FACT"], "size": lp(p).stat().st_size, "md5": m, "content_sig": c, "values_sig": v})
        for pt in WB_PARTNERS:
            p = self.live / wb_rel(pt)
            m, c, v = sig(p)
            rows.append({"kind": "wb", "path": wb_rel(pt), "size": lp(p).stat().st_size, "md5": m, "content_sig": c, "values_sig": v})
        print(f"  {len(rows)} live files recorded ({sum(1 for r in rows if r['kind'] == 'pkg')} package files + {sum(1 for r in rows if r['kind'] == 'wb')} workbooks)")
        if self.a.execute:
            with open(lp(BASELINE), "w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=["kind", "path", "size", "md5", "content_sig", "values_sig"])
                w.writeheader()
                w.writerows(rows)
            print("  wrote", BASELINE)
        else:
            print("  (dry-run) would write", BASELINE)

    def load_baseline(self):
        if not lp(BASELINE).exists():
            return None
        return {r["path"]: (r["md5"], r["content_sig"], r["values_sig"]) for r in read_csv(BASELINE)}

    # ---------------- checks ----------------
    def staging_checks(self):
        man = {r["path"]: r["md5"] for r in read_csv(MANIFEST)}
        need = [(r["path_under_package_FACT"], STAGE / r["path_under_package_FACT"]) for r in self.rows if r["action"] == "REPLACE"]
        need += [(f"workbooks/{wb_rel(p)}", STAGE / "workbooks" / wb_rel(p)) for p in self.wb_partners]
        bad = [k for k, p in need if not lp(p).is_file() or man.get(k) != md5(p)]
        self.check(not bad, f"all {len(need)} staged files to be copied exist and match MANIFEST_staged_files.csv md5" + (f" | MISMATCH {bad[:3]}" if bad else ""))
        c = {a: sum(1 for r in self.rows if r["action"] == a) for a in ("REPLACE", "ARCHIVE", "REMOVE", "IDENTICAL")}
        print(f"  plan for scope {self.a.scope}: {c['REPLACE']} REPLACE, {c['ARCHIVE']} ARCHIVE (move), {c['REMOVE']} REMOVE, {c['IDENTICAL']} identical (untouched); "
              f"{len(self.wb_partners)} workbooks to replace: {self.wb_partners}")
        self.check(not any(r["action"].startswith("UNEXPECTED") or r["action"] == "ADD" for r in self.rows), "change list has no UNEXPECTED_* and no ADD rows")
        exp = {"A": (18, 46, 3), "A_B": (41, 81, 7)}[self.a.scope]
        self.check((c["REPLACE"], c["ARCHIVE"], c["REMOVE"]) == exp, f"counts equal what the register page states for this scope (REPLACE/ARCHIVE/REMOVE = {exp})")
        self.check(len(self.wb_partners) == 11, f"11 workbooks differ from live ({self.wb_partners})")

    def live_checks(self):
        base = self.load_baseline()
        if not self.check(base is not None, f"live baseline present ({BASELINE.name}); run 'baseline --execute' first if not"):
            return
        drift = []
        for r in self.rows:
            rel = r["path_under_package_FACT"]
            if r["action"] in ("REPLACE", "REMOVE", "ARCHIVE", "IDENTICAL") and rel in base:
                p = self.live_path(rel)
                if not lp(p).is_file() or not same_as_baseline(sig(p), base[rel]):
                    drift.append(rel)
            elif r["action"] in ("REPLACE", "REMOVE", "ARCHIVE"):
                drift.append(rel + " (no baseline)")
        self.check(not drift, f"every live package file in the change list is unchanged since the baseline ({len(self.rows)} files checked)" + (f" | DRIFT: {drift[:4]}" if drift else ""))
        wbd = []
        for pt in self.wb_partners:
            p = self.live / wb_rel(pt)
            if not lp(p).is_file() or not same_as_baseline(sig(p), base[wb_rel(pt)]):
                wbd.append(pt)
        self.check(not wbd, "live workbooks to be replaced are unchanged since the baseline (nobody saved new content over them; metadata rewrites tolerated)" + (f" | CHANGED: {wbd}" if wbd else ""))
        clash = [r["path_under_package_FACT"] for r in self.rows if r["action"] == "ARCHIVE"
                 and (lp(self.live_path(r["path_under_package_FACT"])).parent / ARCHIVE_NAME / os.path.basename(r["path_under_package_FACT"])).exists()]
        self.check(not clash, f"no ARCHIVE destination exists yet (nothing can be overwritten in {ARCHIVE_NAME}/)" + (f" | CLASH: {clash[:3]}" if clash else ""))
        self.check(self.backup_root.is_dir(), f"backup root exists: {plain(self.backup_root)}")
        dst = self.backup_root / f"{self.ts}_conpad_r3_{self.a.scope}" / "pre_change_files" / max((r["path_under_package_FACT"] for r in self.rows), key=len)
        print(f"  info: longest backup destination = {len(plain(dst))} chars (over 260 => extended-length paths are required and used)")

    # ---------------- phases ----------------
    def preflight(self):
        print(f"== preflight (read-only) | scope {self.a.scope} ==")
        self.staging_checks()
        self.live_checks()

    def write_before_manifest(self, entries, name):
        lp(self.trace).mkdir(parents=True, exist_ok=True)
        path = self.trace / name
        with open(lp(path), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["relative_path", "size_bytes", "md5"])
            for label, p in entries:
                w.writerow([label, lp(p).stat().st_size, md5(p)])
        return path

    def verify_backup(self, mpath, root, name):
        res, good_all = [], True
        for r in read_csv(mpath):
            f = lp(root / r["relative_path"].replace("/", os.sep))
            good = f.is_file() and f.stat().st_size == int(r["size_bytes"]) and md5(f) == r["md5"]
            good_all &= good
            res.append({**r, "backup_path": plain(f), "backup_verified": good})
        with open(lp(self.trace / name), "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(res[0]))
            w.writeheader()
            w.writerows(res)
        return good_all, len(res)

    def tree_snapshot(self, sub):
        out = {}
        root = self.live / sub
        for dp, dn, fs in os.walk(lp(root)):
            for f in fs:
                full = Path(dp) / f
                out[plain(full.relative_to(lp(self.live))).replace(BS, "/")] = lp(full).stat().st_size
        return out

    def run(self):
        print(f"== run | scope {self.a.scope} | {'EXECUTE' if self.a.execute else 'DRY-RUN'} ==")
        self.staging_checks()
        self.live_checks()
        if self.a.execute:
            self.check(bool(self.a.jack_go.strip()), "--jack-go given (recorded in the log)")
        if self.fails:
            return self.stop()
        touched = [r for r in self.rows if r["action"] in ("REPLACE", "REMOVE", "ARCHIVE")]
        bdir = self.backup_root / f"{self.ts}_conpad_r3_{self.a.scope}"
        print(f"  plan: manifest + verified backup of {len(touched)} package files and {len(self.wb_partners)} workbooks under {plain(bdir)}, then apply, then verify")
        if not self.a.execute:
            return print("  (dry-run) nothing copied, moved or removed")
        entries = [(r["path_under_package_FACT"], self.live_path(r["path_under_package_FACT"])) for r in touched]
        entries += [(f"WORKBOOKS/{wb_rel(p)}", self.live / wb_rel(p)) for p in self.wb_partners]
        before = self.tree_snapshot("FACT")
        mpath = self.write_before_manifest(entries, f"R3_{self.a.scope}_manifest_BEFORE_CHANGES.csv")
        for label, src in entries:
            dst = lp(bdir / "pre_change_files" / label.replace("/", os.sep))
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(lp(src), dst)
        ok, n = self.verify_backup(mpath, bdir / "pre_change_files", f"R3_{self.a.scope}_backup_verification.csv")
        self.log("backup", "BACKUP_VERIFIED", "", plain(bdir), f"{n} files", ok and n == len(entries))
        if not (ok and n == len(entries)):
            self.fails.append("backup verification failed; nothing was changed")
            return self.stop()
        print(f"  backup verified: {n} files at {plain(bdir)}")
        self.log("backup", "PHASE_DONE")
        # ---- packages
        for r in touched:
            rel = r["path_under_package_FACT"]
            live = lp(self.live_path(rel))
            if r["action"] == "REPLACE":
                os.chmod(live, stat.S_IWRITE)  # a live file may carry ReadOnly
                shutil.copy2(lp(STAGE / rel), live)
                good = md5(live) == md5(STAGE / rel)
                self.log("packages", "REPLACED", plain(STAGE / rel), plain(live), md5(live), good)
            elif r["action"] == "REMOVE":
                os.chmod(live, stat.S_IWRITE)
                os.remove(live)
                good = not live.exists()
                self.log("packages", "REMOVED", plain(live), "", "", good)
            else:  # ARCHIVE = move into the archive folder inside the same Cluster_guide dir
                adir = live.parent / ARCHIVE_NAME
                adir.mkdir(parents=True, exist_ok=True)
                dest = adir / live.name
                if dest.exists():
                    self.fails.append(f"archive destination already exists: {plain(dest)}")
                    return self.stop()
                h = md5(live)
                shutil.move(str(live), str(dest))
                good = dest.is_file() and not live.exists() and md5(dest) == h
                self.log("packages", "ARCHIVED", plain(live), plain(dest), h, good)
            if not good:
                self.fails.append(f"post-action check failed for {rel} ({r['action']})")
                return self.stop()
        self.log("packages", "PHASE_DONE")
        print(f"  packages done: {sum(1 for r in touched if r['action'] == 'REPLACE')} replaced, {sum(1 for r in touched if r['action'] == 'ARCHIVE')} archived, {sum(1 for r in touched if r['action'] == 'REMOVE')} removed")
        # ---- workbooks
        for pt in self.wb_partners:
            live, src = lp(self.live / wb_rel(pt)), STAGE / "workbooks" / wb_rel(pt)
            os.chmod(live, stat.S_IWRITE)
            shutil.copy2(lp(src), live)
            good = md5(live) == md5(src)
            self.log("workbooks", "REPLACED", plain(src), plain(live), md5(live), good)
            if not good:
                self.fails.append(f"workbook copy check failed for {pt}")
                return self.stop()
        self.log("workbooks", "PHASE_DONE")
        print(f"  workbooks done: {len(self.wb_partners)} replaced")
        # ---- verify (uses the pre-change tree snapshot)
        self.verify(before)

    def verify(self, before=None):
        print(f"== verify (read-only) | scope {self.a.scope} ==")
        from cluster_exclusions import kml_active_ids_for_partner
        base = self.load_baseline()
        man = {r["path"]: r["md5"] for r in read_csv(MANIFEST)}
        bad = []
        for r in self.rows:
            rel, live = r["path_under_package_FACT"], lp(self.live_path(r["path_under_package_FACT"]))
            if r["action"] == "REPLACE" and not (live.is_file() and md5(live) == man[rel]):
                bad.append(rel + " (REPLACE)")
            elif r["action"] == "REMOVE" and live.exists():
                bad.append(rel + " (REMOVE)")
            elif r["action"] == "ARCHIVE":
                dest = live.parent / ARCHIVE_NAME / live.name
                if live.exists() or not dest.is_file() or (base and rel in base and not same_as_baseline(sig(dest), base[rel])):
                    bad.append(rel + " (ARCHIVE)")
            elif r["action"] == "IDENTICAL" and base and rel in base and not same_as_baseline(sig(live), base[rel]):
                bad.append(rel + " (IDENTICAL changed)")
        self.check(not bad, f"every change-list row in scope is in its expected end state ({len(self.rows)} rows)" + (f" | {bad[:4]}" if bad else ""))
        wbbad = [p for p in self.wb_partners if not (lp(self.live / wb_rel(p)).is_file() and md5(self.live / wb_rel(p)) == man[f"workbooks/{wb_rel(p)}"])]
        self.check(not wbbad, f"the {len(self.wb_partners)} changed workbooks are byte-identical to the staged ones" + (f" | {wbbad}" if wbbad else ""))
        idw = [p for p in WB_PARTNERS if p not in self.wb_partners and base and not same_as_baseline(sig(self.live / wb_rel(p)), base[wb_rel(p)])]
        self.check(not idw, "the 6 workbooks that were identical to staged are untouched" + (f" | {idw}" if idw else ""))
        # live KML point sets == WORKING for the LGAs in scope
        import csv as _c
        work = [r for r in _c.DictReader(open(lp(WORKING), encoding="utf-8-sig", newline=""))]
        name2pc = {(r["adm1_name"], r["adm2_name"]): r["adm2_pcode"] for r in work}
        lgas = sorted({(r["state"], r["lga"]) for r in self.rows})
        badk = []
        for st, lga in lgas:
            Kn, Ki = kml_active_ids_for_partner(str(lp(self.live / "FACT" / st / lga)))  # extended path: replicas/deep roots exceed 260 chars
            rows = [x for x in work if x["adm2_pcode"] == name2pc[(st, lga)] and x.get("sampling_method") != "MSNA Light"]
            Fn = {x["survey_id"] for x in rows if x["pop_type"] == "non_idp" and x["status"] in ("primary", "reserve")}
            Fi = {x["cluster_id"] for x in rows if x["pop_type"] == "idp" and x["status"] == "primary"}
            if Kn != Fn or Ki != Fi:
                badk.append(f"{lga} (KML {len(Kn)}/{len(Ki)} vs WORKING {len(Fn)}/{len(Fi)})")
            prim = {x["cluster_id"] for x in rows if x["status"] == "primary"}
            live_fs = {f[:-len("_factsheet.docx")] for dp, dn, fs in os.walk(lp(self.live / "FACT" / st / lga)) if ARCHIVE_NAME not in dp and "_archive" not in dp for f in fs if f.endswith("_factsheet.docx")}
            if not live_fs <= prim:
                badk.append(f"{lga}: live factsheets for clusters with no primary row: {sorted(live_fs - prim)[:3]}")
        self.check(not badk, f"live KML point sets equal WORKING for all {len(lgas)} LGAs in scope, and no live factsheet is for a cluster that left WORKING" + (f" | {badk[:3]}" if badk else ""))
        if before is not None:
            after = self.tree_snapshot("FACT")
            exp_removed = {("FACT/" + r["path_under_package_FACT"][5:]) for r in self.rows if r["action"] in ("REMOVE", "ARCHIVE")}
            exp_added = {("FACT/" + r["path_under_package_FACT"][5:]).rsplit("/", 1)[0] + f"/{ARCHIVE_NAME}/" + r["path_under_package_FACT"].rsplit("/", 1)[1] for r in self.rows if r["action"] == "ARCHIVE"}
            exp_changed = {r["path_under_package_FACT"] for r in self.rows if r["action"] == "REPLACE"}
            if "FACT" in self.wb_partners:
                exp_changed.add(wb_rel("FACT"))  # FACT's workbook sits at FACT/FACT_sampling_points_summary.xlsx
            removed, added = set(before) - set(after), set(after) - set(before)
            unexpected_changed = [k for k in set(before) & set(after) if before[k] != after[k] and k not in exp_changed]
            self.check(removed == exp_removed and added == exp_added and not unexpected_changed,
                       f"whole FACT tree: removed {len(removed)}, added {len(added)}, resized {len(unexpected_changed)} unexpected (expected {len(exp_removed)}/{len(exp_added)}/0)")
        if self.fails:
            return self.stop()
        self.log("verify", "PHASE_DONE")
        print("  VERIFIED")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("phase", choices=["baseline", "preflight", "run", "verify"])
    ap.add_argument("--scope", choices=["A", "A_B"], required=True)
    ap.add_argument("--execute", action="store_true")
    ap.add_argument("--jack-go", default="")
    ap.add_argument("--live-root", default=str(WS / "3. External coordination" / "NGA MSNA 2026 Package"))
    ap.add_argument("--backup-root", default=str(OUT / "_archive_partner_package_backups"))
    ap.add_argument("--log-dir", default=str(TRACE_DEFAULT))
    ap.add_argument("--trace-dir", default=str(TRACE_DEFAULT))
    a = ap.parse_args()
    r = Run(a)
    print(("EXECUTE" if a.execute else "DRY-RUN"), "|", a.phase, "| scope", a.scope, "| live root:", a.live_root)
    {"baseline": r.baseline, "preflight": r.preflight, "run": r.run, "verify": lambda: r.verify(None)}[a.phase]()
    if r.fails:
        r.stop()
    print("\nno failed checks" if not a.execute else "\ndone")


if __name__ == "__main__":
    main()
