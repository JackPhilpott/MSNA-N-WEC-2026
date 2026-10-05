# ==============================================================================
# 2026-10-04 STAGING ONLY: inputs for Jack's 4 Oct LGA reallocations (14 LGAs to
# FACT, Tangaza to ZOA; relayed by the Coordinator - live execution needs his own
# word in Resampling's window). Writes ONLY into
# resampling/output/staged_reallocation_2026-10-04/:
#   moves.csv     for reallocate_lga_coverage_generic_2026-10-04.R (frame patch)
#   credited.csv  per stratum: credited_at_move = min(canonical achieved, target_sample)
#   Partnerscoverage_REALLOC_2026-10-04_STAGED.xlsx  the coverage workbook with the
#                 moves applied (the live workbook is untouched), cell-diffed and
#                 re-parsed exactly the way build_partner_dc_packages.py parses it
#   README.md     what was staged and from which inputs
# Usage: python stage_reallocation_2026-10-04.py [--shared "Borno|Bama|PLAN|FACT" ...]
#   --shared adds a joint-coverage LGA (Jack's "shared between both FACT and PLAN"
#   caveat) - only once he names the LGAs; none of the 15 below is PLAN's today.
# ==============================================================================
import argparse
import csv
import hashlib
import os
import shutil
import sys
from collections import defaultdict

import openpyxl


def _msna_shared_dir():
    for start in (os.environ.get("MSNA_WORKSPACE", "").strip(), os.path.dirname(os.path.abspath(__file__)), os.getcwd()):
        d = os.path.abspath(start) if start else ""
        while d:
            cand = os.path.join(d, "1_sampling", "scripts", "shared")
            if os.path.isfile(os.path.join(cand, "msna_paths.py")):
                return cand
            parent = os.path.dirname(d)
            d = "" if parent == d else parent
    raise SystemExit("Cannot find 1_sampling/scripts/shared/msna_paths.py - set MSNA_WORKSPACE.")


sys.path.insert(0, _msna_shared_dir())
import msna_paths  # noqa: E402

S = msna_paths.sampling_dir()
M = msna_paths.monitoring_dir()
OUT = os.path.join(S, "resampling", "output", "staged_reallocation_2026-10-04")
EXCEL = os.path.join(S, "input_data", "boundaries", "partner_coverage", "Partnerscoverage.xlsx")
STRATA_FULL = os.path.join(S, "output", "data", "data_collection", "NGA_MSNA_2026_strata_level_sampling_frame_v14_FULL.csv")
SUBS = os.path.join(M, "data", "real_submissions.csv")
OVERLAY = os.path.join(M, "data", "CONFIRMED_DELETIONS_OVERLAY.csv")
DECIDED_BY = "Jack Philpott, 2026-10-04 ~20:45 (Coordinator msna-n-wec-2026-7b window); executed only on his own word in Resampling's window"
REASON = ("Big resampling night 2026-10-04: LGA handed to {to} for Round 2 collection{field}; credit/target/Still Needed follow the "
          "current LGA owner (same rule as the 25 Sep moves)")
# the field note is FACT's (Lina's 2 Oct email); it must not be written onto the ZOA move
FIELD_NOTE = {"FACT": " (FACT took the LGAs over in the field from 2 Oct, using the outgoing partner's KMLs, per Lina CAMPEROS's email of 2 Oct)"}
SHARED_REASON = ("Big resampling night 2026-10-04: LGA now SHARED between {existing} and {added} (Jack D5: both still "
                 "continuing data collection); {existing} keeps its coverage, {added} is added alongside")
# (sheet, state, LGA) -> (frame tokens being replaced (';'-separated), coverage-sheet column(s) to clear, new owner)
MOVES = [
    ("NE", "Adamawa", "Madagali", "Street Child of Nigeria", ["Street Child of Nigeria"], "FACT"),
    ("NE", "Borno", "Mafa", "FHI 360", ["FHI 360"], "FACT"),
    ("NE", "Borno", "Mobbar", "FHI 360", ["FHI 360"], "FACT"),
    ("NE", "Borno", "Monguno", "Street Child of Nigeria", ["Street Child of Nigeria"], "FACT"),
    ("NW", "Katsina", "Mashi", "Malteser", ["Malteser"], "FACT"),
    ("NW", "Katsina", "Batagarawa", "IRC", ["IRC"], "FACT"),
    ("NW", "Katsina", "Jibia", "IRC", ["IRC"], "FACT"),
    ("NW", "Kebbi", "Zuru", "IRC; LHI", ["IRC/LHI"], "FACT"),
    ("NW", "Kebbi", "Gwandu", "Solidarités", ["Solidarités"], "FACT"),
    ("NW", "Kebbi", "Birnin Kebbi", "CARE", ["CARE"], "FACT"),
    ("NW", "Kebbi", "Suru", "CARE", ["CARE"], "FACT"),
    ("NW", "Kebbi", "Argungu", "CARE", ["CARE"], "FACT"),
    ("NW", "Zamfara", "Tsafe", "IRC", ["IRC"], "FACT"),
    ("NW", "Zamfara", "Anka", "Solidarités", ["Solidarités"], "FACT"),
    ("NW", "Sokoto", "Tangaza", "IRC; LHI", ["IRC/LHI"], "ZOA"),
]
SPLITS = {"IRC/LHI": ["IRC", "LHI"]}  # build_partner_dc_packages.py's COMBINED_PARTNER_SPLITS


def md5(p):
    return hashlib.md5(open(p, "rb").read()).hexdigest()


def parse_coverage(path):
    """{(state, lga): set(partners)} exactly as build_partner_dc_packages.py section 2 reads it."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        cidx = header.index("COUNT")
        for r in rows[1:]:
            if r[2] is None:
                continue
            ps = set()
            for i in range(3, cidx):
                if r[i]:
                    name = str(header[i]).strip()
                    ps.update(SPLITS.get(name, [name]))
            out[(str(r[1]).strip(), str(r[2]).strip())] = ps
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--shared", action="append", default=[], help='"State|LGA|EXISTING|ADDED", e.g. "Borno|Bama|PLAN|FACT"')
    args = ap.parse_args()
    moves = list(MOVES)
    for s in args.shared:
        st, lga, existing, added = [x.strip() for x in s.split("|")]
        moves.append(("NE" if st in ("Borno", "Adamawa", "Yobe") else "NW", st, lga, existing, [], f"{existing}, {added}"))
    os.makedirs(OUT, exist_ok=True)

    # ---- 1. moves.csv (frame patch input) ----
    with open(os.path.join(OUT, "moves.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["adm1_name", "adm2_name", "from", "to", "reason", "decided_by"])
        for sheet, st, lga, frm, _, to in moves:
            reason = (SHARED_REASON.format(existing=frm, added=to.split(",")[-1].strip()) if "," in to
                      else REASON.format(to=to, field=FIELD_NOTE.get(to, "")))
            w.writerow([st, lga, frm, to, reason, DECIDED_BY])

    # ---- 2. credited.csv: min(canonical achieved, target_sample) per stratum of the moved LGAs ----
    with open(OVERLAY, encoding="utf-8-sig", newline="") as f:
        dele = {r["uuid"] for r in csv.DictReader(f) if r["status"] in ("confirmed", "contested")}
    ach = defaultdict(int)
    with open(SUBS, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if r["interview_outcome"] == "completed" and r["matched_survey_id"] not in ("", "NA") and r["submission_uuid"] not in dele:
                ach[r["matched_strata_id"]] += 1
    basis = (f"canonical achieved (completed, matched, not a confirmed/contested deletion) by matched_strata_id from real_submissions "
             f"{md5(SUBS)[:8]} + overlay {md5(OVERLAY)[:8]}, min(achieved, target_sample); computed when staged")
    moved = {(st, lga) for _, st, lga, *_ in moves}
    with open(STRATA_FULL, encoding="utf-8-sig", newline="") as f:
        strata = [r for r in csv.DictReader(f) if (r["adm1_name"], r["adm2_name"]) in moved]
    with open(os.path.join(OUT, "credited.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["strata_id", "credited_at_move", "credited_basis"])
        for r in strata:
            t = float(r["target_sample"]) if r["target_sample"] not in ("", "NA") else 0
            w.writerow([r["strata_id"], int(min(ach.get(r["strata_id"], 0), t)), basis])
    print(f"moves: {len(moves)} LGAs | strata: {len(strata)} | credited total {sum(min(ach.get(r['strata_id'], 0), float(r['target_sample'] or 0)) for r in strata):.0f}")

    # ---- 3. staged coverage workbook ----
    staged = os.path.join(OUT, "Partnerscoverage_REALLOC_2026-10-04_STAGED.xlsx")
    shutil.copy2(EXCEL, staged)
    wb = openpyxl.load_workbook(staged)
    expected_cells = set()
    for sheet, st, lga, _, clear_cols, to in moves:
        ws = wb[sheet]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        row = next((r for r in range(2, ws.max_row + 1)
                    if str(ws.cell(r, 2).value or "").strip() == st and str(ws.cell(r, 3).value or "").strip() == lga), None)
        if row is None:
            raise SystemExit(f"STOP: {st}/{lga} not found on sheet {sheet}")
        for col_name in clear_cols:
            c = header.index(col_name) + 1
            if not ws.cell(row, c).value:
                raise SystemExit(f"STOP: {st}/{lga} has no '{col_name}' mark to clear")
            ws.cell(row, c).value = None
            expected_cells.add((sheet, row, c))
        for owner in [x.strip() for x in to.split(",")]:
            if owner not in header:
                raise SystemExit(f"STOP: no '{owner}' column on sheet {sheet}")
            c = header.index(owner) + 1
            if not ws.cell(row, c).value:
                ws.cell(row, c).value = owner
                expected_cells.add((sheet, row, c))
    wb.save(staged)

    # cell diff: exactly the expected cells changed (array formulas compare by object
    # identity in openpyxl, so compare their range + text instead)
    from openpyxl.worksheet.formula import ArrayFormula

    def cv(cell):
        v = cell.value
        return ("ARRAY", v.ref, v.text) if isinstance(v, ArrayFormula) else v

    a, b = openpyxl.load_workbook(EXCEL), openpyxl.load_workbook(staged)
    diff = set()
    for s in a.sheetnames:
        for r in range(1, max(a[s].max_row, b[s].max_row) + 1):
            for c in range(1, max(a[s].max_column, b[s].max_column) + 1):
                if cv(a[s].cell(r, c)) != cv(b[s].cell(r, c)):
                    diff.add((s, r, c))
    if diff != expected_cells:
        raise SystemExit(f"STOP: staged workbook differs in unexpected cells: {sorted(diff ^ expected_cells)[:10]}")
    # parse check: only the moved LGAs' partner sets change, each to exactly its new owner(s)
    before, after = parse_coverage(EXCEL), parse_coverage(staged)
    changed = {k for k in set(before) | set(after) if before.get(k) != after.get(k)}
    want = {(st, lga): {x.strip() for x in to.split(",")} for _, st, lga, _, _, to in moves}
    bad = [k for k in changed if k not in want] + [k for k, v in want.items() if after.get(k) != v]
    if bad:
        raise SystemExit(f"STOP: parsed coverage is wrong for {bad[:10]}")
    print(f"staged workbook: {len(diff)} cells changed (exactly as expected); parsed coverage changes only the {len(changed)} moved LGAs:")
    for k in sorted(changed):
        print(f"   {k[0]:8s} {k[1]:13s} {sorted(before.get(k, []))} -> {sorted(after.get(k, []))}")
    losing = sorted({p for k in changed for p in before.get(k, set()) - after.get(k, set())})
    still = {p: sorted(lga for (st, lga), ps in after.items() if p in ps) for p in losing}
    print("outgoing partners' remaining LGAs after the moves:", {p: len(v) for p, v in still.items()})

    with open(os.path.join(OUT, "README.md"), "w", encoding="utf-8") as f:
        f.write(f"# Staged: 4 Oct 2026 reallocations ({len(moves)} LGAs). NOTHING here is live.\n\n"
                f"- `moves.csv` + `credited.csv` -> `scripts/one_off_analyses/reallocate_lga_coverage_generic_2026-10-04.R <moves> <credited> dry|execute <stage_dir> [snapshot_dir]`\n"
                f"- `Partnerscoverage_REALLOC_2026-10-04_STAGED.xlsx`: live workbook {md5(EXCEL)[:8]} with {len(diff)} cells changed (cell-diff + parse verified). "
                f"Goes to all copies (1_sampling input_data, 2_monitoring input_data) only on Jack's go.\n"
                f"- Credited basis: {basis}\n- Decided by: {DECIDED_BY}\n")
    print("staged in", OUT)


if __name__ == "__main__":
    main()
