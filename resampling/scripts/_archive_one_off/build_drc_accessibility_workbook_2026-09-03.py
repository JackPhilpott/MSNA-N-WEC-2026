# ==============================================================================
# Builds DRC's first-ever accessibility report workbook, 2026-09-03 - same
# standard template shape (README + Ward Accessibility + Cluster
# Accessibility) as every other partner, pre-filled with what we already
# know from DRC's "inaccessible location.xlsx" submission and the live
# sampling frame. See build_drc_accessibility_report_2026-09-03.R (+
# _step2.R) for the reconciliation logic.
# ==============================================================================
import csv

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
WARD_CSV = PROJECT_DIR + r"\resampling\output\resample_runs\DRC\2026-09-03\drc_accessibility_ward_reconciled.csv"
CLUSTER_CSV = PROJECT_DIR + r"\resampling\output\resample_runs\DRC\2026-09-03\drc_accessibility_cluster_reconciled.csv"
OUT_PATH = PROJECT_DIR + r"\resampling\output\resample_runs\DRC\2026-09-03\DRC_accessibility_report_2026-09-03.xlsx"

HEADER_FILL_REF = PatternFill("solid", fgColor="1B2A4A")
HEADER_FILL_INPUT = PatternFill("solid", fgColor="2C5F8A")
NEEDS_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_COLUMNS = {"Accessible (Y/N)", "Reason category", "Reason notes"}
ACCESSIBLE_OPTIONS = ["Yes", "No"]
REASON_OPTIONS = [
    "N/A - fully accessible", "Insecurity / conflict", "Physical access (terrain, flooding, roads)",
    "Population absent / relocated", "Building-footprint issue (no eligible structures)",
    "Access denied by authorities / community", "Other",
]

WARD_COLUMNS = ["State", "LGA", "Ward (GRID3)", "Ward (OCHA/COD)", "Non-IDP clusters", "IDP clusters",
                "Total target HHs (primary)", "Accessible (Y/N)", "Reason category", "Reason notes",
                "Needs your input?", "Why flagged"]
CLUSTER_COLUMNS = ["State", "LGA", "Ward (GRID3)", "Pop Type", "Cluster ID", "IDP Category",
                    "Target HHs (primary)", "Reserve HHs", "Accessible (Y/N)", "Reason category",
                    "Reason notes", "Needs your input?", "Why flagged"]


def read_rows(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def is_true(v):
    return str(v).strip().upper() in ("TRUE", "1")


def write_sheet(wb, sheet_name, columns, rows, row_map):
    ws = wb.create_sheet(sheet_name)
    ws.append(columns)
    for c, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL_INPUT if col_name in INPUT_COLUMNS else HEADER_FILL_REF

    n_flagged = 0
    for row in rows:
        ws.append(row_map(row))
        r = ws.max_row
        if is_true(row.get("needs_attention", "")):
            n_flagged += 1
            for c in range(1, len(columns) + 1):
                ws.cell(row=r, column=c).fill = NEEDS_INPUT_FILL

    n_rows = len(rows)
    access_col = columns.index("Accessible (Y/N)") + 1
    reason_col = columns.index("Reason category") + 1
    dv_access = DataValidation(type="list", formula1=f'"{",".join(ACCESSIBLE_OPTIONS)}"', allow_blank=True)
    dv_reason = DataValidation(type="list", formula1=f'"{",".join(REASON_OPTIONS)}"', allow_blank=True)
    ws.add_data_validation(dv_access)
    ws.add_data_validation(dv_reason)
    dv_access.add(f"{get_column_letter(access_col)}2:{get_column_letter(access_col)}{n_rows + 1}")
    dv_reason.add(f"{get_column_letter(reason_col)}2:{get_column_letter(reason_col)}{n_rows + 1}")

    ws.freeze_panes = "A2"
    for i, col_name in enumerate(columns, start=1):
        width = 40 if col_name in ("Why flagged", "Reason notes") else max(13, min(26, len(col_name) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width
    if n_rows:
        tbl = Table(displayName=sheet_name.replace(" ", ""), ref=f"A1:{get_column_letter(len(columns))}{n_rows + 1}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    return n_flagged


def main():
    ward_rows = read_rows(WARD_CSV)
    cluster_rows = read_rows(CLUSTER_CSV)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("README")
    readme.column_dimensions["A"].width = 108
    n_ward_flag = sum(1 for r in ward_rows if is_true(r.get("needs_attention")))
    n_cluster_flag = sum(1 for r in cluster_rows if is_true(r.get("needs_attention")))
    lines = [
        ("DRC - NGA MSNA 2026 - Accessibility report (first report)", True, 14),
        ("", False, 11),
        ("WHAT THIS IS", True, 12),
        (f"Your full assignment - {len(ward_rows)} wards, {len(cluster_rows)} clusters across Illela, Isa, Sokoto "
         "North, Sokoto South, and Kaura Namoda. This is the same accessibility report template every partner "
         "gets - we haven't received one from your team before, so we've pre-filled it using the sampling frame's "
         "current status plus the 'inaccessible location.xlsx' file your team sent 3 September.", False, 11),
        ("", False, 11),
        ("WHAT WE'VE ALREADY APPLIED FOR YOU", True, 12),
        ("Kaura Namoda: the 4 sites your team reported inaccessible are marked No / Insecurity here.", False, 11),
        ("Illela and Sokoto South: your team flagged specific households (never a whole cluster) - these clusters "
         "are left as Accessible, with a note on how many households were flagged and why.", False, 11),
        ("Isa: your submission said 'All location' with no further detail. We haven't marked anything inaccessible "
         "off that alone, since Isa is also covered by IRC and LHI and a whole-LGA claim needs more detail before "
         "we can act on it.", False, 11),
        ("", False, 11),
        ("WHAT WE NEED FROM YOU", True, 12),
        (f"Only {n_ward_flag} of {len(ward_rows)} wards ({n_cluster_flag} of {len(cluster_rows)} clusters) are "
         "shaded amber - everything else can be left as-is for now. Three things specifically:", False, 11),
        ("1. Kaura Namoda (4 sites) - were you able to reach target using reserve households, or do these need "
         "supplementary/replacement clusters?", False, 11),
        ("2. Illela / Sokoto South - please confirm the flagged households are genuinely specific to those "
         "households, not the wider ward.", False, 11),
        ("3. Isa - please tell us specifically which wards/clusters are affected and why, so we can apply this "
         "properly rather than as a blanket LGA-wide assumption.", False, 11),
        ("", False, 11),
        ("A NOTE ON THE DATA RECOVERY WORKBOOK YOU ALREADY HAVE", True, 12),
        ("This is a separate document from DRC_data_recovery_workbook_2026-08-30.xlsx (the one covering your 46 "
         "interviews needing input and the confirmed deletions) - that one's still open and needs a response "
         "too. This file is specifically about location/ward accessibility, not individual interviews.", False, 11),
    ]
    for i, (text, bold, size) in enumerate(lines, start=1):
        cell = readme.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=size)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if len(text) > 90:
            readme.row_dimensions[i].height = 16 + 14 * (len(text) // 90)

    def ward_map(r):
        return [r["State"], r["LGA"], r["Ward (GRID3)"], r.get("Ward (OCHA/COD)") or None,
                int(float(r["Non-IDP clusters"])) if r.get("Non-IDP clusters") not in (None, "", "NA") else None,
                int(float(r["IDP clusters"])) if r.get("IDP clusters") not in (None, "", "NA") else None,
                int(float(r["Total target HHs (primary)"])) if r.get("Total target HHs (primary)") not in (None, "", "NA") else None,
                r.get("Accessible") or None, r.get("Reason") or None, r.get("Notes") or None,
                "Yes" if is_true(r.get("needs_attention")) else "", r.get("flag_reason") or ""]

    def cluster_map(r):
        return [r["State"], r["LGA"], r["Ward (GRID3)"], r["Pop Type"], r["Cluster ID"], r.get("IDP Category") or None,
                int(float(r["Target HHs (primary)"])) if r.get("Target HHs (primary)") not in (None, "", "NA") else None,
                int(float(r["Reserve HHs"])) if r.get("Reserve HHs") not in (None, "", "NA") else None,
                r.get("Accessible (Y/N)") or None, r.get("Reason category") or None, r.get("Reason notes") or None,
                "Yes" if is_true(r.get("needs_attention")) else "", r.get("flag_reason") or ""]

    n1 = write_sheet(wb, "Ward Accessibility", WARD_COLUMNS, ward_rows, ward_map)
    n2 = write_sheet(wb, "Cluster Accessibility", CLUSTER_COLUMNS, cluster_rows, cluster_map)

    wb.move_sheet("README", offset=-len(wb.sheetnames))
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")
    print(f"Ward Accessibility: {len(ward_rows)} rows, {n1} flagged")
    print(f"Cluster Accessibility: {len(cluster_rows)} rows, {n2} flagged")


if __name__ == "__main__":
    main()
