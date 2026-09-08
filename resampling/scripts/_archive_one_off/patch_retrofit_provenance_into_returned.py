# ==============================================================================
# One-time, git-committed retrofit: adds the "Reported by (Partner /
# IMPACT-default)" and "Source channel" columns (added 2026-08-27 to
# 01_generate_accessibility_reports.py, see that script's PROVENANCE_COLUMNS)
# to every already-returned partner file in accessibility_reports_returned/ -
# every one of them predates the column addition, so none currently have it.
#
# Values are populated by looking up each row's classification from the
# ALREADY-MIGRATED master log (resampling_requests_log.csv's reported_by/
# source_channel, see patch_add_reported_by_column.py) - not re-derived here,
# so this stays consistent with the master log as the single source of
# truth. A row the master log still has as "" (needs review) stays blank
# here too - never guessed.
#
# Rebuilds each sheet fully (columns + Table + DataValidation) rather than
# appending in place, matching the technique write_input_sheet() established
# in generate_fact_cleaned_listing_with_issues.py after the 2026-08-26 bug
# where copying cell values alone silently dropped DataValidation - the same
# risk applies here, so the same fix pattern is reused.
#
# Only touches the canonical `<Partner>_accessibility_report.xlsx` files -
# dated snapshot copies (e.g. IMC_accessibility_report_2026-08-23.xlsx) are
# a deliberate audit-trail record per resampling/README.md and are left
# untouched, matching that convention.
# ==============================================================================
import csv
import glob
import os

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
RETURNED_DIR = PROJECT_DIR + r"\resampling\input\accessibility_reports_returned"
LOG_CSV = PROJECT_DIR + r"\resampling\output\resampling_requests_log.csv"

REPORTED_BY_COL = "Reported by (Partner / IMPACT-default)"
SOURCE_CHANNEL_COL = "Source channel"
REPORTED_BY_OPTIONS = ["Partner", "IMPACT (default - accessible until reported otherwise)"]
SOURCE_CHANNEL_OPTIONS = [
    "Partner's own template", "Email", "WhatsApp/verbal (coordinator-transcribed)",
    "Point-level file annotation", "N/A",
]
INPUT_COLUMNS = {
    "Accessible (Y/N)", "Reason category", "Reason notes", "% of target achieved so far", "Date reported",
    REPORTED_BY_COL, SOURCE_CHANNEL_COL,
}


def load_provenance_lookup():
    """{(partner, report_level, state, lga, ward_name, cluster_id): (reported_by, source_channel)}
    from the LATEST (highest request_id) log row per key - mirrors 02_ingest's
    own latest_by_key() logic, so a superseded older classification never
    wins over a partner's more recent update."""
    with open(LOG_CSV, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    latest = {}
    for r in rows:
        key = (r["partner"], r["report_level"], r["state"], r["lga"], r["ward_name"], r["cluster_id"])
        if key not in latest or int(r["request_id"]) > int(latest[key]["request_id"]):
            latest[key] = r
    return {k: (r["reported_by"], r["source_channel"]) for k, r in latest.items()}


def rebuild_sheet(wb, old_ws, sheet_name, table_name, extra_lookup_fn):
    raw_headers = [c.value for c in old_ws[1]]
    # Trim trailing None header cells - Excel formatting bleeding into extra
    # empty columns beyond the real header row (found on FACT's file, whose
    # used range extends to column Z though only 12 columns have headers;
    # a known-quirky file already, see this workflow's other FACT patches).
    n_real = len(raw_headers)
    while n_real > 0 and raw_headers[n_real - 1] is None:
        n_real -= 1
    old_headers = raw_headers[:n_real]
    old_idx = {h: i for i, h in enumerate(old_headers)}

    if REPORTED_BY_COL in old_idx:
        # Idempotency guard: this sheet was already retrofitted on a prior
        # (possibly interrupted) run - old_headers would otherwise already
        # contain the new columns and this function would double-append them.
        print(f"    ({sheet_name} already has the provenance columns, skipping)")
        return sum(1 for _ in old_ws.iter_rows(min_row=2, values_only=True)), None

    old_rows = [list(row)[:n_real] for row in old_ws.iter_rows(min_row=2, values_only=True)]

    new_headers = old_headers + [REPORTED_BY_COL, SOURCE_CHANNEL_COL]
    del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(new_headers)
    header_fill_ref = openpyxl.styles.PatternFill("solid", fgColor="1B2A4A")
    header_fill_input = openpyxl.styles.PatternFill("solid", fgColor="2C5F8A")
    for c, col_name in enumerate(new_headers, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = header_fill_input if col_name in INPUT_COLUMNS else header_fill_ref

    n_classified = 0
    for row in old_rows:
        accessible = row[old_idx["Accessible (Y/N)"]] if "Accessible (Y/N)" in old_idx else None
        reason = row[old_idx["Reason category"]] if "Reason category" in old_idx else None
        reported_by, source_channel = ("", "")
        if accessible or reason:
            reported_by, source_channel = extra_lookup_fn(row, old_idx)
            if reported_by:
                n_classified += 1
        ws.append(row + [reported_by, source_channel])

    n_rows = len(old_rows)
    accessible_col = new_headers.index("Accessible (Y/N)") + 1
    reason_col = new_headers.index("Reason category") + 1
    reported_by_col = new_headers.index(REPORTED_BY_COL) + 1
    source_channel_col = new_headers.index(SOURCE_CHANNEL_COL) + 1

    validations = [
        (accessible_col, ["Yes", "No"]),
        (reason_col, ["N/A - fully accessible", "Insecurity / conflict", "Physical access (terrain, flooding, roads)",
                       "Population absent / relocated", "Building-footprint issue (no eligible structures)",
                       "Access denied by authorities / community", "Other"]),
        (reported_by_col, REPORTED_BY_OPTIONS),
        (source_channel_col, SOURCE_CHANNEL_OPTIONS),
    ]
    for col, options in validations:
        dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{openpyxl.utils.get_column_letter(col)}2:{openpyxl.utils.get_column_letter(col)}{n_rows + 1}")

    ws.freeze_panes = "A2"
    for i, col_name in enumerate(new_headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(30, len(col_name) + 4))

    if n_rows:
        tbl = Table(displayName=table_name, ref=f"A1:{openpyxl.utils.get_column_letter(len(new_headers))}{n_rows + 1}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)

    return n_rows, n_classified


def retrofit_file(path, provenance):
    partner = os.path.basename(path)[: -len("_accessibility_report.xlsx")]
    wb = openpyxl.load_workbook(path, data_only=True)
    summary = []

    if "Ward Accessibility" in wb.sheetnames:
        def lookup_ward(row, idx):
            key = (partner, "ward", row[idx["State"]] or "", row[idx["LGA"]] or "", row[idx["Ward (GRID3)"]] or "", "")
            return provenance.get(key, ("", ""))
        n, n_c = rebuild_sheet(wb, wb["Ward Accessibility"], "Ward Accessibility", "WardAccessibility", lookup_ward)
        summary.append(f"Ward: {n_c}/{n} rows classified" if n_c is not None else f"Ward: {n} rows (unchanged)")

    if "Cluster Accessibility" in wb.sheetnames:
        def lookup_cluster(row, idx):
            # Cluster Accessibility carries its own "Ward (GRID3)" column
            # too (see 01_generate's CLUSTER_COLUMNS) - 02_ingest logs it as
            # part of the key, it is NOT blank for cluster-level rows (only
            # cluster_id/pop_type are blank on WARD-level rows, the reverse
            # isn't true) - matching that exactly here.
            key = (partner, "cluster", row[idx["State"]] or "" if "State" in idx else "",
                   row[idx["LGA"]] or "" if "LGA" in idx else "",
                   row[idx["Ward (GRID3)"]] or "" if "Ward (GRID3)" in idx else "",
                   row[idx["Cluster ID"]] or "" if "Cluster ID" in idx else "")
            return provenance.get(key, ("", ""))
        n, n_c = rebuild_sheet(wb, wb["Cluster Accessibility"], "Cluster Accessibility", "ClusterAccessibility", lookup_cluster)
        summary.append(f"Cluster: {n_c}/{n} rows classified" if n_c is not None else f"Cluster: {n} rows (unchanged)")

    wb.save(path)
    print(f"  {partner}: {', '.join(summary)}")


def main():
    provenance = load_provenance_lookup()
    files = sorted(glob.glob(os.path.join(RETURNED_DIR, "*_accessibility_report.xlsx")))
    print(f"Retrofitting {len(files)} returned file(s)...")
    for path in files:
        try:
            retrofit_file(path, provenance)
        except PermissionError as e:
            print(f"  WARNING: {os.path.basename(path)} appears open/locked, skipped: {e}")
    print("\nDONE.")


if __name__ == "__main__":
    main()
