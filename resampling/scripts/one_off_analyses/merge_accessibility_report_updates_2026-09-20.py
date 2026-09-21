# ==============================================================================
# Tonight's (2026-09-20) partner accessibility-report prep batch: IRC, IMC,
# CRS, and Street Child of Nigeria sent full updated reports (raw comms,
# _2009 suffix - full "Assigned to you: N wards, M clusters" resubmissions,
# not partial deltas). FACT sent a small 17-row Cluster Accessibility-only
# delta (Dandume/Zurmi/Shinkafi, "_update_18_Sept_2026_2009"). Save the
# Children's update and FACT's Dikwa update were already placed directly in
# accessibility_reports_returned/ by Jack - nothing to merge for those two.
#
# Same merge_sheet() logic as merge_accessibility_report_updates_2026-09-16.py
# (copied, not imported - this project's standalone-script convention):
# MERGE, never overwrite. A non-blank raw value always wins (a partner's
# newer report reflects their current answer); a blank raw cell never
# overwrites an existing non-blank value in the main file - this correctly
# preserves both the original partner answers already on record AND
# tonight's earlier GIS-eligible-ward-append fix (those new zero-cluster
# rows have blank Accessible/Reason cells, so they're only touched if the
# partner's raw file happens to have a real answer for that exact ward - in
# which case their answer correctly fills in what was a placeholder).
#
# Backs up each partner's main file first, to accessibility_reports_returned/
# _archive/, per this project's standing convention.
# ==============================================================================
import shutil
from datetime import date

import openpyxl

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling\resampling"
RETURNED_DIR = PROJECT_DIR + r"\input\accessibility_reports_returned"
ARCHIVE_DIR = RETURNED_DIR + r"\_archive"
RAW_DIR = PROJECT_DIR + r"\input\partner_raw_comms"
TODAY = "2026-09-20"

UPDATE_COLS = ["Accessible (Y/N)", "Reason category", "Reason notes", "% of target achieved so far", "Date reported"]


def backup(main_path, partner, reason):
    bak = f"{ARCHIVE_DIR}\\{partner}_accessibility_report_pre_{reason}_{TODAY}.xlsx"
    shutil.copy2(main_path, bak)
    print(f"  backed up -> {bak}")


def header_idx(ws):
    header = [c.value for c in ws[1]]
    return {h: i for i, h in enumerate(header) if h}


def merge_sheet(ws_main, idx_main, ws_raw, idx_raw, key_cols):
    """Layer raw's non-blank UPDATE_COLS values onto ws_main's matching rows
    (matched by key_cols). Never blanks out an existing value; a non-blank
    raw value always wins."""
    raw_rows = {}
    for r in ws_raw.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in r):
            continue
        key = tuple(r[idx_raw[c]] for c in key_cols)
        raw_rows[key] = r

    n_updated = 0
    n_unmatched = 0
    for row_cells in ws_main.iter_rows(min_row=2):
        vals = [c.value for c in row_cells]
        if not any(v is not None for v in vals):
            continue
        key = tuple(vals[idx_main[c]] for c in key_cols)
        rr = raw_rows.pop(key, None)
        if rr is None:
            continue
        changed = False
        for col in UPDATE_COLS:
            if col not in idx_raw or col not in idx_main:
                continue
            new_val = rr[idx_raw[col]]
            if new_val is None or new_val == "":
                continue
            cell = row_cells[idx_main[col]]
            if cell.value != new_val:
                cell.value = new_val
                changed = True
        if changed:
            n_updated += 1
    n_unmatched = len(raw_rows)
    return n_updated, n_unmatched, list(raw_rows.keys())


def merge_partner(partner, raw_filename, raw_sheet_map=None, reason="0920_merge"):
    main_path = f"{RETURNED_DIR}\\{partner}_accessibility_report.xlsx"
    raw_path = f"{RAW_DIR}\\{partner}\\{raw_filename}"
    print(f"\n=== {partner} ===")
    backup(main_path, partner, reason)

    wb_main = openpyxl.load_workbook(main_path)
    wb_raw = openpyxl.load_workbook(raw_path, data_only=True)

    sheet_map = raw_sheet_map or {"Ward Accessibility": "Ward Accessibility", "Cluster Accessibility": "Cluster Accessibility"}
    for main_sheet, raw_sheet in sheet_map.items():
        ws_main = wb_main[main_sheet]
        ws_raw = wb_raw[raw_sheet]
        idx_main = header_idx(ws_main)
        idx_raw = header_idx(ws_raw)
        key_cols = ["Cluster ID"] if "Cluster ID" in idx_main else ["State", "LGA", "Ward (GRID3)"]
        n, n_unmatched, unmatched_keys = merge_sheet(ws_main, idx_main, ws_raw, idx_raw, key_cols)
        print(f"  {main_sheet}: {n} row(s) updated, {n_unmatched} raw row(s) had no matching key in the main file")
        if unmatched_keys:
            print(f"    unmatched keys (first 10): {unmatched_keys[:10]}")

    wb_main.save(main_path)
    print(f"  saved -> {main_path}")
    return main_path


if __name__ == "__main__":
    merge_partner("IRC", "IRC_accessibility_report_2009.xlsx")
    merge_partner("IMC", "IMC_accessibility_report_2009.xlsx")
    merge_partner("CRS", "CRS_accessibility_report_2009.xlsx")
    merge_partner("Street Child of Nigeria", "Street Child of Nigeria_accessibility_report_2009.xlsx")
    merge_partner("FACT", "FACT_accessibility_report_update_18_Sept_2026_2009.xlsx",
                  raw_sheet_map={"Cluster Accessibility": "Cluster Accessibility"}, reason="0920_fact_book_merge")

    print("\nAll merges done. Save the Children and FACT's Dikwa update were already placed directly in "
          "the returned folder - nothing to merge for those.")
