# ==============================================================================
# 2026-09-21 resampling round, small-partner batch: stages every remaining
# partner end-to-end (extract shortfalls -> scope filter -> dedupe dual-
# coverage strata -> Non-IDP + IDP draws -> ward-status stamp -> post-draw
# MoE projection), one partner at a time, sequentially. Nothing here merges
# - each partner's staging dir is left for Jack's review, same as FACT/NRC/
# CRS/CARE/COOPI/INTERSOS earlier tonight.
#
# Scope (Jack, tonight): closeable + exceeds-pool strata; "no remaining
# pool" strata dropped (nothing to draw). Unfiltered extracts kept as
# *_ALL_unfiltered.csv.
#
# Dual-coverage rule (Jack, tonight): Zuru + Tangaza (IRC, LHI) and Isa
# (DRC, IRC, LHI) are covered by more than one partner. extract_partner_
# shortfalls.R greps the partner name inside "Partners covering", so each
# shared stratum lands in EVERY covering partner's list and would be drawn
# more than once. Each shared stratum is drawn ONCE, under IRC (attributed
# there at merge); it is removed from every other partner's CSV. LHI ends
# up with nothing of its own and is skipped.
#
# Run only when nothing else is reading/writing the FULL frame (the draws
# read FULL for their already-used checks).
# Usage: python stage_small_partners_2026-09-21.py
# ==============================================================================
import csv
import os
import re
import subprocess
import sys
from collections import defaultdict

import openpyxl

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
SCRIPTS = os.path.join(PROJECT_DIR, "resampling", "scripts")
RUNS = os.path.join(PROJECT_DIR, "resampling", "output", "resample_runs")
WORKBOOK = os.path.join(PROJECT_DIR, "resampling", "output", "NGA_MSNA_2026_accessibility_impact_workbook.xlsx")
RSCRIPT = r"C:\Users\JackPHILPOTT\AppData\Local\Programs\R\R-4.6.0\bin\Rscript.exe"
BATCH = "2026-09-21_post_feasibility_fix"
SEED = "20260921"
DROP_FEAS = {"Not closeable - no remaining pool left"}
SHARED_OWNER = "IRC"

PARTNERS = ["Save the Children", "Street Child of Nigeria", "IMC", "MDM", "PLAN", "IRC", "LHI",
            "ACF", "ZOA", "DRC", "Malteser", "FHI 360", "Solidarités"]


def slug(p):
    return re.sub(r"[^A-Za-z0-9]", "_", p).lower()


def run(cmd, log_path=None, cwd=PROJECT_DIR):
    with open(log_path, "w", encoding="utf-8") if log_path else open(os.devnull, "w") as out:
        rc = subprocess.run(cmd, cwd=cwd, stdout=out if log_path else subprocess.PIPE, stderr=subprocess.STDOUT).returncode
    return rc


def read_rows(p):
    if not os.path.exists(p):
        return []
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def write_rows(p, rows, fields):
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Strata Level"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {h: i for i, h in enumerate(hdr) if h}
    feas = {r[ix["Strata ID"]]: r[ix["Feasibility"]] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[ix["Strata ID"]]}
    wb.close()

    # 1. extract + scope-filter every partner
    csvs = {}
    for p in PARTNERS:
        d = os.path.join(RUNS, p, BATCH)
        os.makedirs(d, exist_ok=True)
        rc = run([RSCRIPT, os.path.join(SCRIPTS, "extract_partner_shortfalls.R"), p, d], os.path.join(d, "extract_log.txt"))
        print(f"[{p}] extract rc={rc}")
        for kind in ("", "_idp"):
            fn = os.path.join(d, f"{slug(p)}_shortfalls{kind}.csv")
            rows = read_rows(fn)
            if not rows:
                csvs[(p, kind)] = (fn, [], None)
                continue
            fields = list(rows[0].keys())
            write_rows(fn.replace(".csv", "_ALL_unfiltered.csv"), rows, fields)
            kept = [r for r in rows if feas.get(r["strata_id"]) not in DROP_FEAS]
            csvs[(p, kind)] = (fn, kept, fields)

    # 2. dedupe dual-coverage strata: keep under SHARED_OWNER, drop elsewhere
    owners = defaultdict(list)
    for (p, kind), (fn, rows, fields) in csvs.items():
        for r in rows:
            owners[r["strata_id"]].append(p)
    shared = {sid: ps for sid, ps in owners.items() if len(set(ps)) > 1}
    for sid, ps in sorted(shared.items()):
        keep = SHARED_OWNER if SHARED_OWNER in ps else sorted(set(ps))[0]
        print(f"  shared stratum {sid} covered by {sorted(set(ps))} -> drawn once under {keep}")
        for (p, kind), (fn, rows, fields) in csvs.items():
            if p != keep:
                csvs[(p, kind)] = (fn, [r for r in rows if r["strata_id"] != sid], fields)
    for (p, kind), (fn, rows, fields) in csvs.items():
        if fields:
            write_rows(fn, rows, fields)

    # 3. draws + stamp + projection, one partner at a time
    for p in PARTNERS:
        d = os.path.join(RUNS, p, BATCH)
        n_non = len(csvs[(p, "")][1]); n_idp = len(csvs[(p, "_idp")][1])
        need = sum(int(float(r["additional_clusters_needed"])) for k in ("", "_idp") for r in csvs[(p, k)][1])
        print(f"\n===== {p}: {n_non} Non-IDP + {n_idp} IDP strata in scope, {need} clusters =====")
        if n_non == 0 and n_idp == 0:
            print(f"[{p}] nothing in scope after filtering/dedupe - skipped")
            continue
        if n_non:
            rc = run([RSCRIPT, os.path.join(SCRIPTS, "draw_supplementary_clusters_batch.R"), csvs[(p, "")][0], d, SEED], os.path.join(d, "draw_non_idp_run_log.txt"))
            print(f"[{p}] Non-IDP draw rc={rc}")
            hh = os.path.join(d, "new_households.csv")
            if rc == 0 and os.path.exists(hh) and os.path.getsize(hh) > 0:
                print(f"[{p}] stamp Non-IDP rc={run([sys.executable, os.path.join(SCRIPTS, 'stamp_ward_accessible_status.py'), hh], os.path.join(d, 'stamp_non_idp_log.txt'), cwd=SCRIPTS)}")
        if n_idp:
            rc = run([RSCRIPT, os.path.join(SCRIPTS, "draw_supplementary_idp_sites_batch.R"), csvs[(p, "_idp")][0], d, SEED], os.path.join(d, "draw_idp_run_log.txt"))
            print(f"[{p}] IDP draw rc={rc}")
            hh = os.path.join(d, "new_households_idp_sitelevel.csv")
            if rc == 0 and os.path.exists(hh) and os.path.getsize(hh) > 0:
                print(f"[{p}] stamp IDP rc={run([sys.executable, os.path.join(SCRIPTS, 'stamp_ward_accessible_status.py'), hh], os.path.join(d, 'stamp_idp_log.txt'), cwd=SCRIPTS)}")
        proj = subprocess.run([sys.executable, os.path.join(SCRIPTS, "one_off_analyses", "project_post_draw_moe_2026-09-21.py"), p, d],
                              cwd=PROJECT_DIR, capture_output=True, text=True)
        print(proj.stdout.strip() or proj.stderr.strip())

    print("\nALL SMALL PARTNERS STAGED - nothing merged.")


if __name__ == "__main__":
    main()
