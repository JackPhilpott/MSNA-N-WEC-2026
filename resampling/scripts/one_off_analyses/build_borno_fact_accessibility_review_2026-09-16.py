# ==============================================================================
# Borno State accessibility review workbook, for FACT (ACTED's main Borno
# partner) - built 2026-09-16, relayed via the coordinating session on Jack's
# instinct ("makes more sense on the side that owns the actual accessibility
# pipeline/data directly"). Purpose: let FACT decide, LGA by LGA and ward by
# ward, where they can support/supplement/take over coverage from other
# partners across the whole state. Dikwa's known pending handover (currently
# Street Child of Nigeria, will move to FACT once Jack confirms) is NOT
# actioned here - shows as normal current-state context like every other row.
#
# Hard requirement (Jack, explicit, non-negotiable): EVERY ward in EVERY
# Borno LGA, not just wards currently in the sampling frame/WORKING universe.
# Checked directly before building anything (see export_borno_ward_backbone_
# 2026-09-16.R's own header): the existing master_accessibility_status_ward_
# level.csv is scoped to LGAs with coverage_status != "not_covered" in the
# current frame - Marte LGA has ZERO assigned partner and is therefore
# completely absent from it and every other current accessibility output.
# This script's ward backbone instead comes from GRID3's national admin3
# boundary directly (every real administrative ward, 310 across Borno's 27
# LGAs), independent of whether a partner/cluster has ever touched it.
#
# A ward with zero accessibility reports shows "Not yet assessed" -
# deliberately NOT "Accessible" (the existing default-unreported convention
# used elsewhere in this project, e.g. 04_build_master_accessibility_status.
# py) - Jack was explicit this must be a genuine third state here, since the
# whole point of this workbook is for FACT to see what's actually unknown,
# not have it silently look fine.
#
# Structure (Jack's own words: "we need to show the current decisions, and
# then a space for FACT to make another decision as well, and then we'll
# compare and see what gets overwritten") - two column blocks, never sharing
# a cell, each with its own group header row:
#   - CURRENT STATUS (read-only): State/LGA/Ward(GRID3+OCHA-COD)/Current
#     Accessibility Status/Current Reason category+notes/Current Reporting
#     partner(s)/Last reported date/Current covering partner(s)/existing
#     cluster+target+achieved context.
#   - FACT REVIEW (blank, dropdown-validated, for FACT to fill independently):
#     FACT's Assessed Accessibility Status/FACT's Recommended Decision/FACT
#     Reasons & Notes.
# ==============================================================================
import csv
from collections import defaultdict

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"

WARD_BACKBONE_CSV = PROJECT_DIR + r"\resampling\output\national_ward_backbone_2026-09-16.csv"
LOG_CSV = PROJECT_DIR + r"\resampling\output\resampling_requests_log.csv"
COVERAGE_XLSX = PROJECT_DIR + r"\input_data\boundaries\partner_coverage\Partnerscoverage.xlsx"
STAGE2_FULL_CSV = PROJECT_DIR + r"\output\data\data_collection\NGA_MSNA_2026_stage2_sampling_frame_v10_FULL.csv"
REAL_SUBMISSIONS_CSV = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\2_monitoring\data\real_submissions.csv"
CONFIRMED_DELETIONS_OVERLAY_CSV = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\2_monitoring\data\CONFIRMED_DELETIONS_OVERLAY.csv"
OUT_XLSX = PROJECT_DIR + r"\resampling\output\resample_runs\FACT\2026-09-17_borno_accessibility_review\FACT_Borno_Accessibility_Review_2026-09-17.xlsx"

# ---------------------------------------------------------------------------
# 1. Ward backbone - EVERY real Borno ward, independent of sampling frame.
# Sourced from the corrected NATIONAL backbone (export_national_ward_
# backbone_2026-09-16.R - real GRID3 national product, LGA-reconciled
# against the canonical sampling frame), filtered to Borno here. Earlier
# version of this script read the wrong shapefile (an OCHA/COD NE-only
# product mistaken for GRID3's national one) - see that R script's own
# header for the full correction story. Not sent to FACT yet, so rebuilding
# on the corrected source before anything went out was still possible.
# ---------------------------------------------------------------------------
with open(WARD_BACKBONE_CSV, encoding="utf-8") as f:
    backbone = [r for r in csv.DictReader(f) if r["adm1_name"] == "Borno"]
print(f"Loaded {len(backbone)} real Borno wards across {len({r['adm2_name'] for r in backbone})} LGAs (GRID3, full administrative list).")

# ---------------------------------------------------------------------------
# 2. Latest ward-level accessibility report per (partner, state, lga, ward) -
# duplicated from 04_build_master_accessibility_status.py's own
# latest_ward_reports(), per this project's standalone-script convention -
# this is the RAW report log, not scoped to any sampling universe, which is
# exactly why it's safe to join against the full backbone above.
# ---------------------------------------------------------------------------
with open(LOG_CSV, encoding="utf-8", newline="") as f:
    log_rows = list(csv.DictReader(f))

latest = {}
for r in log_rows:
    if r["report_level"] != "ward":
        continue
    key = (r["partner"], r["state"], r["lga"], r["ward_name"])
    if key not in latest or int(r["request_id"]) > int(latest[key]["request_id"]):
        latest[key] = r

reports_by_ward = defaultdict(list)
for (partner, state, lga, ward), rep in latest.items():
    reports_by_ward[(state, lga, ward)].append(rep)
print(f"Loaded {len(log_rows)} accessibility log rows -> {len(reports_by_ward)} distinct (state, lga, ward) combinations with at least one ward-level report.")

# ---------------------------------------------------------------------------
# 3. Current partner coverage, per Borno LGA (Partnerscoverage.xlsx, NE sheet
# - LGA-grain, same source/format as every other partner-coverage lookup in
# this project). A ward with no LGA-level partner at all (Marte) legitimately
# shows blank here - not an error, exactly the case this workbook exists to
# surface.
# ---------------------------------------------------------------------------
wb_cov = openpyxl.load_workbook(COVERAGE_XLSX, data_only=True)
ws_ne = wb_cov["NE"]
cov_rows = list(ws_ne.iter_rows(values_only=True))
cov_header = cov_rows[0]
partner_col_idx = list(range(3, cov_header.index("COUNT")))
covering_partner_by_lga = {}
for row in cov_rows[1:]:
    if row[1] != "Borno" or not row[2]:
        continue
    partners_here = [str(cov_header[i]).strip() for i in partner_col_idx if row[i]]
    covering_partner_by_lga[row[2]] = "; ".join(partners_here) if partners_here else ""
print(f"Loaded partner coverage for {len(covering_partner_by_lga)} Borno LGAs.")

# ---------------------------------------------------------------------------
# 4. Cluster / target / achieved context, per (state, lga, ward) - from the
# CURRENT sampling frame (FULL), for wards that happen to be in it. Wards
# outside the frame entirely (no cluster ever drawn there) correctly show
# 0 clusters / 0 target / 0 achieved - real, not a data gap.
# ---------------------------------------------------------------------------
with open(STAGE2_FULL_CSV, encoding="utf-8") as f:
    full_rows = list(csv.DictReader(f))
borno_full = [r for r in full_rows if r["adm1_name"] == "Borno"]

with open(CONFIRMED_DELETIONS_OVERLAY_CSV, encoding="utf-8") as f:
    confirmed_deleted = {r["uuid"] for r in csv.DictReader(f) if r["status"] in ("confirmed", "contested")}
with open(REAL_SUBMISSIONS_CSV, encoding="utf-8") as f:
    real_subs = list(csv.DictReader(f))


def is_achieved(r):
    return (
        r.get("interview_outcome") == "completed"
        and r.get("matched_survey_id") not in (None, "", "NA")
        and r.get("submission_uuid") not in confirmed_deleted
    )


achieved_survey_ids = {r["matched_survey_id"] for r in real_subs if r.get("pop_type") == "non_idp" and is_achieved(r)}
cluster_achieved_n = defaultdict(int)
for r in real_subs:
    cid = r.get("matched_cluster_id")
    if cid and cid != "NA" and is_achieved(r):
        cluster_achieved_n[cid] += 1

# Group Borno FULL rows by (state, lga, ward) - a cluster can straddle >1
# ward (real, ~35% of Non-IDP clusters nationally - see 01_generate_
# accessibility_reports.py), so this groups by WARD-PORTION, same as every
# other ward-grain script in this project, not by whole cluster.
#
# IDP GOTCHA (caught before this shipped, not after): the household-level
# FULL frame has ONE ROW PER PLANNED INTERVIEW SLOT for an IDP cluster (not
# one row per cluster) - e.g. idp_NG008008_1 (target_households=30) appears
# as 30 near-identical primary rows. target_households/achieved are CLUSTER-
# level facts repeated on every one of those rows - summing them per row (as
# a first draft of this script did) inflates target/achieved by roughly the
# cluster's own size (Dikwa ward's Total target HHs first came out as 3612,
# not a plausible number for one ward). Fixed by only adding an IDP
# cluster's target/achieved the FIRST time that (ward, cluster_id) pair is
# seen - mirrors non_idp_cluster_summary_rows()/idp_cluster_summary_row()'s
# own one-row-per-cluster grouping in build_partner_dc_packages.py, just
# done inline here since this script doesn't need that function's full
# per-cluster row-builder.
ward_context = defaultdict(lambda: {
    "non_idp_clusters": set(), "idp_clusters": set(), "target_hh": 0, "achieved": 0,
    "ward_cod": "", "ward_cod_pcode": "",
})
seen_idp_cluster_in_ward = set()
for r in borno_full:
    key = (r["adm1_name"], r["adm2_name"], r["adm3_name"])
    ctx = ward_context[key]
    ward_cod = r.get("admin3_cod_name")
    if ward_cod and ward_cod != "NA":
        ctx["ward_cod"] = ward_cod
        ctx["ward_cod_pcode"] = r.get("admin3_cod_pcode") or ""
    if r["status"] != "primary":
        continue
    if r["pop_type"] == "non_idp":
        ctx["non_idp_clusters"].add(r["cluster_id"])
        ctx["target_hh"] += 1
        if r["survey_id"] in achieved_survey_ids:
            ctx["achieved"] += 1
    else:
        ctx["idp_clusters"].add(r["cluster_id"])
        cluster_ward_key = (key, r["cluster_id"])
        if cluster_ward_key in seen_idp_cluster_in_ward:
            continue
        seen_idp_cluster_in_ward.add(cluster_ward_key)
        nominal_target = int(r["target_households"]) if r.get("target_households") not in (None, "", "NA") else 0
        ctx["target_hh"] += nominal_target
        n_ach = cluster_achieved_n.get(r["cluster_id"], 0)
        ctx["achieved"] += min(n_ach, nominal_target) if nominal_target else n_ach

print(f"Built cluster/target/achieved context for {len(ward_context)} (state, lga, ward) combinations actually present in the current sampling frame.")

# ---------------------------------------------------------------------------
# 5. Assemble one row per real Borno ward.
# ---------------------------------------------------------------------------
rows_out = []
for w in backbone:
    state, lga, ward = "Borno", w["adm2_name"], w["adm3_name"]
    key = (state, lga, ward)
    reps = reports_by_ward.get(key, [])
    any_no = any((r["accessible"] or "").strip().lower() == "no" for r in reps)
    any_yes = any((r["accessible"] or "").strip().lower() == "yes" for r in reps)
    if any_no:
        status = "Inaccessible"
    elif any_yes:
        status = "Accessible"
    else:
        # HARD REQUIREMENT (Jack): a ward with zero reports is a genuine
        # third state, never silently defaulted to Accessible.
        status = "Not yet assessed"

    reason_categories = sorted({(r.get("reason_category") or "").strip() for r in reps if (r.get("reason_category") or "").strip()})
    reason_notes = sorted({(r.get("reason_notes") or "").strip() for r in reps if (r.get("reason_notes") or "").strip()})
    reporting_partners = sorted({r["partner"] for r in reps})
    dates = sorted({(r.get("date_reported_by_partner") or "").strip() for r in reps if (r.get("date_reported_by_partner") or "").strip()})

    ctx = ward_context.get(key)
    ctx = ctx or {"non_idp_clusters": set(), "idp_clusters": set(), "target_hh": 0, "achieved": 0, "ward_cod": "", "ward_cod_pcode": ""}

    rows_out.append({
        "State": state, "LGA": lga,
        "Ward (GRID3)": ward, "Ward (OCHA/COD)": ctx["ward_cod"],
        "Current Accessibility Status": status,
        "Current Reason category": "; ".join(reason_categories),
        "Current Reason notes": "; ".join(reason_notes),
        "Current Reporting partner(s)": "; ".join(reporting_partners),
        "Last reported date": "; ".join(dates),
        "Current covering/assigned partner(s)": covering_partner_by_lga.get(lga, ""),
        "Non-IDP clusters": len(ctx["non_idp_clusters"]),
        "IDP clusters": len(ctx["idp_clusters"]),
        "Total target HHs (primary)": ctx["target_hh"],
        "Achieved so far (real interviews)": ctx["achieved"],
    })

n_not_assessed = sum(1 for r in rows_out if r["Current Accessibility Status"] == "Not yet assessed")
n_no_partner = sum(1 for r in rows_out if not r["Current covering/assigned partner(s)"])
print(f"Assembled {len(rows_out)} ward rows. {n_not_assessed} 'Not yet assessed' (zero reports). {n_no_partner} with no LGA-level partner assigned at all.")

# ---------------------------------------------------------------------------
# 6. Write the workbook - two structurally distinct column blocks.
# ---------------------------------------------------------------------------
CURRENT_COLS = [
    "State", "LGA", "Ward (GRID3)", "Ward (OCHA/COD)",
    "Current Accessibility Status", "Current Reason category", "Current Reason notes",
    "Current Reporting partner(s)", "Last reported date", "Current covering/assigned partner(s)",
    "Non-IDP clusters", "IDP clusters", "Total target HHs (primary)", "Achieved so far (real interviews)",
]
FACT_COLS = ["FACT's Assessed Accessibility Status", "FACT's Recommended Decision", "FACT Reasons / Notes"]
ALL_COLS = CURRENT_COLS + FACT_COLS

wb_out = openpyxl.Workbook()
ws_r = wb_out.active
ws_r.title = "README"
ws_r.column_dimensions["A"].width = 100
readme_lines = [
    ("FACT - Borno State Accessibility Review", True, 14),
    ("", False, 11),
    ("Purpose: a state-wide, ward-by-ward view of current accessibility and partner coverage across all of "
     "Borno, for FACT to review and propose where FACT could support, supplement, or take over coverage from "
     "another partner. Nothing here is a decision yet - the CURRENT STATUS columns are read-only, reflecting "
     "what is known today; the FACT REVIEW columns are blank, for FACT's own team to complete independently. "
     "Once returned, Current vs FACT's columns will be compared directly to see exactly what is being proposed "
     "to change.", False, 11),
    ("", False, 11),
    ("EVERY real ward in EVERY Borno LGA is listed - 310 wards across all 27 LGAs - not just wards currently "
     "assigned to a partner or already part of an active survey cluster. A ward with 'Not yet assessed' under "
     "Current Accessibility Status has ZERO accessibility reports on file from any partner - this is a genuine "
     "'we don't know yet' state, not a finding, and it is never the same as 'Accessible'. Some LGAs (e.g. Marte) "
     "currently have no partner assigned at all - 'Current covering/assigned partner(s)' is correctly blank for "
     "those, not an omission.", False, 11),
    ("", False, 11),
    ("Ward (OCHA/COD) is only populated for wards that already have at least one survey cluster in the current "
     "sampling frame - it comes from a cross-reference built at draw time, which a never-yet-drawn ward has no "
     "occasion to have computed. Ward (GRID3) is populated for every row and is the reliable match key throughout.", False, 11),
    ("", False, 11),
    ("Dikwa is shown here as normal current-state context (currently Street Child of Nigeria, not started) - its "
     "planned handover to FACT is a separate, already-agreed decision pending Jack's final confirmation, not "
     "something this workbook is asking FACT to weigh in on.", False, 11),
    ("", False, 11),
    ("Column definitions - CURRENT STATUS (read-only)", True, 12),
    ("Current Accessibility Status", False, 11, "Accessible / Inaccessible / Not yet assessed - 'any partner report says Inaccessible' wins over any 'Accessible' report for the same ward, same conservative rule used elsewhere in this project."),
    ("Current Reason category / notes", False, 11, "From the most recent partner report(s) on file for this ward, if any."),
    ("Current Reporting partner(s)", False, 11, "Every partner who has filed an accessibility report specifically for this ward, if any."),
    ("Current covering/assigned partner(s)", False, 11, "The LGA-level partner assignment from the master partner-coverage list - independent of whether that partner has reported on this specific ward yet."),
    ("Non-IDP / IDP clusters, Total target HHs, Achieved so far", False, 11, "Only populated where this ward already has at least one survey cluster in the current sampling frame - 0/0 is real and expected for a ward never yet drawn into the frame."),
    ("", False, 11),
    ("Column definitions - FACT REVIEW (please complete)", True, 12),
    ("FACT's Assessed Accessibility Status", False, 11, "FACT's own current read on this ward - use the dropdown."),
    ("FACT's Recommended Decision", False, 11, "Retain with current partner / FACT supplements / FACT takes over - use the dropdown."),
    ("FACT Reasons / Notes", False, 11, "Free text - why this recommendation."),
]
r = 1
for line in readme_lines:
    text = line[0]
    bold = line[1]
    size = line[2] if len(line) > 2 else 11
    cell = ws_r.cell(row=r, column=1, value=text)
    cell.font = openpyxl.styles.Font(bold=bold, size=size)
    cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    if len(line) > 3:
        ws_r.cell(row=r, column=2, value=line[3]).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        ws_r.column_dimensions["B"].width = 90
    r += 1

ws = wb_out.create_sheet("Borno Ward Review")
# Group header row (row 1): one merged cell per block.
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CURRENT_COLS))
cell = ws.cell(row=1, column=1, value="CURRENT STATUS (read-only - do not edit)")
cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF", size=12)
cell.fill = openpyxl.styles.PatternFill("solid", fgColor="1B2A4A")
cell.alignment = openpyxl.styles.Alignment(horizontal="center")

ws.merge_cells(start_row=1, start_column=len(CURRENT_COLS) + 1, end_row=1, end_column=len(ALL_COLS))
cell = ws.cell(row=1, column=len(CURRENT_COLS) + 1, value="FACT REVIEW (please complete)")
cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF", size=12)
cell.fill = openpyxl.styles.PatternFill("solid", fgColor="A5281B")
cell.alignment = openpyxl.styles.Alignment(horizontal="center")

# Column header row (row 2).
for c, col in enumerate(ALL_COLS, start=1):
    cell = ws.cell(row=2, column=c, value=col)
    cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    is_fact_col = c > len(CURRENT_COLS)
    cell.fill = openpyxl.styles.PatternFill("solid", fgColor="C0392B" if is_fact_col else "2C5F8A")

for i, row in enumerate(rows_out, start=3):
    for c, col in enumerate(CURRENT_COLS, start=1):
        ws.cell(row=i, column=c, value=row[col])
    for c in range(len(CURRENT_COLS) + 1, len(ALL_COLS) + 1):
        ws.cell(row=i, column=c, value=None).fill = openpyxl.styles.PatternFill("solid", fgColor="FFF2CC")

# Conditional-format Current Accessibility Status for a quick visual scan.
status_col_letter = openpyxl.utils.get_column_letter(CURRENT_COLS.index("Current Accessibility Status") + 1)
status_range = f"{status_col_letter}3:{status_col_letter}{len(rows_out) + 2}"
from openpyxl.formatting.rule import CellIsRule
ws.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Accessible"'], fill=openpyxl.styles.PatternFill("solid", fgColor="C6E0B4")))
ws.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Inaccessible"'], fill=openpyxl.styles.PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Not yet assessed"'], fill=openpyxl.styles.PatternFill("solid", fgColor="D9D9D9")))

# Dropdown validation on the two FACT status/decision columns.
dv_status = DataValidation(type="list", formula1='"Accessible,Inaccessible,Not yet assessed,Unable to determine"', allow_blank=True)
dv_status.error = "Choose from the list."
ws.add_data_validation(dv_status)
fact_status_col = len(CURRENT_COLS) + 1
dv_status.add(f"{openpyxl.utils.get_column_letter(fact_status_col)}3:{openpyxl.utils.get_column_letter(fact_status_col)}{len(rows_out) + 2}")

dv_decision = DataValidation(type="list", formula1='"Retain with current partner,FACT supplements,FACT takes over"', allow_blank=True)
dv_decision.error = "Choose from the list."
ws.add_data_validation(dv_decision)
fact_decision_col = len(CURRENT_COLS) + 2
dv_decision.add(f"{openpyxl.utils.get_column_letter(fact_decision_col)}3:{openpyxl.utils.get_column_letter(fact_decision_col)}{len(rows_out) + 2}")

for c, col in enumerate(ALL_COLS, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(14, min(32, len(col) + 4))
ws.freeze_panes = "E3"

import os
os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
wb_out.save(OUT_XLSX)
print(f"\nWrote {OUT_XLSX} - {len(rows_out)} ward rows.")
