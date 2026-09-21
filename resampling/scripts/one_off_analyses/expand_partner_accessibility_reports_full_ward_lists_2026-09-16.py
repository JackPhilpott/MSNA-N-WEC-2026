# ==============================================================================
# Expands all 19 partners' accessibility reports (Ward Accessibility sheet)
# to a COMPLETE ward list per covered LGA - built 2026-09-16, Jack's direct,
# explicit go-ahead (see 1_sampling/CLAUDE.md). Root problem this fixes: a
# ward with zero accessibility reports on file silently defaults to
# "Accessible" everywhere downstream (04_build_master_accessibility_status.py's
# own documented rule) - real, quantified nationally at 370 ward-portions /
# 4,518 target households currently in that exact state (verified against
# master_accessibility_status_ward_level.csv before this was built). Jack's
# own diagnosis, from following up with INTERSOS on Maru: their report only
# ever covered wards already assigned/in-frame - 6 more wards in that LGA had
# genuinely never been asked about at all, not confirmed accessible by
# anyone. The fix: give every partner the chance to see and answer for every
# real ward, not just ones already in the frame.
#
# HARD REQUIREMENTS (Jack, explicit, confirmed via AskUserQuestion before
# this was built):
#   1. LGA scope per partner comes from the CURRENT SAMPLING FRAME's own
#      partners_covering column (coverage_status=="covered" &
#      exclusion_reason=="none"), NOT Partnerscoverage.xlsx - confirmed with
#      the coordinating session before proceeding: coverage reinstatements
#      (e.g. Dandume/Faskari) get patched directly into the frame's own
#      coverage_status/exclusion_reason columns, never back into the source
#      spreadsheet, so the frame is the only source that reflects every
#      coverage decision made throughout this project's history.
#   2. ALL existing information/answers already in a partner's file are kept
#      completely untouched - this script only ever APPENDS new blank rows
#      for wards that weren't previously listed. Never rewrites, reorders, or
#      clears an existing row.
#   3. A partner with no returned file yet builds from their generated
#      (outgoing) file instead (accessibility_reports_generated/).
#   4. Output overwrites accessibility_reports_returned/<Partner>_
#      accessibility_report.xlsx in place (backed up first, to _archive/,
#      matching this folder's own existing backup-before-fix convention) -
#      so that folder is unambiguously "where all the main partner reports
#      sit" (Jack's own words), for every one of the 19 partners, regardless
#      of whether they've actually returned anything yet.
#
# Ward backbone: resampling/output/national_ward_backbone_2026-09-16.csv
# (export_national_ward_backbone_2026-09-16.R) - the REAL, GRID3-sourced,
# LGA-reconciled national ward list. Read that script's own header before
# touching this one again - the first version of this whole effort (including
# the already-built Borno/FACT workbook) used the wrong shapefile (an
# OCHA/COD NE-only product mistaken for GRID3's national one), caught and
# fixed before anything was sent.
#
# SCHEMA-ADAPTIVE by necessity: checked every partner's actual current Ward
# Accessibility sheet header before writing this - 18 of 19 match the
# standard WARD_COLUMNS template (some missing the two provenance columns,
# expected/known variance), but FACT's has evolved through a followup-
# workbook process into a genuinely different column set (extra "Needs your
# input?"/"Why flagged"/"Column 1" columns, missing several standard ones).
# This script reads each file's OWN header row and matches by COLUMN NAME,
# never a fixed position - a new row fills whatever columns it recognises
# (State/LGA/Ward names, cluster/target context) and leaves anything it
# doesn't recognise genuinely blank, rather than forcing every file into one
# canonical schema (which risked silently breaking a file's own established
# structure/formatting for a partner-specific reason not visible from here).
# ==============================================================================
import csv
import os
import shutil
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
WARD_BACKBONE_CSV = PROJECT_DIR + r"\resampling\output\national_ward_backbone_2026-09-16.csv"
STAGE2_FULL_CSV = PROJECT_DIR + r"\output\data\data_collection\NGA_MSNA_2026_stage2_sampling_frame_v9_FULL.csv"
STRATA_FULL_CSV = PROJECT_DIR + r"\output\data\data_collection\NGA_MSNA_2026_strata_level_sampling_frame_v9_FULL.csv"
RETURNED_DIR = PROJECT_DIR + r"\resampling\input\accessibility_reports_returned"
GENERATED_DIR = PROJECT_DIR + r"\resampling\input\accessibility_reports_generated"
ARCHIVE_DIR = os.path.join(RETURNED_DIR, "_archive")

MULTI_LGA_WARD_NOTE = (
    "This ward's GRID3 polygon spans more than one LGA - you only need to report on the part of the ward that "
    "falls within your own LGA coverage. See the 'Other LGA(s) sharing this ward' column for where the rest "
    "of it sits."
)
NEW_ROW_NOTE_PREFIX = "Newly added 2026-09-16 - not previously listed. "

ACCESSIBLE_OPTIONS = ["Yes", "No"]
REASON_OPTIONS = [
    "N/A - fully accessible", "Insecurity / conflict", "Physical access (terrain, flooding, roads)",
    "Population absent / relocated", "Building-footprint issue (no eligible structures)",
    "Access denied by authorities / community", "Other",
]
REPORTED_BY_OPTIONS = ["Partner", "IMPACT (default - accessible until reported otherwise)"]
SOURCE_CHANNEL_OPTIONS = [
    "Partner's own template", "Email", "WhatsApp/verbal (coordinator-transcribed)",
    "Point-level file annotation", "N/A",
]
DATE_REPORTED_MIN = date(2026, 7, 1)
DATE_REPORTED_FORMAT = "dd-mmm-yyyy"

# ---------------------------------------------------------------------------
# 1. Ward backbone + ward-to-LGA-list (spans-multiple-LGAs detection),
# national, real GRID3 source.
# ---------------------------------------------------------------------------
with open(WARD_BACKBONE_CSV, encoding="utf-8") as f:
    backbone = list(csv.DictReader(f))
print(f"Loaded {len(backbone)} real national ward rows.")

ward_to_lgas = defaultdict(set)
for r in backbone:
    ward_to_lgas[(r["adm1_name"], r["adm3_name"])].add(r["adm2_name"])
ward_to_lgas = {k: sorted(v) for k, v in ward_to_lgas.items()}

wards_by_lga = defaultdict(list)
for r in backbone:
    wards_by_lga[(r["adm1_name"], r["adm2_name"])].append(r["adm3_name"])

# ---------------------------------------------------------------------------
# 2. Partner LGA coverage - from the CURRENT sampling frame (Jack's confirmed
# call, not Partnerscoverage.xlsx - see header).
# ---------------------------------------------------------------------------
with open(STRATA_FULL_CSV, encoding="utf-8") as f:
    strata_rows = list(csv.DictReader(f))

partner_lgas = defaultdict(set)
for r in strata_rows:
    if r.get("coverage_status") == "covered" and r.get("exclusion_reason") == "none":
        for p in (r.get("partners_covering") or "").split(","):
            p = p.strip()
            if p:
                partner_lgas[p].add((r["adm1_name"], r["adm2_name"]))

PARTNERS = sorted(partner_lgas.keys())
print(f"{len(PARTNERS)} partners with real current coverage: {PARTNERS}")

# ---------------------------------------------------------------------------
# 3. Cluster/target context per (state, lga, ward) - national, from FULL.
# Same IDP one-row-per-interview-slot gotcha as the Borno workbook (see that
# script's own comment) - deduped per (ward, cluster_id).
# ---------------------------------------------------------------------------
with open(STAGE2_FULL_CSV, encoding="utf-8") as f:
    full_rows = list(csv.DictReader(f))

ward_context = defaultdict(lambda: {"non_idp_clusters": set(), "idp_clusters": set(), "target_hh": 0, "ward_cod": ""})
seen_idp_cluster_in_ward = set()
for r in full_rows:
    key = (r["adm1_name"], r["adm2_name"], r["adm3_name"])
    ctx = ward_context[key]
    ward_cod = r.get("admin3_cod_name")
    if ward_cod and ward_cod != "NA":
        ctx["ward_cod"] = ward_cod
    if r["status"] != "primary":
        continue
    if r["pop_type"] == "non_idp":
        ctx["non_idp_clusters"].add(r["cluster_id"])
        ctx["target_hh"] += 1
    else:
        ctx["idp_clusters"].add(r["cluster_id"])
        cluster_ward_key = (key, r["cluster_id"])
        if cluster_ward_key in seen_idp_cluster_in_ward:
            continue
        seen_idp_cluster_in_ward.add(cluster_ward_key)
        nominal_target = int(r["target_households"]) if r.get("target_households") not in (None, "", "NA") else 0
        ctx["target_hh"] += nominal_target
print(f"Built cluster/target context for {len(ward_context)} (state, lga, ward) combinations.")


def safe_folder_name(s):
    import re
    return re.sub(r'[<>:"/\\|?*]', "-", s).strip()


def find_col(header, name):
    try:
        return header.index(name) + 1
    except ValueError:
        return None


def extend_or_add_validation(ws, col_idx, kind, existing_dvs, n_rows_before, new_row_range):
    """Reuses an existing DataValidation covering this column if one exists
    on the sheet (extends its range to include the new rows); otherwise adds
    a fresh one with this project's standard option list for that column
    kind. Keeps whatever validation a partner's own file already had (e.g. a
    file regenerated at an earlier point in this project's history might
    have a slightly different option list already in place) rather than
    silently replacing it."""
    from openpyxl.utils import range_boundaries
    for dv in existing_dvs:
        for cell_range in dv.sqref.ranges:
            min_col, _, max_col, _ = range_boundaries(str(cell_range))
            if min_col <= col_idx <= max_col:
                dv.add(new_row_range)
                return
    formula_map = {
        "accessible": f'"{",".join(ACCESSIBLE_OPTIONS)}"',
        "reason": f'"{",".join(REASON_OPTIONS)}"',
        "reported_by": f'"{",".join(REPORTED_BY_OPTIONS)}"',
        "source_channel": f'"{",".join(SOURCE_CHANNEL_OPTIONS)}"',
    }
    if kind == "date":
        dv = DataValidation(
            type="date", operator="between", formula1=DATE_REPORTED_MIN, formula2=date.today(),
            allow_blank=True, showErrorMessage=True, errorTitle="Invalid date",
            error=f"Enter an actual date between {DATE_REPORTED_MIN:%d %b %Y} and today.",
        )
    else:
        dv = DataValidation(type="list", formula1=formula_map[kind], allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(new_row_range)


def process_partner(partner):
    fname = f"{safe_folder_name(partner)}_accessibility_report.xlsx"
    returned_path = os.path.join(RETURNED_DIR, fname)
    generated_path = os.path.join(GENERATED_DIR, fname)

    if os.path.exists(returned_path):
        src_path = returned_path
        src_kind = "returned"
    elif os.path.exists(generated_path):
        src_path = generated_path
        src_kind = "generated (no return on file yet)"
    else:
        print(f"  {partner}: SKIPPED - no file found in either returned/ or generated/.")
        return

    wb = openpyxl.load_workbook(src_path)
    if "Ward Accessibility" not in wb.sheetnames:
        print(f"  {partner}: SKIPPED - no 'Ward Accessibility' sheet in {src_path}.")
        return
    ws = wb["Ward Accessibility"]
    header = [c.value for c in ws[1]]

    col_state = find_col(header, "State")
    col_lga = find_col(header, "LGA")
    col_ward = find_col(header, "Ward (GRID3)")
    if not (col_state and col_lga and col_ward):
        print(f"  {partner}: SKIPPED - Ward Accessibility sheet missing State/LGA/Ward (GRID3) columns, can't safely process.")
        return

    n_rows_before = ws.max_row - 1  # excl header
    existing_wards = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        existing_wards.add((row[col_state - 1], row[col_lga - 1], row[col_ward - 1]))

    col_ocha = find_col(header, "Ward (OCHA/COD)")
    col_spans = find_col(header, "Ward spans multiple LGAs (Y/N)")
    col_other_lga = find_col(header, "Other LGA(s) sharing this ward")
    col_note = find_col(header, "Note")
    col_non_idp = find_col(header, "Non-IDP clusters")
    col_idp = find_col(header, "IDP clusters")
    col_target = find_col(header, "Total target HHs (primary)")

    lgas = sorted(partner_lgas.get(partner, set()))
    new_rows = []
    for (state, lga) in lgas:
        for ward in sorted(wards_by_lga.get((state, lga), [])):
            key = (state, lga, ward)
            if key in existing_wards:
                continue
            ctx = ward_context.get(key, {"non_idp_clusters": set(), "idp_clusters": set(), "target_hh": 0, "ward_cod": ""})
            other_lgas = [x for x in ward_to_lgas.get((state, ward), [lga]) if x != lga]
            new_rows.append({
                "state": state, "lga": lga, "ward": ward,
                "ocha": ctx["ward_cod"],
                "spans": "Yes" if other_lgas else "No",
                "other_lga": "; ".join(other_lgas),
                "note": (NEW_ROW_NOTE_PREFIX + MULTI_LGA_WARD_NOTE) if other_lgas else NEW_ROW_NOTE_PREFIX.strip(),
                "non_idp": len(ctx["non_idp_clusters"]), "idp": len(ctx["idp_clusters"]), "target_hh": ctx["target_hh"],
            })
            existing_wards.add(key)  # guard against a (state,lga,ward) dup within the backbone itself

    if not new_rows:
        print(f"  {partner}: 0 new ward rows needed - already complete against {len(lgas)} covered LGA(s). ({src_kind})")
        return

    new_rows.sort(key=lambda r: (r["state"], r["lga"], r["ward"]))
    start_row = ws.max_row + 1
    for i, nr in enumerate(new_rows):
        row_i = start_row + i
        ws.cell(row=row_i, column=col_state, value=nr["state"])
        ws.cell(row=row_i, column=col_lga, value=nr["lga"])
        ws.cell(row=row_i, column=col_ward, value=nr["ward"])
        if col_ocha:
            ws.cell(row=row_i, column=col_ocha, value=nr["ocha"])
        if col_spans:
            ws.cell(row=row_i, column=col_spans, value=nr["spans"])
        if col_other_lga:
            ws.cell(row=row_i, column=col_other_lga, value=nr["other_lga"])
        if col_note:
            ws.cell(row=row_i, column=col_note, value=nr["note"])
        if col_non_idp:
            ws.cell(row=row_i, column=col_non_idp, value=nr["non_idp"])
        if col_idp:
            ws.cell(row=row_i, column=col_idp, value=nr["idp"])
        if col_target:
            ws.cell(row=row_i, column=col_target, value=nr["target_hh"])
        # every other existing column on this row is left genuinely blank -
        # unreported, exactly like any other never-answered row.

    end_row = start_row + len(new_rows) - 1

    # Extend the Excel Table's own range, if one exists, so filters/styling
    # cover the new rows too.
    for tbl_name in list(ws.tables.keys()):
        tbl = ws.tables[tbl_name]
        from openpyxl.utils import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        tbl.ref = f"{openpyxl.utils.get_column_letter(min_col)}{min_row}:{openpyxl.utils.get_column_letter(max_col)}{end_row}"

    # Extend/add dropdown validation on the input columns for the new rows.
    existing_dvs = list(ws.data_validations.dataValidation)
    col_accessible = find_col(header, "Accessible (Y/N)")
    col_reason = find_col(header, "Reason category")
    col_reported_by = find_col(header, "Reported by (Partner / IMPACT-default)")
    col_source_channel = find_col(header, "Source channel")
    col_date = find_col(header, "Date reported")
    if col_accessible:
        extend_or_add_validation(ws, col_accessible, "accessible", existing_dvs, n_rows_before,
                                  f"{openpyxl.utils.get_column_letter(col_accessible)}{start_row}:{openpyxl.utils.get_column_letter(col_accessible)}{end_row}")
    if col_reason:
        extend_or_add_validation(ws, col_reason, "reason", existing_dvs, n_rows_before,
                                  f"{openpyxl.utils.get_column_letter(col_reason)}{start_row}:{openpyxl.utils.get_column_letter(col_reason)}{end_row}")
    if col_reported_by:
        extend_or_add_validation(ws, col_reported_by, "reported_by", existing_dvs, n_rows_before,
                                  f"{openpyxl.utils.get_column_letter(col_reported_by)}{start_row}:{openpyxl.utils.get_column_letter(col_reported_by)}{end_row}")
    if col_source_channel:
        extend_or_add_validation(ws, col_source_channel, "source_channel", existing_dvs, n_rows_before,
                                  f"{openpyxl.utils.get_column_letter(col_source_channel)}{start_row}:{openpyxl.utils.get_column_letter(col_source_channel)}{end_row}")
    if col_date:
        extend_or_add_validation(ws, col_date, "date", existing_dvs, n_rows_before,
                                  f"{openpyxl.utils.get_column_letter(col_date)}{start_row}:{openpyxl.utils.get_column_letter(col_date)}{end_row}")
        for row_i in range(start_row, end_row + 1):
            ws.cell(row=row_i, column=col_date).number_format = DATE_REPORTED_FORMAT

    # Backup-before-fix, matching this folder's own established convention.
    os.makedirs(ARCHIVE_DIR, exist_ok=True)
    if os.path.exists(returned_path):
        backup_name = f"{safe_folder_name(partner)}_accessibility_report_pre_ward_completeness_expansion_2026-09-16.xlsx"
        shutil.copy2(returned_path, os.path.join(ARCHIVE_DIR, backup_name))

    wb.save(returned_path)
    print(f"  {partner}: {n_rows_before} existing rows + {len(new_rows)} new ward row(s) added -> {n_rows_before + len(new_rows)} total. ({src_kind}) -> {returned_path}")


print()
for partner in PARTNERS:
    process_partner(partner)
print("\nDONE.")
