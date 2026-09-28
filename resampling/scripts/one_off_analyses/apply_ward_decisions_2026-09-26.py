# ==============================================================================
# 2026-09-26 (final accessibility changes): apply an explicit list of ward decisions to ONE partner's returned accessibility report,
# same discipline as apply_conpad_accessibility_decisions_2026-09-26.py but driven by a plan CSV, and able to APPEND a ward row when the
# partner's own file has none (Chibok and Damboa moved IMC -> FACT on 25 Sep, so FACT's file has no rows for their wards).
#
# plan CSV columns: action (update|append), state, lga, ward, accessible (Yes|No), reason_category, detail, date (YYYY-MM-DD)
#   update: the row must exist exactly once in 'Ward Accessibility'; only Accessible, Reason category, Reason notes, Date reported change
#   append: the row must NOT exist; a new row is added at the end (State/LGA/Ward, informational columns copied from the master, the four
#           answer columns, and a note in the informational column 'Why flagged' or 'Note')
# Reason notes = "<detail> | Source: <source>" (detail 'KEEP' keeps an existing real note; empty detail leaves only the source).
# Built from the pre-change backup; result written to a temp file, diffed cell by cell against the backup, moved into place only if the diff is
# EXACTLY the plan. Dry run by default; --execute needs the backup to exist (the script makes it if missing).
# ==============================================================================
import argparse
import copy
import csv
import os
import re
import shutil
import sys
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling\resampling"
RETURNED = PROJECT_DIR + r"\input\accessibility_reports_returned"
ARCHIVE = RETURNED + r"\_archive"
MASTER = PROJECT_DIR + r"\output\master_accessibility_status_ward_level.csv"
EDIT_COLS = ("Accessible (Y/N)", "Reason category", "Reason notes", "Date reported")


def is_real_note(v):
    return v is not None and str(v).strip() not in ("", "NA", "N/A", "nan", "None")


def note(existing, detail, source):
    parts = []
    if "{KEEP}" in detail:  # e.g. "Restored ...; previous note: {KEEP}" - keep the partner's earlier real note inside the new one, or drop the clause if there is none
        if is_real_note(existing):
            detail = detail.replace("{KEEP}", str(existing).strip().rstrip("."))
        else:
            detail = re.sub(r";?\s*previous note:\s*\{KEEP\}", "", detail).strip()
    if detail == "KEEP":
        if is_real_note(existing):
            parts.append(str(existing).strip().rstrip("."))
    elif detail:
        parts.append(detail)
    parts.append("Source: " + source)
    return " | ".join(parts)


def hidx(ws):
    return {c.value: i for i, c in enumerate(ws[1]) if c.value}


def last_row(ws):
    r = ws.max_row
    while r > 1 and all(c.value is None for c in ws[r]):
        r -= 1
    return r


def extend_validations(ws, new_last):
    by = {}
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            c0, r0, c1, r1 = range_boundaries(rng)
            by.setdefault((c0, dv.formula1, dv.type), dv)
    ws.data_validations.dataValidation = []
    for (col, f1, typ), dv in by.items():
        nd = DataValidation(type=typ, formula1=f1, allow_blank=True, operator=dv.operator)
        L = get_column_letter(col)
        nd.add(f"{L}2:{L}{new_last}")
        ws.add_data_validation(nd)


def sheet_values(ws):
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--partner", required=True)
    ap.add_argument("--plan", required=True)
    ap.add_argument("--tag", required=True)
    ap.add_argument("--source", required=True)
    ap.add_argument("--execute", action="store_true")
    a = ap.parse_args()
    main_path = RETURNED + rf"\{a.partner}_accessibility_report.xlsx"
    backup = ARCHIVE + rf"\{a.partner}_accessibility_report_pre_{a.tag}_2026-09-26.xlsx"
    tmp = main_path.replace(".xlsx", f"_tmp_{a.tag}.xlsx")
    plan = list(csv.DictReader(open(a.plan, encoding="utf-8-sig", newline="")))
    if not os.path.exists(backup):
        if not a.execute:
            src = main_path
        else:
            shutil.copy2(main_path, backup)
            src = backup
    else:
        src = backup
    if not a.execute and not os.path.exists(backup):
        print("dry run: using the live file as the 'before' (the backup is made on --execute)")
    wb = openpyxl.load_workbook(src)
    ws = wb["Ward Accessibility"]; ix = hidx(ws)
    master = {(r["State"], r["LGA"], r["Ward (GRID3)"]): r for r in csv.DictReader(open(MASTER, encoding="utf-8-sig", newline=""))}
    rows_by_key = {}
    for n in range(2, ws.max_row + 1):
        k = (ws.cell(n, ix["State"] + 1).value, ws.cell(n, ix["LGA"] + 1).value, ws.cell(n, ix["Ward (GRID3)"] + 1).value)
        if k[0]:
            rows_by_key.setdefault(k, []).append(n)
    log, n_upd, n_app = [], 0, 0
    nxt = last_row(ws) + 1
    info_col = "Why flagged" if "Why flagged" in ix else ("Note" if "Note" in ix else None)
    problems = []
    for p in plan:
        k = (p["state"], p["lga"], p["ward"]); date = datetime.strptime(p["date"], "%Y-%m-%d")
        if p["action"] == "update":
            rr = rows_by_key.get(k, [])
            if len(rr) != 1:
                problems.append(f"update {k}: matched {len(rr)} rows"); continue
            r = rr[0]
            old = {c: ws.cell(r, ix[c] + 1).value for c in EDIT_COLS}
            ws.cell(r, ix["Accessible (Y/N)"] + 1).value = p["accessible"]
            ws.cell(r, ix["Reason category"] + 1).value = p["reason_category"]
            ws.cell(r, ix["Reason notes"] + 1).value = note(old["Reason notes"], p["detail"], a.source)
            ws.cell(r, ix["Date reported"] + 1).value = date
            log.append(("update", r, k, old)); n_upd += 1
        elif p["action"] == "append":
            if k in rows_by_key:
                problems.append(f"append {k}: row already exists at {rows_by_key[k]}"); continue
            m = master.get(k)
            if m is None:
                problems.append(f"append {k}: not in the master"); continue
            r = nxt; nxt += 1
            vals = {"State": k[0], "LGA": k[1], "Ward (GRID3)": k[2], "Ward (OCHA/COD)": "NA", "Non-IDP clusters": int(m["Non-IDP clusters"] or 0), "IDP clusters": int(m["IDP clusters"] or 0),
                    "Total target HHs (primary)": int(float(m["Total target HHs (primary)"] or 0)), "Accessible (Y/N)": p["accessible"], "Reason category": p["reason_category"],
                    "Reason notes": note(None, p["detail"], a.source), "Date reported": date}
            if "Ward spans multiple LGAs (Y/N)" in ix:
                vals["Ward spans multiple LGAs (Y/N)"] = m["Ward spans multiple LGAs (Y/N)"]
            if info_col:
                vals[info_col] = f"Added 2026-09-26: this ward moved to your coverage (Chibok and Damboa, IMC to FACT on 25 Sep) and had no row in your file; row added to record your answer."
            for c, v in vals.items():
                if c in ix:
                    ws.cell(r, ix[c] + 1).value = v
            if "Date reported" in ix:
                # inherit the date cell format of the row above
                ws.cell(r, ix["Date reported"] + 1).number_format = ws.cell(r - 1, ix["Date reported"] + 1).number_format
            log.append(("append", r, k, None)); n_app += 1
        else:
            problems.append(f"unknown action {p['action']}")
    if problems:
        for x in problems: print("PLAN PROBLEM:", x)
        sys.exit("STOP: plan does not match the file; nothing written")
    if n_app:
        extend_validations(ws, nxt - 1)
    wb.save(tmp)
    # ---- diff vs the before-file
    wa = openpyxl.load_workbook(src); wb2 = openpyxl.load_workbook(tmp)
    ok = True
    for sh in wa.sheetnames:
        va, vb = sheet_values(wa[sh]), sheet_values(wb2[sh])
        d = {k: (va.get(k), vb.get(k)) for k in set(va) | set(vb) if va.get(k) != vb.get(k)}
        if sh != "Ward Accessibility":
            if d: ok = False; print("UNEXPECTED change on sheet", sh, len(d))
            continue
        rows_changed = sorted({k[0] for k in d})
        exp_rows = sorted({r for kind, r, k, o in log})
        if rows_changed != exp_rows: ok = False; print("row set differs: got", rows_changed[:8], "expected", exp_rows[:8])
        bad = [k for k in d if k[0] <= wa[sh].max_row and k[0] in {r for kind, r, kk, o in log if kind == "update"} and k[1] not in {ix[c] + 1 for c in EDIT_COLS}]
        if bad: ok = False; print("UNEXPECTED column touched:", bad[:5])
        for kind, r, k, old in log:
            if kind == "append" and any(wa[sh].cell(r, c).value is not None for c in range(1, wa[sh].max_column + 1)):
                ok = False; print("append row was not empty", r)
    dva = sorted((str(dv.sqref), dv.type, dv.formula1) for dv in wa["Ward Accessibility"].data_validations.dataValidation)
    print(f"plan: {n_upd} update(s), {n_app} append(s); diff vs before == plan: {ok}")
    for kind, r, k, old in log[:60]:
        print(f"  {kind:6s} row {r}: {k}" + (f"   before: {old['Accessible (Y/N)']!r}/{old['Reason category']!r}/{str(old['Reason notes'])[:30]!r}" if old else ""))
    if not ok:
        os.remove(tmp); sys.exit("STOP: diff is not exactly the plan; nothing written")
    if not a.execute:
        os.remove(tmp); print("dry run: nothing written"); return
    os.replace(tmp, main_path); print("written ->", main_path, "| backup:", backup)


if __name__ == "__main__":
    main()
