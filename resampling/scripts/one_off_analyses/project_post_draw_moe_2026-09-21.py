# ==============================================================================
# Pre-merge review aid for the 2026-09-21 resampling round: for every stratum
# in a partner's staged shortfall CSVs, project the RIGOROUS (Kish cluster-
# size-aware) MoE after merging the actually-drawn clusters, using the same
# realized_moe_unequal() formula as 05_build_accessibility_impact_workbook.py
# and the real drawn cluster sizes (not the m_used-sized units the search
# assumes - real PPS-weighted IDP sites are larger, which is why some "not
# closeable" strata clear). Read-only. Prints before -> after per stratum so
# Jack can approve a merge on what it actually buys, not on "households
# closed" from the draw log.
#
# Usage: python project_post_draw_moe_2026-09-21.py <Partner> <staging_dir>
# ==============================================================================
import csv
import math
import os
import re
import sys
from collections import defaultdict

import openpyxl

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
WORKBOOK = os.path.join(PROJECT_DIR, r"resampling\output\NGA_MSNA_2026_accessibility_impact_workbook.xlsx")
STRATA_FULL = os.path.join(PROJECT_DIR, r"output\data\data_collection\NGA_MSNA_2026_strata_level_sampling_frame_v10_FULL.csv")
CLUSTER_STATUS = os.path.join(PROJECT_DIR, r"output\data\data_collection\NGA_MSNA_2026_cluster_status_v10.csv")
TARGET_MOE_PCT = 10.0


def realized_moe_unequal(achieved, N_hh, sizes, ICC=0.06, Z=1.6448536269514722, p=0.5):
    sizes = [s for s in sizes if s and s > 0]
    if not sizes or achieved <= 0:
        return None
    if N_hh <= achieved:
        return 0.0
    m_bar = sum(sizes) / len(sizes)
    cv = (math.sqrt(sum((s - m_bar) ** 2 for s in sizes) / (len(sizes) - 1)) / m_bar) if len(sizes) > 1 else 0.0
    deff = 1 + ((cv ** 2 + 1) * m_bar - 1) * ICC
    n0 = achieved * (N_hh - 1) / (N_hh - achieved) / deff
    return math.sqrt(Z ** 2 * p * (1 - p) / n0) * 100 if n0 > 0 else None


def main(partner, staging_dir):
    slug = re.sub(r"[^A-Za-z0-9]", "_", partner).lower()
    shortfall_files = [os.path.join(staging_dir, f"{slug}_shortfalls.csv"), os.path.join(staging_dir, f"{slug}_shortfalls_idp.csv")]
    targets = set()
    for f in shortfall_files:
        if os.path.exists(f):
            targets |= {r["strata_id"] for r in csv.DictReader(open(f, encoding="utf-8"))}

    N_hh_design = {r["strata_id"]: float(r["N_hh"]) for r in csv.DictReader(open(STRATA_FULL, encoding="utf-8"))}
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Strata Level"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {h: i for i, h in enumerate(hdr) if h}
    strata = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[ix["Strata ID"]]:
            strata[r[ix["Strata ID"]]] = {"pct": r[ix["% of population remaining"]], "state": r[ix["State"]], "lga": r[ix["LGA"]],
                                          "pop": r[ix["Pop type"]], "feas": r[ix["Feasibility"]]}

    sizes_by = defaultdict(list)
    for c in csv.DictReader(open(CLUSTER_STATUS, encoding="utf-8")):
        sid = re.sub(r"_[^_]+$", "", c["cluster_id"])
        acc, ach = int(c["n_accessible_primary_post_threshold"]), int(c["n_achieved"])
        sizes_by[sid].append(max(acc, ach) if c["status"] in ("completed", "partially_completed_access_lost") else acc)

    # Sizes come from the STAMPED household files, not new_clusters.csv's
    # nominal target_households. Lesson from CRS (2026-09-21): projected 10 of
    # 15 strata to clear, only 4 did - the merge correctly held 172 rows
    # FULL-only (131 in inaccessible wards, 41 in below-threshold clusters),
    # none of which the nominal-size projection saw. Mirror the merge:
    # count only primary rows stamped Accessible, and drop a Non-IDP cluster
    # entirely if fewer than 4 of its primary rows survive. Run
    # stamp_ward_accessible_status.py on both staged files first.
    per_cluster = defaultdict(lambda: {"acc": 0, "pop": None})
    for fn in ("new_households.csv", "new_households_idp_sitelevel.csv"):
        p = os.path.join(staging_dir, fn)
        if not os.path.exists(p):
            continue
        for r in csv.DictReader(open(p, encoding="utf-8")):
            cid = r.get("cluster_id")
            if not cid or r.get("status") != "primary":
                continue
            c = per_cluster[cid]
            c["pop"] = "idp" if cid.startswith("idp_") else "non_idp"
            if (r.get("ward_accessible_status") or "").strip() == "Accessible":
                c["acc"] += 1
    new_by = defaultdict(list)
    for cid, c in per_cluster.items():
        n = c["acc"]
        if c["pop"] == "non_idp" and n < 4:
            n = 0
        if n > 0:
            new_by[re.sub(r"_[^_]+$", "", cid)].append(n)

    out = []
    for sid in sorted(targets):
        s = strata.get(sid)
        if s is None or not isinstance(s["pct"], (int, float)):
            continue
        N = N_hh_design[sid] * s["pct"] / 100
        base, new = sizes_by[sid], new_by.get(sid, [])
        before = realized_moe_unequal(sum(base), N, base)
        after = realized_moe_unequal(sum(base) + sum(new), N, base + new)
        out.append((s["state"], s["lga"], s["pop"], None if before is None else round(before, 1), len(new), sum(new),
                    None if after is None else round(after, 1), s["feas"].split(" - ")[0]))

    clears = sum(1 for o in out if o[6] is not None and o[6] <= TARGET_MOE_PCT)
    print(f"{partner}: {len(out)} strata drawn against; projected to clear {TARGET_MOE_PCT:.0f}% after merge: {clears}; still above: {len(out) - clears}")
    print("State | LGA | Pop | MoE before | new clusters | new HH | MoE after | category")
    for o in sorted(out, key=lambda x: (x[6] is not None and x[6] <= TARGET_MOE_PCT, x[0], x[1])):
        print(" | ".join(str(v) for v in o))


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("Usage: python project_post_draw_moe_2026-09-21.py <Partner> <staging_dir>")
    main(sys.argv[1], sys.argv[2])
