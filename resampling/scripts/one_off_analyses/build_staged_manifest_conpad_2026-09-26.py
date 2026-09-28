# ==============================================================================
# 2026-09-26: manifest + live-change list for the STAGED Conpad-accessibility release
# (resampling/output/staged_packages_2026-09-26_conpad_accessibility/). READ-ONLY on every live folder;
# writes only inside the staging root.
#
#   MANIFEST_staged_files.csv            every staged file: path, size, md5
#   LIVE_CHANGE_LIST_FACT_package.csv    per FACT LGA file: ADD / REPLACE / ARCHIVE / REMOVE / IDENTICAL vs the live package
#   WORKBOOK_COMPARISON_vs_live.csv      per partner and sheet: rows added / removed / changed cells, LGA sets
#   WORKBOOK_STRATA_SUMMARY_CHANGES.csv  every Strata Summary figure that moves, per partner
#
# Content comparison, not md5: Microsoft 365 rewrites xlsx/docx metadata, and rebuilt docx carry new timestamps.
# ==============================================================================
import csv
import hashlib
import os
import re
import zipfile
from collections import Counter, defaultdict

import openpyxl

B = r"C:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026"
ST = B + r"\1_sampling\resampling\output\staged_packages_2026-09-26_conpad_accessibility"
LIVE = r"C:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\3. External coordination\NGA MSNA 2026 Package"
WORKING = B + r"\1_sampling\output\data\data_collection\NGA_MSNA_2026_stage2_sampling_frame_v13_WORKING.csv"
GROUP_A = {"Dandume", "Faskari", "Funtua", "Malumfashi", "Matazu"}     # Conpad accessibility change
GROUP_B = {"Biu", "Kaga", "Batsari", "Baure", "Danja", "Daura", "Dutsin-Ma", "Kafur", "Mani", "Rimi"}  # completions since the live KML build
LGAS = [("Katsina", x) for x in sorted(GROUP_A | GROUP_B) if x not in ("Biu", "Kaga")] + [("Borno", "Biu"), ("Borno", "Kaga")]
ARCHIVE_DIR_NAME = "_archived_dropped_clusters_2026-09-27"


def lp(p):
    p = os.path.abspath(p)
    return p if p.startswith("\\\\?\\") else "\\\\?\\" + p


def md5(p):
    h = hashlib.md5()
    with open(lp(p), "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def walk_files(root, skip_archives=True):
    out = {}
    if not os.path.isdir(lp(root)):
        return out
    for dp, dn, fn in os.walk(lp(root)):
        if skip_archives:
            dn[:] = [d for d in dn if not d.startswith("_archive")]
        for f in fn:
            full = os.path.join(dp, f)
            out[os.path.relpath(full, lp(root))] = full
    return out


def docx_same(a, b):
    """Same content ignoring docProps timestamps: word/* parts and media compared by md5."""
    def parts(p):
        with zipfile.ZipFile(lp(p)) as z:
            return {n: hashlib.md5(z.read(n)).hexdigest() for n in z.namelist() if not n.startswith(("docProps/", "customXml", "[trash]")) and n != "[Content_Types].xml"}
    return parts(a) == parts(b)


# ---- 1. FACT package: staged vs live -------------------------------------------------------------------------
working_primary_clusters = set()
with open(lp(WORKING), encoding="utf-8-sig", newline="") as f:
    for r in csv.DictReader(f):
        if r["status"] == "primary" and r.get("sampling_method") != "MSNA Light":
            working_primary_clusters.add(r["cluster_id"])
rows = []
for state, lga in LGAS:
    grp = "A_conpad_accessibility" if lga in GROUP_A else "B_completions_since_last_kml_build"
    live_root = os.path.join(LIVE, "FACT", state, lga)
    st_root = os.path.join(ST, "FACT", state, lga)
    L, S = walk_files(live_root), walk_files(st_root)
    for rel in sorted(set(L) | set(S)):
        relp = "FACT/" + state + "/" + lga + "/" + rel.replace("\\", "/")
        l, s = L.get(rel), S.get(rel)
        if l and s:
            same = docx_same(l, s) if rel.endswith(".docx") else md5(l) == md5(s)
            act = "IDENTICAL" if same else "REPLACE"
            note = ""
        elif s:
            act, note = "ADD", ""
        else:
            if rel.endswith("_factsheet.docx"):
                cid = os.path.basename(rel)[: -len("_factsheet.docx")]
                if cid in working_primary_clusters:
                    act, note = "UNEXPECTED_MISSING_FROM_STAGED", "cluster still has a primary row in WORKING"
                else:
                    act, note = "ARCHIVE", f"move into Cluster_guide/{ARCHIVE_DIR_NAME}/ (cluster has no primary row left in WORKING; standing sweep rule, never deleted)"
            else:
                act, note = "REMOVE", "backup first; builder no longer emits this file (no such points/clusters left in WORKING)"
        rows.append({"lga": lga, "state": state, "group": grp, "path_under_package_FACT": relp, "action": act,
                     "live_size": os.path.getsize(l) if l else "", "staged_size": os.path.getsize(s) if s else "", "note": note})
os.makedirs(lp(ST), exist_ok=True)
with open(lp(os.path.join(ST, "LIVE_CHANGE_LIST_FACT_package.csv")), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
print("FACT package change list (files under the 15 LGA folders; existing _archive*/ folders ignored):")
c = Counter((r["group"], r["action"]) for r in rows)
for k in sorted(c):
    print(f"  {k[0]:38s} {k[1]:32s} {c[k]}")
print("  by LGA:")
byl = defaultdict(Counter)
for r in rows:
    byl[r["lga"]][r["action"]] += 1
for l in sorted(byl):
    print(f"    {l:12s} {dict(byl[l])}")

# ---- 2. workbooks: staged vs live -----------------------------------------------------------------------------
PARTNERS = ["CARE", "COOPI", "CRS", "DRC", "FACT", "FHI 360", "INTERSOS", "IRC", "LHI", "MDM", "Malteser", "NRC", "PLAN", "Save the Children", "Solidarités", "Street Child of Nigeria", "ZOA"]
KEYS = {"Strata Summary": ("State", "LGA", "Population Type"), "Sampling Points": ("Cluster ID", "Survey ID"),
        "Available to Collect": ("Cluster ID", "Survey ID"), "Cluster Summary": ("Cluster ID",)}


def read_sheet(path, name):
    wb = openpyxl.load_workbook(lp(path), read_only=True, data_only=True)
    if name not in wb.sheetnames:
        return None, None, wb.sheetnames
    it = wb[name].iter_rows(values_only=True)
    header = [h for h in next(it)]
    data = [dict(zip(header, r)) for r in it if any(v is not None for v in r)]
    return header, data, wb.sheetnames


def norm(v):
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, str):
        return v.strip()
    return v


cmp_rows, strata_changes = [], []
for p in PARTNERS:
    lw = os.path.join(LIVE, p, f"{p}_sampling_points_summary.xlsx")
    sw = os.path.join(ST, "workbooks", p, f"{p}_sampling_points_summary.xlsx")
    _, _, lnames = read_sheet(lw, "README")
    _, _, snames = read_sheet(sw, "README")
    for sheet in list(KEYS) + [n for n in snames if n.startswith("MSNA Light")]:
        key = KEYS.get(sheet)
        hl, dl, _ = read_sheet(lw, sheet)
        hs, ds, _ = read_sheet(sw, sheet)
        if dl is None or ds is None:
            cmp_rows.append({"partner": p, "sheet": sheet, "note": f"sheet missing (live {dl is not None}, staged {ds is not None})"})
            continue
        if key is None:
            key = tuple(h for h in ("Cluster ID", "Survey ID") if h in hs) or (hs[0],)
        kf = lambda r: tuple(str(r.get(k) or "") for k in key)
        Dl, Ds = defaultdict(list), defaultdict(list)
        for r in dl: Dl[kf(r)].append(r)
        for r in ds: Ds[kf(r)].append(r)
        added, removed = set(Ds) - set(Dl), set(Dl) - set(Ds)
        changed, cols = 0, Counter()
        for k in set(Dl) & set(Ds):
            a, b = Dl[k], Ds[k]
            if len(a) != len(b):
                changed += 1; cols["(row count for key)"] += 1; continue
            rowchg = False
            for ra, rb in zip(a, b):
                for h in hs:
                    if h in ra and norm(ra.get(h)) != norm(rb.get(h)):
                        cols[h] += 1; rowchg = True
                        if sheet == "Strata Summary":
                            strata_changes.append({"partner": p, "state": rb.get("State"), "lga": rb.get("LGA"), "pop_type": rb.get("Population Type"), "column": h, "live": ra.get(h), "staged": rb.get(h)})
            changed += rowchg
        lga_l = {str(r.get("LGA")) for r in dl if r.get("LGA")}; lga_s = {str(r.get("LGA")) for r in ds if r.get("LGA")}
        cmp_rows.append({"partner": p, "sheet": sheet, "rows_live": len(dl), "rows_staged": len(ds), "keys_added": len(added), "keys_removed": len(removed),
                         "rows_with_changed_cells": changed, "columns_that_moved": "; ".join(f"{h}={n}" for h, n in cols.most_common()),
                         "LGA_set_identical": lga_l == lga_s, "note": ""})
        for k in list(added)[:0]:
            pass
    print(f"{p:26s} " + " | ".join(f"{r['sheet'][:14]}: +{r.get('keys_added','?')}/-{r.get('keys_removed','?')}/~{r.get('rows_with_changed_cells','?')}" for r in cmp_rows if r["partner"] == p))
with open(lp(os.path.join(ST, "WORKBOOK_COMPARISON_vs_live.csv")), "w", encoding="utf-8-sig", newline="") as f:
    fn_ = ["partner", "sheet", "rows_live", "rows_staged", "keys_added", "keys_removed", "rows_with_changed_cells", "columns_that_moved", "LGA_set_identical", "note"]
    w = csv.DictWriter(f, fieldnames=fn_)
    w.writeheader(); w.writerows(cmp_rows)
with open(lp(os.path.join(ST, "WORKBOOK_STRATA_SUMMARY_CHANGES.csv")), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["partner", "state", "lga", "pop_type", "column", "live", "staged"])
    w.writeheader(); w.writerows(strata_changes)

# ---- 3. md5 manifest of every staged file ------------------------------------------------------------------------
man = []
for dp, dn, fn in os.walk(lp(ST)):
    for f_ in fn:
        full = os.path.join(dp, f_)
        rel = os.path.relpath(full, lp(ST)).replace("\\", "/")
        if rel.startswith("MANIFEST_") or rel in ("LIVE_CHANGE_LIST_FACT_package.csv", "WORKBOOK_COMPARISON_vs_live.csv", "WORKBOOK_STRATA_SUMMARY_CHANGES.csv"):
            continue
        man.append((rel, os.path.getsize(full), md5(full)))
man.sort()
with open(lp(os.path.join(ST, "MANIFEST_staged_files.csv")), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f); w.writerow(["path", "size_bytes", "md5"]); w.writerows(man)
print(f"\nMANIFEST_staged_files.csv: {len(man)} files, {sum(m[1] for m in man)/1e6:.1f} MB")
