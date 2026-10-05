# ==============================================================================
# 2026-10-04 STAGING (dry run by default): apply FACT's accessibility reports for
# the LGAs FACT takes over on 4 Oct, as log rows in resampling_requests_log.csv,
# AFTER the reallocation and ONLY for those LGAs (Jack, 4 Oct; relayed by the
# Coordinator - live use needs his own word in Resampling's window).
#
# Source: one email from FACT (Isah Suleiman, isah@factfoundation-int.org,
# 4 Oct 2026 07:18, "Re: Follow up: Request for Overall Cluster Verification"),
# with 7 attached reports named after the OUTGOING partners (FACT filled our
# templates for those partners' LGAs). Saved under
# resampling/input/partner_raw_comms/04102026_bigresampling/.
#
# Rules:
#   - only FACT's 13 new LGAs (Tangaza goes to ZOA, not FACT; Mashi has no report: Jack D8,
#     it keeps its current statuses);
#   - EDIT-ONLY (Coordinator's D6 finding, re-verified by Resampling 4 Oct ~21:30): a row is
#     FACT's answer only if FACT changed its Accessible / Reason category / Reason notes /
#     Date reported versus the file FACT was given (BASES below). An untouched pre-filled
#     value is the template's / outgoing partner's, not FACT's, and is never cited as FACT's.
#     The "X (1).xlsx" files are NOT newer versions: a second FACT person filled the same
#     template for OTHER LGAs (IRC: LiGHuD did Zuru, Talatu did Tsafe; Solidarites: LiGHuD
#     did Gwandu, Talatu did Anka), so every file is used. Two files editing the same row
#     differently stops the run;
#   - a blank answer (even if "edited") leaves the current status as it is (Jack D7);
#   - date_reported_by_partner is written as YYYY-MM-DD. Talatu's two files store the date she typed,
#     03/10/2026 (3 Oct), month-first as 10 Mar 2026 (cell format mm-dd-yy): corrected to 2026-10-03 with a
#     note. LiGHuD's files have no date: the date of FACT's email (2026-10-04) is used, with a note. Any
#     other date outside 2026-09-01..2026-10-04 stops the run. (The master and the overlay order rows by
#     request_id, so this field never changes a status; it feeds the "date reported" display columns.)
#   - OPTION --d7-slivers (Jack's D7 decision, not yet taken): the CARE-template wards FACT left blank in
#     Argungu / Birnin Kebbi / Suru are border slivers with 0 clusters. Where the SAME ward's portion in a
#     neighbouring LGA is Inaccessible (after tonight's changes), an IMPACT-default "No" row (partner FACT,
#     reported_by IMPACT) is added for the sliver, citing the neighbouring report, for FACT to confirm.
#     Stops if a sliver has any row in the FULL frame (the note says 0 clusters). With --d7-require-evidence,
#     a neighbouring portion only counts if its report gives a reason AND a known date (drops Jadadi and
#     Lani Shiba, which rest on a CRS row with no reason and an unknown date: 17 -> 15).
#   - OPTION --isa-gapfill FILE (Jack D4, 4 Oct): Bama + Kala/Balge (now "PLAN, FACT", no FACT report tonight) default
#     to Isa Jugudum's FACT Borno review of 21 Sep, ONLY for wards where PLAN has no row (PLAN's own rows are kept, so
#     Soye stays open on PLAN's Yes). Jack's rule: "if any of the LGAs absorbed by FACT tonight doesn't have an
#     accessibility report and is Borno, default back to Isa's full Borno report from last week."
#   - OPTION --isa-imc-docs FILE (Jack D5, 4 Oct): a FACT "No" row (Isa, 21 Sep) for every Chibok/Damboa ward that is
#     Inaccessible only on IMC's report; Isa marked them all Inaccessible, so the status does not change - the rows
#     document FACT's own position. Stops if Isa's answer for any of them is not Inaccessible, or if the set is not 14.
#   - every row is status "new", like every other row in the log.
#   - ward level: 04_build_master_accessibility_status.py takes each partner's latest
#     row per ward and marks the ward Inaccessible if ANY partner's says No. So FACT's
#     row is appended, and where the OUTGOING partner's latest row for the same ward
#     disagrees with FACT, a superseding row for that partner (same answer as FACT,
#     supersedes_request_id = its old row) is appended too - otherwise its stale No
#     would keep a ward FACT now reports accessible closed (precedent: the 14 Sep
#     Mairari rows);
#   - cluster level: build_cluster_accessibility_overlay.py takes the latest row per
#     (cluster, ward) regardless of partner, so FACT's row alone is enough.
# Every appended row says where it came from (Jack: "we should be clear in our
# internal documentation that any updated information is from this change").
#
# Usage: python stage_fact_accessibility_after_reallocation_2026-10-04.py [--d7-slivers [--d7-require-evidence]]
#            [--isa-gapfill ISA.xlsx] [--isa-imc-docs ISA.xlsx] [--execute]
#   dry run: writes rows_to_append.csv + predicted_changes.csv to
#            resampling/output/staged_fact_accessibility_2026-10-04/ only.
#   --execute: backs up the log, appends the rows, verifies (needs Jack's own go).
# ==============================================================================
import argparse
import csv
import glob
import hashlib
import os
import shutil
import sys
from collections import defaultdict
from datetime import date

import openpyxl

RS = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
IN_DIR = os.path.join(RS, "input", "partner_raw_comms", "04102026_bigresampling")
LOG = os.path.join(RS, "output", "resampling_requests_log.csv")
MASTER = os.path.join(RS, "output", "master_accessibility_status_ward_level.csv")
OUT = os.path.join(RS, "output", "staged_fact_accessibility_2026-10-04")
TODAY = date.today().isoformat()
FACT_LGAS = {"Madagali", "Mafa", "Mobbar", "Monguno", "Batagarawa", "Jibia", "Zuru", "Gwandu",
             "Birnin Kebbi", "Suru", "Argungu", "Tsafe", "Anka"}
PKG = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(RS)))), "3. External coordination", "NGA MSNA 2026 Package")
# file FACT sent -> the file FACT was given (cell-diff verified 4 Oct: same row set, fewest differences). CARE's own
# 28 Sep return (partner_raw_comms/CARE/CARE_accessibility_report_2809.xlsx) is cloud-only and OneDrive is not serving
# downloads tonight, so our returned CARE copy (which merged that 28 Sep return on 29 Sep) stands in; in the 3 moved
# LGAs every difference against it is a blank FACT filled, so nothing is misattributed.
BASES = {
    "CARE_accessibility_report_2809.xlsx": os.path.join(RS, "input", "accessibility_reports_returned", "CARE_accessibility_report.xlsx"),
    "FHI 360_accessibility_report.xlsx": os.path.join(PKG, "FHI 360", "FHI 360_accessibility_report.xlsx"),
    "IRC_accessibility_report.xlsx": os.path.join(PKG, "IRC", "IRC_accessibility_report.xlsx"),
    "IRC_accessibility_report (1).xlsx": os.path.join(PKG, "IRC", "IRC_accessibility_report.xlsx"),
    "Solidarités_accessibility_report.xlsx": os.path.join(PKG, "Solidarités", "Solidarités_accessibility_report.xlsx"),
    "Solidarités_accessibility_report (1).xlsx": os.path.join(PKG, "Solidarités", "Solidarités_accessibility_report.xlsx"),
    "Street Child of Nigeria_accessibility_report.xlsx": os.path.join(PKG, "Street Child of Nigeria", "Street Child of Nigeria_accessibility_report.xlsx"),
}
FIELDS = ("Accessible (Y/N)", "Reason category", "Reason notes", "Date reported")
EMAIL_DATE = "2026-10-04"
TALATU_FILES = {"IRC_accessibility_report (1).xlsx", "Solidarités_accessibility_report (1).xlsx"}
D7_LGAS = {"Argungu", "Birnin Kebbi", "Suru"}
ISA_GAPFILL_LGAS = {"Bama", "Kala/Balge"}
ISA_O, ISA_Q = "FACT's Assessed Accessibility Status", "FACT Reasons / Notes"
ISA_NO_REASON = "No reason given in FACT's Borno review"  # the 26 Sep wording
ISA_YES_NOTE = "Accessible in FACT's Borno review"


def read_isa(path):
    """Isa's 'Borno Ward Review' sheet -> {(LGA, ward): row}, plus a provenance string."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    raw = list(wb["Borno Ward Review"].iter_rows(values_only=True))
    wb.close()
    hi = next(i for i, r in enumerate(raw) if sum(c not in (None, "") for c in r) >= 5)
    h = [norm(c) for c in raw[hi]]
    if any(x not in h for x in ("LGA", "Ward (GRID3)", ISA_O, ISA_Q)):
        raise SystemExit(f"STOP: {path} is not Isa's review layout (header {h})")
    out = {}
    for r in raw[hi + 1:]:
        d = dict(zip(h, r))
        if norm(d.get("LGA")):
            k = (norm(d["LGA"]), norm(d["Ward (GRID3)"]))
            if k in out:
                raise SystemExit(f"STOP: Isa's sheet has {k} twice")
            out[k] = d
    return out, f"{os.path.basename(path)} md5 {md5(path)[:8]}"


def isa_answer(d):
    return {"accessible": "Yes", "inaccessible": "No"}.get(norm(d.get(ISA_O)).lower())


def isa_reason(note, a):
    n = norm(note).lower()
    if a == "Yes":
        return "N/A - fully accessible"
    if "insecur" in n or "conflict" in n:
        return "Insecurity / conflict"
    if "flood" in n or "terrain" in n:
        return "Physical access (terrain, flooding, roads)"
    return "Other"  # incl. "only wards in garrison town are accessible" and no reason (the 26 Sep convention)


def reported_date(fn, v):
    """(YYYY-MM-DD, note to append) for a 'Date reported' cell of file fn."""
    import datetime as _dt
    if v in (None, "") or norm(v) == "":
        return EMAIL_DATE, " [no date entered in the file; date of FACT's email]"
    if isinstance(v, (_dt.datetime, _dt.date)):
        d = (v.date() if isinstance(v, _dt.datetime) else v).isoformat()
    else:
        d = norm(v)[:10]
    note = ""
    if fn in TALATU_FILES and d == "2026-03-10":
        d, note = "2026-10-03", " [typed 03/10/2026 (3 Oct); Excel stored it month-first as 10 Mar 2026]"
    if not ("2026-09-01" <= d <= EMAIL_DATE):
        raise SystemExit(f"STOP: implausible 'Date reported' {v!r} in {fn} - check the cell before applying")
    return d, note


def norm(v):
    return "" if v is None or str(v).strip().lower() in ("", "none", "nan") else str(v).strip()


def edited(row, base_row):
    return base_row is None or any(norm(row.get(x)) != norm(base_row.get(x)) for x in FIELDS)


SOURCE = ("FACT accessibility report for its newly assigned LGAs (LGA moved to FACT on 2026-10-04), "
          "received 2026-10-04 07:18 from Isah Suleiman (isah@factfoundation-int.org), "
          "'Re: Follow up: Request for Overall Cluster Verification', file {file}")


def md5(p):
    return hashlib.md5(open(p, "rb").read()).hexdigest()


def answer(v):
    s = str(v).strip().lower() if v not in (None, "") else ""
    return {"yes": "Yes", "y": "Yes", "no": "No", "n": "No"}.get(s)


def read_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    h = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(rows[0])]
    return [dict(zip(h, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--execute", action="store_true")
    ap.add_argument("--d7-slivers", action="store_true", help="OPTION (Jack D7): IMPACT-default No for blank border slivers")
    ap.add_argument("--d7-require-evidence", action="store_true",
                    help="with --d7-slivers: a neighbouring portion counts only if its report has a reason and a known date")
    ap.add_argument("--isa-gapfill", metavar="ISA_XLSX", help="OPTION (Jack D4): Bama + Kala/Balge from Isa's review where PLAN has no row")
    ap.add_argument("--isa-imc-docs", metavar="ISA_XLSX", help="OPTION (Jack D5): FACT No rows for the Chibok/Damboa wards closed only on IMC's word")
    args = ap.parse_args()
    if args.d7_require_evidence and not args.d7_slivers:
        raise SystemExit("--d7-require-evidence only applies with --d7-slivers")
    if args.isa_gapfill and args.isa_imc_docs and os.path.abspath(args.isa_gapfill) != os.path.abspath(args.isa_imc_docs):
        raise SystemExit("--isa-gapfill and --isa-imc-docs must name the same file")
    os.makedirs(OUT, exist_ok=True)
    with open(LOG, encoding="utf-8-sig", newline="") as f:
        rd = csv.DictReader(f)
        cols, log = rd.fieldnames, list(rd)
    next_id = max(int(r["request_id"]) for r in log) + 1
    latest_ward = {}
    for r in log:
        if r["report_level"] == "ward":
            k = (r["partner"], r["state"], r["lga"], r["ward_name"])
            if k not in latest_ward or int(r["request_id"]) > int(latest_ward[k]["request_id"]):
                latest_ward[k] = r

    new_rows = []

    def add(**kw):
        nonlocal next_id
        row = {c: "" for c in cols}
        row.update(kw, request_id=str(next_id), date_added=TODAY, status="new")
        new_rows.append(row)
        next_id += 1
        return row["request_id"]

    files = sorted(glob.glob(os.path.join(IN_DIR, "*.xlsx")))
    if {os.path.basename(f) for f in files} != set(BASES):
        raise SystemExit(f"STOP: report files differ from the 7 reviewed: {sorted({os.path.basename(f) for f in files} ^ set(BASES))}")
    claimed = {}  # (sheet, key) -> (file, answer): stop if two files edit the same row differently
    n_untouched = 0
    for f in files:
        fn = os.path.basename(f)
        src = SOURCE.format(file=fn)
        base_w = {(r["LGA"], r["Ward (GRID3)"]): r for r in read_sheet(BASES[fn], "Ward Accessibility")}
        base_c = {r["Cluster ID"]: r for r in read_sheet(BASES[fn], "Cluster Accessibility") if r.get("Cluster ID")}
        for r in read_sheet(f, "Ward Accessibility"):
            if r["LGA"] not in FACT_LGAS:
                continue
            if not edited(r, base_w.get((r["LGA"], r["Ward (GRID3)"]))):
                n_untouched += 1
                continue
            if not answer(r.get("Accessible (Y/N)")):
                continue
            k = ("ward", r["LGA"], r["Ward (GRID3)"])
            if k in claimed and claimed[k][1] != answer(r["Accessible (Y/N)"]):
                raise SystemExit(f"STOP: {k} edited differently in {claimed[k][0]} and {fn}")
            claimed[k] = (fn, answer(r["Accessible (Y/N)"]))
            a = answer(r["Accessible (Y/N)"])
            reason = r.get("Reason category") or ("N/A - fully accessible" if a == "Yes" else "")
            d, dnote = reported_date(fn, r.get("Date reported"))
            add(partner="FACT", report_level="ward", state=r["State"], lga=r["LGA"], ward_name=r["Ward (GRID3)"],
                source_channel="Partner's own template (FACT filled the outgoing partner's template)", reported_by="Partner",
                accessible=a, reason_category=reason, reason_notes=((str(r.get("Reason notes")) + " | ") if r.get("Reason notes") else "") + src + dnote,
                pct_target_achieved=str(r.get("% of target achieved so far") or ""),
                date_reported_by_partner=d)
            for (p, st, lga, ward), old in latest_ward.items():
                if (st, lga, ward) == (r["State"], r["LGA"], r["Ward (GRID3)"]) and p != "FACT" and old["accessible"].strip().capitalize() != a:
                    add(partner=p, report_level="ward", state=st, lga=lga, ward_name=ward, source_channel="Reallocation supersession",
                        reported_by="IMPACT (reallocation)", accessible=a, reason_category=reason,
                        reason_notes=(f"Superseded: this LGA moved from {p} to FACT on 2026-10-04, and FACT's own report now answers "
                                      f"for it ({a}). {p}'s earlier answer ({old['accessible']}, request {old['request_id']}) no longer "
                                      f"applies. " + src),
                        date_reported_by_partner=d, supersedes_request_id=old["request_id"])
        for r in read_sheet(f, "Cluster Accessibility"):
            if r["LGA"] not in FACT_LGAS or not r.get("Cluster ID"):
                continue
            if not edited(r, base_c.get(r["Cluster ID"])):
                n_untouched += 1
                continue
            if not answer(r.get("Accessible (Y/N)")):
                continue
            k = ("cluster", r["Cluster ID"])
            if k in claimed and claimed[k][1] != answer(r["Accessible (Y/N)"]):
                raise SystemExit(f"STOP: {k} edited differently in {claimed[k][0]} and {fn}")
            claimed[k] = (fn, answer(r["Accessible (Y/N)"]))
            a = answer(r["Accessible (Y/N)"])
            d, dnote = reported_date(fn, r.get("Date reported"))
            add(partner="FACT", report_level="cluster", state=r["State"], lga=r["LGA"], ward_name=r["Ward (GRID3)"],
                cluster_id=r["Cluster ID"], pop_type=r.get("Pop Type") or "",
                source_channel="Partner's own template (FACT filled the outgoing partner's template)", reported_by="Partner",
                accessible=a, reason_category=r.get("Reason category") or ("N/A - fully accessible" if a == "Yes" else ""),
                reason_notes=((str(r.get("Reason notes")) + " | ") if r.get("Reason notes") else "") + src + dnote,
                pct_target_achieved=str(r.get("% of target achieved so far") or ""), date_reported_by_partner=d)

    print(f"edit-only: {len(claimed)} rows FACT edited with an answer; {n_untouched} untouched template rows ignored")
    with open(MASTER, encoding="utf-8-sig", newline="") as f:
        master_rows = {(r["State"], r["LGA"], r["Ward (GRID3)"]): r for r in csv.DictReader(f)}

    def predicted_status(rows):
        lw_ = {}
        for r in rows:
            if r["report_level"] == "ward":
                k = (r["partner"], r["state"], r["lga"], r["ward_name"])
                if k not in lw_ or int(r["request_id"]) > int(lw_[k]["request_id"]):
                    lw_[k] = r
        bw = defaultdict(list)
        for (p, st, lga, ward), r in lw_.items():
            bw[(st, lga, ward)].append(r)
        return {k: ("Inaccessible" if any(r["accessible"].strip().lower() == "no" for r in bw[k]) else "Accessible") if bw.get(k)
                else m["Accessible status"] for k, m in master_rows.items()}

    variant = "edit_only"
    if args.d7_slivers:
        variant = "edit_only_with_d7" + ("_evidence_only" if args.d7_require_evidence else "")
        after_fact = predicted_status(log + new_rows)
        full_rows = defaultdict(int)
        with open(os.path.join(os.path.dirname(RS), "output", "data", "data_collection",
                               "NGA_MSNA_2026_stage2_sampling_frame_v14_FULL.csv"), encoding="utf-8-sig", newline="") as f:
            for x in csv.DictReader(f):
                full_rows[(x["adm1_name"], x["adm2_name"], x["adm3_name"])] += 1

        def has_evidence(m):
            d = norm(m.get("Last reported date"))
            return bool(norm(m.get("Reason category"))) and bool(d) and not d.lower().startswith("unknown")

        care = os.path.join(IN_DIR, "CARE_accessibility_report_2809.xlsx")
        n_d7, weak = 0, []
        for r in read_sheet(care, "Ward Accessibility"):
            if r["LGA"] not in D7_LGAS or answer(r.get("Accessible (Y/N)")):
                continue
            k = (r["State"], r["LGA"], r["Ward (GRID3)"])
            others = [x.strip() for x in norm(master_rows.get(k, {}).get("Other LGA(s) sharing this ward")).split(";") if x.strip()]
            closed = [(o, master_rows[(r["State"], o, r["Ward (GRID3)"])]) for o in others
                      if (r["State"], o, r["Ward (GRID3)"]) in master_rows and after_fact[(r["State"], o, r["Ward (GRID3)"])] == "Inaccessible"]
            if closed and args.d7_require_evidence and not any(has_evidence(m) for _, m in closed):
                weak.append(f"{r['LGA']}/{r['Ward (GRID3)']}")
                continue
            if args.d7_require_evidence:
                closed = [(o, m) for o, m in closed if has_evidence(m)]
            if not closed:
                continue
            if full_rows.get(k):
                raise SystemExit(f"STOP: D7 sliver {k} has {full_rows[k]} FULL-frame rows - it is not an empty sliver; review before using the option")
            why = "; ".join(sorted({norm(m.get("Reason category")) or "no reason given" for _, m in closed}))
            cite = "; ".join(f"{o} portion Inaccessible ({norm(m.get('Reason category')) or 'no reason'}; reported by "
                             f"{norm(m.get('Reporting partner(s)')) or 'n/a'}, {norm(m.get('Last reported date')) or 'n/a'})" for o, m in closed)
            add(partner="FACT", report_level="ward", state=r["State"], lga=r["LGA"], ward_name=r["Ward (GRID3)"],
                source_channel="IMPACT default (D7 border sliver)",
                reported_by="IMPACT (default - border portion of a ward whose neighbouring portion is inaccessible; FACT to confirm)",
                accessible="No", reason_category=why,
                reason_notes=(f"Jack D7 option, 2026-10-04: FACT left this border sliver blank (0 clusters, 0 target HHs) in the CARE template it "
                              f"returned on 2026-10-04; the same ward's {cite}. Treated as inaccessible by IMPACT default until FACT confirms."),
                date_reported_by_partner=TODAY)
            n_d7 += 1
        print(f"D7 option: {n_d7} border sliver(s) set Inaccessible by IMPACT default"
              + (f"; {len(weak)} skipped for no reason/date on the neighbouring report: {', '.join(weak)}" if args.d7_require_evidence else ""))

    if args.isa_gapfill or args.isa_imc_docs:
        isa, isa_src = read_isa(args.isa_gapfill or args.isa_imc_docs)
        variant += "_with_isa"
        borno = defaultdict(dict)  # (lga, ward) -> {partner: latest logged ward row}, before tonight's rows
        for (p, st, lga, ward), r in latest_ward.items():
            if st == "Borno":
                borno[(lga, ward)][p] = r
        src_isa = f"Source: Isa Jugudum (FACT) Borno review of 21 Sep ({isa_src})"
        if args.isa_gapfill:
            n_gap, n_plan = 0, 0
            for (lga, ward), d in sorted(isa.items()):
                a = isa_answer(d)
                if lga not in ISA_GAPFILL_LGAS or not a:
                    continue
                if "PLAN" in borno.get((lga, ward), {}):
                    n_plan += 1
                    continue
                add(partner="FACT", report_level="ward", state="Borno", lga=lga, ward_name=ward,
                    source_channel="Partner's own template", reported_by="Partner", accessible=a,
                    reason_category=isa_reason(d.get(ISA_Q), a),
                    reason_notes=(f"{norm(d.get(ISA_Q)) or (ISA_NO_REASON if a == 'No' else ISA_YES_NOTE)} | {src_isa} | "
                                  f"applied 2026-10-04 under Jack's rule for Borno LGAs FACT absorbed without a new report "
                                  f"(LGA now PLAN, FACT); gap-fill (Jack D4): only wards PLAN has not reported, PLAN's own rows kept"),
                    date_reported_by_partner="2026-09-21")
                n_gap += 1
            print(f"Isa gap-fill (D4): {n_gap} FACT row(s) in Bama/Kala-Balge; {n_plan} ward(s) skipped because PLAN has a row")
        if args.isa_imc_docs:
            docs = []
            for (lga, ward), reps in sorted(borno.items()):
                imc = reps.get("IMC")
                if lga not in ("Chibok", "Damboa") or not imc or imc["accessible"].strip().lower() != "no" \
                        or any(p != "IMC" and x["accessible"].strip().lower() == "no" for p, x in reps.items()):
                    continue
                d = isa.get((lga, ward), {})
                if isa_answer(d) != "No":
                    raise SystemExit(f"STOP: Isa's answer for {lga}/{ward} (closed only on IMC's word) is {norm(d.get(ISA_O))!r}, "
                                     f"not Inaccessible - Jack's D5 was decided on all 14 being Inaccessible")
                docs.append((lga, ward, d, imc))
            if len(docs) != 14:
                raise SystemExit(f"STOP: {len(docs)} Chibok/Damboa wards are closed only on IMC's word, not the 14 Jack decided on - re-confirm")
            for lga, ward, d, imc in docs:
                add(partner="FACT", report_level="ward", state="Borno", lga=lga, ward_name=ward,
                    source_channel="Partner's own template", reported_by="Partner", accessible="No",
                    reason_category=isa_reason(d.get(ISA_Q), "No"),
                    reason_notes=(f"{norm(d.get(ISA_Q)) or ISA_NO_REASON} | {src_isa} | "
                                  f"Jack D5, 2026-10-04: documents FACT's own position on a ward closed so far only on IMC's report "
                                  f"(request {imc['request_id']}, {imc['date_added']}); no status change"),
                    date_reported_by_partner="2026-09-21")
            print(f"Isa IMC-only documentation rows (D5): {len(docs)}")
    with open(os.path.join(OUT, f"rows_to_append_{variant}.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(new_rows)

    # ---- predict: 04's ward rule and the overlay's cluster rule on log + new rows ----
    allr = log + new_rows
    lw = {}
    for r in allr:
        if r["report_level"] == "ward":
            k = (r["partner"], r["state"], r["lga"], r["ward_name"])
            if k not in lw or int(r["request_id"]) > int(lw[k]["request_id"]):
                lw[k] = r
    by_ward = defaultdict(list)
    for (p, st, lga, ward), r in lw.items():
        by_ward[(st, lga, ward)].append(r)
    with open(MASTER, encoding="utf-8-sig", newline="") as f:
        master = {(r["State"], r["LGA"], r["Ward (GRID3)"]): r["Accessible status"] for r in csv.DictReader(f)}
    pred = []
    touched = FACT_LGAS | {r["lga"] for r in new_rows}
    for k, now in master.items():
        if k[1] not in touched or not by_ward.get(k):
            continue
        new = "Inaccessible" if any(r["accessible"].strip().lower() == "no" for r in by_ward[k]) else "Accessible"
        if new != now:
            pred.append({"level": "ward", "state": k[0], "lga": k[1], "ward_or_cluster": k[2], "now": now, "after": new})

    def cluster_excluded(rows):
        latest = {}
        for r in rows:
            if r["report_level"] == "cluster" and r["cluster_id"]:
                k = (r["cluster_id"], r["ward_name"])
                if k not in latest or int(r["request_id"]) > int(latest[k]["request_id"]):
                    latest[k] = r
        return {cid for (cid, _), r in latest.items() if r["accessible"].strip().lower() == "no"}

    before, after = cluster_excluded(log), cluster_excluded(allr)
    for cid in sorted(after - before):
        pred.append({"level": "cluster", "state": "", "lga": "", "ward_or_cluster": cid, "now": "not excluded", "after": "excluded"})
    for cid in sorted(before - after):
        pred.append({"level": "cluster", "state": "", "lga": "", "ward_or_cluster": cid, "now": "excluded", "after": "not excluded"})
    with open(os.path.join(OUT, f"predicted_changes_{variant}.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["level", "state", "lga", "ward_or_cluster", "now", "after"])
        w.writeheader()
        w.writerows(pred)
    n_fact = sum(1 for r in new_rows if r["partner"] == "FACT")
    print(f"rows to append: {len(new_rows)} ({n_fact} FACT, {len(new_rows) - n_fact} superseding outgoing-partner rows)")
    print(f"predicted master changes: {sum(1 for p in pred if p['level'] == 'ward')} ward(s), "
          f"{sum(1 for p in pred if p['level'] == 'cluster')} cluster overlay change(s)")
    for p in pred:
        print("  ", p)

    if not args.execute:
        print("\n(dry run) nothing live written; staged files in", OUT)
        return
    bak = os.path.join(RS, "output", "_archive", f"{TODAY}_pre_fact_accessibility_after_reallocation")
    os.makedirs(bak, exist_ok=True)
    shutil.copy2(LOG, os.path.join(bak, os.path.basename(LOG)))
    if md5(LOG) != md5(os.path.join(bak, os.path.basename(LOG))):
        raise SystemExit("STOP: log backup does not match")
    with open(LOG, "a", encoding="utf-8", newline="") as f:
        csv.DictWriter(f, fieldnames=cols).writerows(new_rows)
    with open(LOG, encoding="utf-8-sig", newline="") as f:
        n_after = sum(1 for _ in csv.DictReader(f))
    if n_after != len(log) + len(new_rows):
        raise SystemExit(f"STOP: log has {n_after} rows, expected {len(log) + len(new_rows)} - restore from {bak}")
    print(f"appended {len(new_rows)} rows; log backup in {bak}. Next: 04 master build -> layer -> site frame -> overlay -> chain.")


if __name__ == "__main__":
    main()
