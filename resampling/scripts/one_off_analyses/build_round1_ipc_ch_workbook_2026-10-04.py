# Builds the Round 1 representativity summary workbook for the IPC/CH analysts (MSNA N-WEC 2026, 4 Oct 2026).
#
# Jack's decisions (4 Oct ~21:10): categories (a)-(d) on the SIMPLIFIED Round 1 margin-of-error rule; basis = the data
# officer's final cleaned Round 1 dataset; Kaduna reported at LGA level only (no state row).
#
# Source: the Round 1 tables re-run on the candidate aligned to the DO's cleaned dataset (verified 29/29 on 4 Oct; its
# achieved households in coverage are exactly the DO's cleaned households in coverage, 21,428 = 21,428). Every figure is
# recomputed here from the strata table, and the build STOPS unless the recomputation reproduces the verified strata,
# LGA and State tables exactly. Only then are the IPC/CH categories applied and the workbook written.
#
# Env: R1_MSNA_LIGHT = "not_assessed" (default: MSNA Light strata -> (d), left out of aggregates, matching the weighted
#      tables, which exclude MSNA Light) or "keep" (keep their Round 1 label, with a flag).
#      R1_IPC_FINAL = "1" drops the _DRAFT suffix from the file name (only after review and independent verification).
import csv, datetime, hashlib, math, os, sys, collections
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
WS = os.path.normpath(os.path.join(HERE, "..", "..", "..", ".."))
CAND = os.path.join(WS, "1_sampling", "resampling", "output", "round1_candidate_aligned_to_DO_clean_2026-10-04", "package_staging")
REP = os.path.join(CAND, "03_representativity")
STRATA_CSV = os.path.join(REP, "R1_representativity_strata_NE_NW.csv")
LGA_CSV = os.path.join(REP, "R1_representativity_lga_NE_NW.csv")
STATE_CSV = os.path.join(REP, "R1_representativity_state_NE_NW.csv")
FRAME_CSV = os.path.join(WS, "1_sampling", "output", "data", "data_collection", "NGA_MSNA_2026_strata_level_sampling_frame_v14_FULL.csv")
OUT_DIR = os.environ.get("R1_IPC_OUT") or os.path.join(WS, "2_monitoring", "reports", "round1_ipc_ch_representativity_2026-10-04")
DO_FILE = "IMPACT_NGA_Dataset_MSNA-UNHCR-2026-Round1_October-2026.xlsx"
DO_MD5 = "d603f7a2e4f6599a46c254bc1726c8a7"
DO_HH_IN_COVERAGE = 21428
MSNA_LIGHT = os.environ.get("R1_MSNA_LIGHT", "not_assessed")
FINAL = os.environ.get("R1_IPC_FINAL") == "1"
assert MSNA_LIGHT in ("not_assessed", "keep"), MSNA_LIGHT
# Full-precision stratum values (N_hh, N_acc, variance) captured by capture_round1_full_precision_2026-10-04.py from an
# unchanged re-run of the verified Round 1 build on the candidate snapshot (re-run tables byte-identical to the verified
# ones). The verified CSVs round these, and sums of rounded values drift by 1-2 households, so the exact values are used.
FP_CSV = os.path.join(WS, "2_monitoring", "reports", "round1_ipc_ch_representativity_2026-10-04", "_build_intermediate", "full_precision_strata.csv")
Z = 1.6448536269514722   # the verified build's Z (90% two-sided); asserted against the capture below
REPR, INDIC = "Representative", "Indicative - meets reporting threshold"
CAT_NAME = {"a": "Representative", "b": "Indicative", "c": "Dropped", "d": "Not assessed"}
CAT_DEF = {
    "a": "Margin of error 10% or less (90% confidence).",
    "b": "Margin of error above 10%, with at least 20 interviews. Usable with caution.",
    "c": "Fewer than 20 interviews, and a margin of error above 10%. No estimate is reported.",
    "d": "Not assessed in Round 1: outside the Round 1 coverage, excluded from the sampling design, no accessible "
         "population, or (MSNA Light) not part of the weighted Round 1 analysis. No estimate is reported.",
}
EXCL_TEXT = {
    "accessibility_loss_below_population_threshold": "Excluded from the design: accessible share of the population below the design threshold",
    "insecurity_related_inaccessibility": "Excluded from the design: inaccessible because of insecurity",
    "idp_population_no_longer_present_partner_reported": "Excluded from the design: IDP population no longer present (partner report)",
}
log_lines = []

def log(msg):
    print(msg)
    log_lines.append(msg)

def md5(p):
    h = hashlib.md5()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()

def rd(p):
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def var_h(n, N):
    """Variance of a proportion at p = 0.5, design effect 1, with finite population correction: 0.25 / n_eff,
    n_eff = n (N - 1) / (N - n). Achieved at or above the accessible population -> 0 (MoE floored at 0).
    Same expression as the verified build's var_simple()."""
    if N <= n:
        return 0.0
    return 0.5 * (1 - 0.5) / (n * (N - 1) / (N - n))

def moe_stratum(n, N):
    if n <= 0 or N <= 0:
        return None
    return 100 * Z * math.sqrt(var_h(n, N))

def aggregate(strata, included):
    """MoE = Z * sqrt(sum W_h^2 * Var_h) over included strata, W_h = share of accessible households."""
    inc = [s for s in strata if included(s)]
    hh_all = sum(s["N"] for s in strata)
    acc_all = sum(s["Nacc"] for s in strata)
    rep = sum(s["Nacc"] for s in inc)
    moe = None
    if inc:
        moe = 100 * Z * math.sqrt(sum((s["Nacc"] / rep) ** 2 * var_h(s["n"], s["Nacc"]) for s in inc))
    return {"inc": inc, "n_inc": sum(s["n"] for s in inc), "hh_all": hh_all, "acc_all": acc_all, "rep": rep, "moe": moe}

def fnum(x):
    return float(x) if str(x).strip() != "" else None

def moe_display(m):
    """MoE to 2 decimals, or to as many more as needed so that a value above 10% never displays as 10.00
    (categories use the unrounded value)."""
    if m is None:
        return None
    d = 2
    while m > 10 and round(m, d) <= 10 and d < 8:
        d += 1
    return round(m, d)

def edge_note(m):
    return ("Margin of error %s%% before rounding: just above the 10%% threshold, so (b)" % ("%.4f" % m)) if (m is not None and m > 10 and round(m, 2) <= 10) else ""

# ---------------------------------------------------------------------------------------------------------------------
# 1. Inputs
log("Round 1 IPC/CH workbook build - %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
for p in (STRATA_CSV, LGA_CSV, STATE_CSV, FRAME_CSV):
    log("input %s md5 %s" % (os.path.relpath(p, WS), md5(p)))
raw = rd(STRATA_CSV)
strata = []
for r in raw:
    strata.append({
        "row": r, "state": r["State"], "lga": r["LGA"], "pcode": r["adm2_pcode"], "pop": r["Pop type"], "sid": r["strata_id"],
        "method": r["Sampling method"], "covered": r["Covered in design"], "basis": r["Accessible-population basis"],
        "N": float(r["Households (N_hh)"]), "pct_acc": fnum(r["% population accessible"]), "Nacc": float(r["Accessible households"]),
        "n": int(float(r["Achieved (Round 1)"])), "moe_csv": fnum(r["MoE Round 1 (deff=1) %"]),
        "label": r["Round 1 label"], "reason": r["Round 1 reason"],
    })
log("strata: %d | interviews: %d" % (len(strata), sum(s["n"] for s in strata)))
assert sum(s["n"] for s in strata) == DO_HH_IN_COVERAGE, "interviews in coverage != 21,428"
log("input %s md5 %s" % (os.path.relpath(FP_CSV, WS), md5(FP_CSV)))
fp = {r["strata_id"]: r for r in rd(FP_CSV)}
assert all(float(r["Z"]) == Z for r in fp.values()), "captured Z differs"
for s in strata:
    f = fp[s["sid"]]
    N_full, Nacc_full = float(f["_N_hh"]), float(f["_N_acc"])
    # the full-precision values must round to exactly what the verified table shows
    assert round(N_full) == round(s["N"]) and round(Nacc_full, 1) == s["Nacc"], (s["sid"], N_full, s["N"], Nacc_full, s["Nacc"])
    assert int(f["Achieved (Round 1)"]) == s["n"] and f["Round 1 label"] == s["label"], s["sid"]
    s["N"], s["Nacc"] = N_full, Nacc_full
    s["var_captured"] = None if f["_var"] in ("", "None") else float(f["_var"])
log("full-precision values attached to all %d strata (Z = %r)" % (len(strata), Z))

# ---------------------------------------------------------------------------------------------------------------------
# 2. Self-check: recompute the verified Round 1 tables exactly (Round 1 labels, MSNA Light kept, as in the source)
fails = []
for s in strata:
    excluded = s["reason"].startswith("Dropped - LGA/strata excluded")
    m = None if excluded else moe_stratum(s["n"], s["Nacc"])
    s["moe"] = m
    if s["var_captured"] is not None and (s["n"] <= 0 or var_h(s["n"], s["Nacc"]) != s["var_captured"]):
        fails.append("stratum variance %s: recomputed %r vs captured %r" % (s["sid"], var_h(s["n"], s["Nacc"]) if s["n"] > 0 else None, s["var_captured"]))
    # the verified build's own label order: excluded -> no accessible population -> MoE <= 10 -> >= 20 interviews -> Dropped
    if excluded:
        lab, why = "Dropped", s["reason"]
    elif s["Nacc"] <= 0:
        lab, why = "Dropped", "Dropped - no accessible population"
    elif m is not None and m <= 10:
        lab, why = REPR, "Representative (MoE <= 10%)"
    elif s["n"] >= 20:
        lab, why = INDIC, "Indicative - meets reporting threshold (>=20 samples achieved)"
    else:
        lab, why = "Dropped", "Dropped - <20 samples achieved"
    if (lab, why) != (s["label"], s["reason"]):
        fails.append("stratum label %s: recomputed %s / %s vs %s / %s" % (s["sid"], lab, why, s["label"], s["reason"]))
    if (s["moe_csv"] is None) != (m is None) or (m is not None and round(m, 2) != s["moe_csv"]):
        fails.append("stratum MoE %s: recomputed %s vs %s" % (s["sid"], m, s["moe_csv"]))

def check_rows(ref_rows, groups, keyf, what):
    for r in ref_rows:
        k = keyf(r)
        if k not in groups:
            fails.append("%s row with no strata: %s" % (what, k)); continue
        a = aggregate(groups[k], lambda s: s["label"] != "Dropped")
        lab = "Dropped - no included strata" if a["moe"] is None else (REPR if a["moe"] <= 10 else INDIC)
        exp = {
            "Strata in unit": len(groups[k]), "Strata included": len(a["inc"]), "Achieved (included strata)": a["n_inc"],
            "Households (all strata)": round(a["hh_all"]), "Accessible households (all strata)": round(a["acc_all"]),
            "Households represented (included, accessible)": round(a["rep"]),
            "% of all households represented": round(100 * a["rep"] / a["hh_all"], 1) if a["hh_all"] else None,
            "% of accessible households represented": round(100 * a["rep"] / a["acc_all"], 1) if a["acc_all"] else None,
            "MoE Round 1 (deff=1) %": None if a["moe"] is None else round(a["moe"], 2), "Round 1 label": lab,
        }
        for c, v in exp.items():
            got = r[c]
            if c == "Round 1 label":
                ok = got == v
            elif v is None:
                ok = str(got).strip() == ""
            else:
                ok = str(got).strip() != "" and float(got) == float(v)
            if not ok:
                fails.append("%s %s, %s: recomputed %r vs %r" % (what, k, c, v, got))

by_lga = collections.defaultdict(list)
for s in strata:
    by_lga[(s["state"], s["lga"])].append(s)
check_rows(rd(LGA_CSV), by_lga, lambda r: (r["State"], r["LGA"]), "LGA")
scope_pop = {"IDP + Non-IDP": None, "IDP only": "IDP", "Non-IDP only": "Non-IDP"}
by_state = collections.defaultdict(list)
for s in strata:
    for sc, pop in scope_pop.items():
        if pop is None or s["pop"] == pop:
            by_state[(s["state"], sc)].append(s)
check_rows(rd(STATE_CSV), by_state, lambda r: (r["State"], r["Scope"]), "State")
if len(rd(LGA_CSV)) != len(by_lga):
    fails.append("LGA count %d vs %d" % (len(rd(LGA_CSV)), len(by_lga)))
lab_counts = collections.Counter(s["label"] for s in strata)
if (lab_counts[REPR], lab_counts[INDIC], lab_counts["Dropped"]) != (188, 48, 33):
    fails.append("strata label counts %s" % dict(lab_counts))
if fails:
    for f in fails[:60]:
        log("SELF-CHECK FAIL: " + f)
    log("STOPPED: %d self-check failures; nothing written." % len(fails))
    sys.exit(1)
log("SELF-CHECK PASS: %d strata, %d LGA rows, %d State rows reproduced exactly (MoE, labels, counts, households, shares)."
    % (len(strata), len(by_lga), len(rd(STATE_CSV))))

# ---------------------------------------------------------------------------------------------------------------------
# 3. Frame: coverage of every MSNA state, Kaduna totals, nothing missing from the Round 1 tables
frame = rd(FRAME_CSV)
msna_states = sorted({f["adm1_name"] for f in frame if f["coverage_status"] == "covered"})
r1_states = sorted({s["state"] for s in strata})
outside_states = [st for st in msna_states if st not in r1_states]
frame_only_states = sorted({f["adm1_name"] for f in frame} - set(msna_states))   # in the frame, no LGA in the design
REGION = {"NE": "North-East", "NW": "North-West", "NC": "North-Central"}
state_region = {f["adm1_name"]: REGION.get(f["region"], f["region"]) for f in frame}
log("MSNA states (any covered stratum in the frame): %s" % ", ".join(msna_states))
log("Round 1 states: %s | outside Round 1 (in design): %s | in the frame, no LGA in the design: %s" % (", ".join(r1_states), ", ".join(outside_states), ", ".join(frame_only_states)))
sids = {s["sid"] for s in strata}
for f in frame:
    if f["adm1_name"] in r1_states and f["adm1_name"] != "Kaduna" and f["strata_id"] not in sids:
        fails.append("frame stratum in a Round 1 state missing from the Round 1 table: %s (%s)" % (f["strata_id"], f["coverage_status"]))
kad = [f for f in frame if f["adm1_name"] == "Kaduna"]
kad_lgas = sorted({f["adm2_name"] for f in kad})
kad_cov = sorted({f["adm2_name"] for f in kad if f["coverage_status"] == "covered"})
kad_hh = sum(float(f["N_hh"]) for f in kad)
kad_cov_hh = sum(float(f["N_hh"]) for f in kad if f["adm2_name"] in kad_cov)
if kad_cov != sorted({s["lga"] for s in strata if s["state"] == "Kaduna"}):
    fails.append("Kaduna covered LGAs in frame %s != Round 1 table" % kad_cov)
if any(f["coverage_status"] == "covered" for f in kad if f["adm2_name"] not in kad_cov):
    fails.append("Kaduna: a covered stratum outside the covered LGAs")
log("Kaduna: %d LGAs, %s households in the design frame; in the sampling design: %s = %s households (%.1f%%)"
    % (len(kad_lgas), f"{kad_hh:,.0f}", ", ".join(kad_cov), f"{kad_cov_hh:,.0f}", 100 * kad_cov_hh / kad_hh))
if fails:
    for f in fails:
        log("FRAME CHECK FAIL: " + f)
    sys.exit(1)

# ---------------------------------------------------------------------------------------------------------------------
# 4. IPC/CH categories
def stratum_category(s):
    if s["label"] == REPR or s["label"] == INDIC:
        if s["method"] == "MSNA Light" and MSNA_LIGHT == "not_assessed":
            return "d", "MSNA Light: collected in the accessible area only; not part of the weighted Round 1 analysis"
        note = "" if s["method"] != "MSNA Light" else " (MSNA Light: not part of the weighted Round 1 analysis)"
        return ("a", "Margin of error 10% or less" + note) if s["label"] == REPR else ("b", "Margin of error above 10%, at least 20 interviews" + note)
    if s["reason"] == "Dropped - <20 samples achieved":
        return "c", "Fewer than 20 interviews (%d)" % s["n"]
    if s["reason"] == "Dropped - no accessible population":
        return "d", "No accessible population at the end of Round 1"
    if s["reason"].startswith("Dropped - LGA/strata excluded"):
        code = s["reason"].split("(", 1)[1].rstrip(")")
        return "d", EXCL_TEXT[code]
    raise ValueError("unmapped stratum %s: %s / %s" % (s["sid"], s["label"], s["reason"]))

for s in strata:
    s["cat"], s["why"] = stratum_category(s)
inc_ipc = lambda s: s["cat"] in ("a", "b")
pop_label = lambda ss: " + ".join(sorted({x["pop"] for x in ss}, key=lambda p: 0 if p == "IDP" else 1)) or "none"

def unit_category(a, ss):
    if a["moe"] is not None:
        return "a" if a["moe"] <= 10 else "b"
    return "c" if any(x["cat"] == "c" for x in ss) else "d"

def not_included(ss):
    return "; ".join("%s: %s" % (x["pop"], x["why"]) for x in ss if not inc_ipc(x))

lga_out = []
for (st, lg), ss in sorted(by_lga.items()):
    a = aggregate(ss, inc_ipc)
    lga_out.append({"State": st, "LGA": lg, "LGA pcode": ss[0]["pcode"], "Category": unit_category(a, ss),
                    "Population groups in the estimate": pop_label(a["inc"]) if a["inc"] else "none",
                    "Interviews in the estimate": a["n_inc"], "Margin of error (%)": moe_display(a["moe"]),
                    "Households (all)": round(a["hh_all"]), "Households in accessible areas": round(a["acc_all"]),
                    "Households represented by the estimate": round(a["rep"]),
                    "% of households represented": round(100 * a["rep"] / a["hh_all"], 1) if a["hh_all"] else None,
                    "Not included in the estimate": not_included(ss),
                    "Note": "; ".join(x for x in ["Only Jema'a and Zaria are in the MSNA design in Kaduna; no Kaduna state figure" if st == "Kaduna" else "", edge_note(a["moe"])] if x)})
state_out = []
for st in r1_states:
    if st == "Kaduna":
        continue
    n_lgas = len({s["lga"] for s in strata if s["state"] == st})
    for sc in scope_pop:
        ss = by_state[(st, sc)]
        a = aggregate(ss, inc_ipc)
        state_out.append({"State": st, "Population group": sc.replace(" only", ""), "Category": unit_category(a, ss),
                          "LGAs": n_lgas, "Strata": len(ss), "Strata in the estimate": len(a["inc"]),
                          "Interviews in the estimate": a["n_inc"], "Margin of error (%)": moe_display(a["moe"]),
                          "Households (all)": round(a["hh_all"]), "Households in accessible areas": round(a["acc_all"]),
                          "Households represented by the estimate": round(a["rep"]),
                          "% of households represented": round(100 * a["rep"] / a["hh_all"], 1) if a["hh_all"] else None,
                          "Not included in the estimate": "; ".join("%s %s: %s" % (x["lga"], x["pop"], x["why"]) for x in ss if not inc_ipc(x))})
strata_out = [{"State": s["state"], "LGA": s["lga"], "LGA pcode": s["pcode"], "Population group": s["pop"], "Category": s["cat"],
               "Interviews": s["n"], "Margin of error (%)": moe_display(s["moe"]) if inc_ipc(s) else None,
               "Households (all)": round(s["N"]), "% of households in accessible areas": s["pct_acc"],
               "Households in accessible areas": round(s["Nacc"], 1), "Accessible-population basis": s["basis"],
               "Reason": "; ".join(x for x in [s["why"], edge_note(s["moe"]) if inc_ipc(s) else ""] if x), "Stratum ID": s["sid"]} for s in sorted(strata, key=lambda s: (s["state"], s["lga"], s["pop"]))]

# Not assessed: every LGA of an MSNA state outside Round 1, and Kaduna's LGAs outside the design (from the frame)
na = collections.OrderedDict()
for f in sorted(frame, key=lambda f: (f["adm1_name"], f["adm2_name"])):
    st, lg = f["adm1_name"], f["adm2_name"]
    if st in outside_states or st in frame_only_states or (st == "Kaduna" and lg not in kad_cov):
        d = na.setdefault((st, lg), {"State": st, "LGA": lg, "LGA pcode": f["adm2_pcode"], "Region": f["region"], "statuses": set(), "pops": set(), "hh": 0.0})
        d["statuses"].add(f["coverage_status"]); d["pops"].add("IDP" if f["pop_type"] == "idp" else "Non-IDP"); d["hh"] += float(f["N_hh"])
na_out = []
for (st, lg), d in na.items():
    in_design = bool(d["statuses"] & {"covered", "excluded"})
    if st == "Kaduna":
        why = "Not in the MSNA sampling design (Kaduna: only Jema'a and Zaria are in the design)"
    elif not in_design:
        why = "Not in the MSNA sampling design"
    else:
        why = "In the MSNA design; %s (%s) is not part of the Round 1 analysis coverage" % (st, state_region[st])
    na_out.append({"State": st, "LGA": lg, "LGA pcode": d["LGA pcode"], "Category": "d", "Population groups in the frame": " + ".join(sorted(d["pops"])),
                   "Households (design frame)": round(d["hh"]), "Reason": why, "_hh": d["hh"]})

# ---------------------------------------------------------------------------------------------------------------------
# 5. Consistency checks on the IPC/CH outputs
cc = collections.Counter(s["cat"] for s in strata)
lc = collections.Counter(r["Category"] for r in lga_out)
n_est = sum(s["n"] for s in strata if inc_ipc(s))
r2 = lambda v: None if v is None else round(v, 2)   # compare displayed values with the 2-decimal verified tables
assert sum(cc.values()) == len(strata) == 269 and sum(lc.values()) == len(lga_out) == 138
assert all(r["State"] != "Kaduna" for r in state_out) and len(state_out) == 3 * (len(r1_states) - 1)
assert all((r["Margin of error (%)"] is None) == (r["Category"] in ("c", "d")) for r in lga_out + state_out)
odd = [r for r in lga_out + state_out + strata_out if r["Margin of error (%)"] is not None and (r["Margin of error (%)"] <= 10) != (r["Category"] == "a")]
for r in odd:
    log("DISPLAY/CATEGORY MISMATCH: %s" % {k: r[k] for k in r if k in ("State", "LGA", "Population group", "Category", "Margin of error (%)", "Reason", "Stratum ID")})
assert not odd, "a displayed MoE disagrees with its category"
for r in strata_out:
    if r["Category"] in ("a", "b"):
        assert r["Interviews"] >= 20
if MSNA_LIGHT == "keep":
    for r in lga_out:   # in keep mode the LGA table must equal the verified one
        ref = next(x for x in rd(LGA_CSV) if (x["State"], x["LGA"]) == (r["State"], r["LGA"]))
        assert r2(r["Margin of error (%)"]) == fnum(ref["MoE Round 1 (deff=1) %"]), r
log("IPC/CH categories (MSNA Light = %s): strata %s | LGAs %s | interviews in estimates %d of %d"
    % (MSNA_LIGHT, dict(sorted(cc.items())), dict(sorted(lc.items())), n_est, DO_HH_IN_COVERAGE))
changed = [(s["state"], s["lga"], s["pop"]) for s in strata if s["label"] in (REPR, INDIC) and not inc_ipc(s)]
log("strata moved out of the estimates vs the Round 1 tables: %s" % changed)
for r in lga_out:
    ref = next(x for x in rd(LGA_CSV) if (x["State"], x["LGA"]) == (r["State"], r["LGA"]))
    if r2(r["Margin of error (%)"]) != fnum(ref["MoE Round 1 (deff=1) %"]):
        log("   LGA changed vs Round 1 table: %s / %s MoE %s -> %s, category %s" % (r["State"], r["LGA"], ref["MoE Round 1 (deff=1) %"], r["Margin of error (%)"], r["Category"]))
for r in state_out:
    ref = next(x for x in rd(STATE_CSV) if (x["State"], x["Scope"].replace(" only", "")) == (r["State"], r["Population group"]))
    if r2(r["Margin of error (%)"]) != fnum(ref["MoE Round 1 (deff=1) %"]):
        log("   State changed vs Round 1 table: %s / %s MoE %s -> %s" % (r["State"], r["Population group"], ref["MoE Round 1 (deff=1) %"], r["Margin of error (%)"]))

# ---------------------------------------------------------------------------------------------------------------------
# 6. Workbook
FILL = {"a": "C6EFCE", "b": "FFEB9C", "c": "FFC7CE", "d": "E7E6E6"}
HDR = PatternFill("solid", fgColor="1F4E78")
wb = openpyxl.Workbook()

def add_table(ws, rows, cols, widths, start_row=1, cat_col="Category"):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 14)
    for i, r in enumerate(rows, start_row + 1):
        for j, c in enumerate(cols, 1):
            v = r.get(c)
            if c == "Category name":
                v = CAT_NAME[r["Category"]]
            cell = ws.cell(row=i, column=j, value=v)
            if c in ("Category", "Category name"):
                cell.fill = PatternFill("solid", fgColor=FILL[r["Category"]])
            if isinstance(v, int):
                cell.number_format = "#,##0"
            elif isinstance(v, float):
                cell.number_format = ("0." + "0" * next(d for d in range(2, 9) if round(v, d) == v)) if c.startswith("Margin") else ("0.0" if c.startswith("%") else "#,##0.0")
            if c in ("Reason", "Not included in the estimate", "Note"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = "A%d:%s%d" % (start_row, get_column_letter(len(cols)), start_row + len(rows))

W = {"State": 11, "LGA": 16, "LGA pcode": 11, "Category": 9, "Category name": 15, "Population group": 13, "Population groups in the estimate": 16,
     "Interviews": 11, "Interviews in the estimate": 12, "Margin of error (%)": 10, "Households (all)": 13, "Households in accessible areas": 14,
     "Households represented by the estimate": 15, "% of households represented": 12, "% of households in accessible areas": 12,
     "Accessible-population basis": 30, "Reason": 60, "Stratum ID": 18, "Not included in the estimate": 60, "Note": 40,
     "LGAs": 7, "Strata": 7, "Strata in the estimate": 10, "Population groups in the frame": 16, "Households (design frame)": 14}

# Read me
def names(states):
    return states[0] if len(states) == 1 else ", ".join(states[:-1]) + " and " + states[-1]

def by_region(states):
    groups = collections.OrderedDict()
    for st in sorted(states, key=lambda s: (state_region[s], s)):
        groups.setdefault(state_region[st], []).append(st)
    return ", ".join("%s (%s)" % (names(v), k) for k, v in groups.items())

rm = wb.active; rm.title = "Read me"
rm.column_dimensions["A"].width = 24; rm.column_dimensions["B"].width = 110
title = "MSNA N-WEC 2026 - Round 1 representativity summary (states, LGAs, population groups)"
lines = [
    (title, None, "title"),
    ("Status", "DRAFT for review - not yet for circulation" if not FINAL else "Final", None),
    ("Purpose", "Final summary of Round 1 representativity: for each state, LGA and population group (IDP / Non-IDP), whether "
                "the Round 1 data supports an estimate, and with what margin of error.", None),
    ("Data basis", "The final cleaned Round 1 dataset: %s (md5 %s). %s households in the Round 1 coverage below "
                   "(interviews submitted by 30 September 2026 and kept by the final data cleaning)." % (DO_FILE, DO_MD5, f"{DO_HH_IN_COVERAGE:,}"), None),
    ("Round 1 coverage", "North-East: %s. North-West: %s, and the two Kaduna LGAs in the MSNA design (Jema'a and Zaria). "
                         "Not part of the Round 1 analysis, and listed on the 'Not assessed' sheet: %s, which are in the MSNA "
                         "design; and %s, which appear in the sampling frame but have no LGA in the MSNA sampling design."
                         % (names([s for s in r1_states if state_region[s] == "North-East"]),
                            names([s for s in r1_states if state_region[s] == "North-West" and s != "Kaduna"]),
                            by_region(outside_states), by_region(frame_only_states)), None),
    ("Kaduna", "Only Jema'a and Zaria are in the MSNA sampling design: 2 of Kaduna's %d LGAs, %s of %s households in the "
               "design frame (%.1f%%). Kaduna is therefore reported for these two LGAs only, with no state-level figure."
               % (len(kad_lgas), f"{kad_cov_hh:,.0f}", f"{kad_hh:,.0f}", 100 * kad_cov_hh / kad_hh), None),
    ("Categories", None, "header"),
]
for k in "abcd":
    lines.append(("(%s) %s" % (k, CAT_NAME[k]), CAT_DEF[k], "cat:" + k))
lines += [
    ("Margin of error", "MoE = 1.6449 x sqrt(0.25 / n_eff) x 100, with n_eff = n (N - 1) / (N - n): n = interviews, N = households "
                        "in the accessible areas of the stratum. 90% confidence, p = 0.5 (the most conservative proportion), "
                        "design effect 1, with finite population correction - the simplified rule agreed for Round 1. No margin of error "
                        "is shown for (c) and (d), as no estimate is reported for them.", None),
    ("Strata", "A stratum is one population group (IDP or Non-IDP) in one LGA. Categories are assigned per stratum first.", None),
    ("LGA and state figures", "Computed over the strata in categories (a) and (b) only, each weighted by its share of accessible "
                              "households: MoE = 1.6449 x sqrt(sum of W^2 x Var) x 100. Strata in (c) and (d) are left out of every "
                              "LGA and state figure. An LGA or state is (a) if its MoE is 10% or less and (b) otherwise; an LGA "
                              "with no stratum in (a) or (b) is (c) if any of its strata is (c), otherwise (d).", None),
    ("Accessible households", "N is the households of the stratum living in areas reported accessible at the end of Round 1 "
                              "(design households x the accessible share of the population). For %d strata whose accessible area "
                              "shrank after their interviews were collected, the larger accessible share recorded during collection "
                              "(8 or 26 September) is used; for %d stratum never recorded as accessible, all its households are used "
                              "(both choices give a larger N, so a larger, more conservative MoE)."
                              % (sum(s["basis"].startswith("accessible population during collection") for s in strata),
                                 sum(s["basis"].startswith("total households") for s in strata)), None),
    ("% of households represented", "The share of all households in the unit (including inaccessible areas) that live in the "
                                    "accessible areas of the strata in the estimate. Inaccessible areas and strata in (c)/(d) are "
                                    "not represented. Check 'Population groups in the estimate': some LGA figures cover only "
                                    "IDPs or only Non-IDPs.", None),
    ("MSNA Light", ("Abadam, Guzamala and Nganzai (Borno, Non-IDP) were collected under MSNA Light, in the accessible area only. "
                    "They are not part of the weighted Round 1 analysis, so they are shown as (d) and left out of the LGA and Borno "
                    "figures. On the margin-of-error rule alone they would be (a).") if MSNA_LIGHT == "not_assessed" else
                   ("Abadam, Guzamala and Nganzai (Borno, Non-IDP) were collected under MSNA Light, in the accessible area only. "
                    "They keep their margin-of-error category here, but they are not part of the weighted Round 1 analysis."), None),
    ("Sheets", "Summary - counts by category, per state. State - each state, all households and by population group. "
               "LGA - every Round 1 LGA. Strata - every LGA x population group. Not assessed - areas outside Round 1.", None),
    ("Prepared", datetime.date.today().strftime("%d %B %Y"), None),
]
r0 = 1
for k, v, kind in lines:
    a = rm.cell(row=r0, column=1, value=k); b = rm.cell(row=r0, column=2, value=v)
    a.alignment = Alignment(wrap_text=True, vertical="top"); b.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title":
        a.font = Font(bold=True, size=14); rm.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=2)
    elif kind == "header":
        a.font = Font(bold=True, size=12)
    else:
        a.font = Font(bold=True)
        if kind and kind.startswith("cat:"):
            a.fill = PatternFill("solid", fgColor=FILL[kind[4:]])
    r0 += 1

# Summary
sm = wb.create_sheet("Summary")
rows = []
for st in r1_states:
    ss = [s for s in strata if s["state"] == st]
    ll = [r for r in lga_out if r["State"] == st]
    so = next((r for r in state_out if r["State"] == st and r["Population group"] == "IDP + Non-IDP"), None)
    row = {"State": st, "LGAs": len(ll), "Interviews": sum(s["n"] for s in ss), "Interviews in the estimates": sum(s["n"] for s in ss if inc_ipc(s)),
           "State margin of error (%)": so["Margin of error (%)"] if so else None,
           "% of households represented": so["% of households represented"] if so else None,
           "Note": "No state figure (2 of 23 LGAs in the design)" if st == "Kaduna" else ""}
    for k in "abcd":
        row["LGAs (%s)" % k] = sum(r["Category"] == k for r in ll)
        row["Strata (%s)" % k] = sum(s["cat"] == k for s in ss)
    rows.append(row)
tot = {"State": "Total (Round 1 coverage)", "LGAs": len(lga_out), "Interviews": sum(s["n"] for s in strata), "Interviews in the estimates": n_est, "Note": ""}
for k in "abcd":
    tot["LGAs (%s)" % k] = lc[k]; tot["Strata (%s)" % k] = cc[k]
rows.append(tot)
scols = ["State", "LGAs", "LGAs (a)", "LGAs (b)", "LGAs (c)", "LGAs (d)", "Strata (a)", "Strata (b)", "Strata (c)", "Strata (d)",
         "Interviews", "Interviews in the estimates", "State margin of error (%)", "% of households represented", "Note"]
sm.cell(row=1, column=1, value="Round 1 representativity by state - counts by category (definitions on 'Read me')").font = Font(bold=True, size=12)
for j, c in enumerate(scols, 1):
    cell = sm.cell(row=3, column=j, value=c); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sm.column_dimensions[get_column_letter(j)].width = {"State": 24, "Note": 40}.get(c, 11)
    k = c[-2] if c.endswith(")") and c[-3] == "(" else None
    if k in FILL:
        cell.fill = PatternFill("solid", fgColor=FILL[k]); cell.font = Font(bold=True)
for i, r in enumerate(rows, 4):
    for j, c in enumerate(scols, 1):
        v = r.get(c)
        cell = sm.cell(row=i, column=j, value=v)
        if isinstance(v, int): cell.number_format = "#,##0"
        elif isinstance(v, float): cell.number_format = "0.00" if "margin" in c else "0.0"
        if r["State"].startswith("Total"): cell.font = Font(bold=True)
r_na = len(rows) + 6
sm.cell(row=r_na, column=1, value="Not part of Round 1 (category d) - see 'Not assessed'").font = Font(bold=True)
for i, st in enumerate(outside_states + frame_only_states + ["Kaduna (other LGAs)"], r_na + 1):
    nn = [r for r in na_out if r["State"] == st.split(" ")[0]]
    sm.cell(row=i, column=1, value=st)
    sm.cell(row=i, column=2, value="%d LGAs (%d in the MSNA design), %s households in the design frame"
            % (len(nn), sum(r["Reason"].startswith("In the MSNA design") for r in nn), f"{round(sum(r['_hh'] for r in nn)):,}"))

# State, LGA, Strata, Not assessed
ws = wb.create_sheet("State")
add_table(ws, state_out, ["State", "Population group", "Category", "Category name", "LGAs", "Strata", "Strata in the estimate",
                          "Interviews in the estimate", "Margin of error (%)", "Households (all)", "Households in accessible areas",
                          "Households represented by the estimate", "% of households represented", "Not included in the estimate"], W)
ws.cell(row=len(state_out) + 3, column=1, value="Kaduna has no state row: only Jema'a and Zaria (2 of 23 LGAs) are in the MSNA design - see the LGA sheet.").font = Font(italic=True)
ws = wb.create_sheet("LGA")
add_table(ws, lga_out, ["State", "LGA", "LGA pcode", "Category", "Category name", "Population groups in the estimate", "Interviews in the estimate",
                        "Margin of error (%)", "Households (all)", "Households in accessible areas", "Households represented by the estimate",
                        "% of households represented", "Not included in the estimate", "Note"], W)
ws = wb.create_sheet("Strata")
add_table(ws, strata_out, ["State", "LGA", "LGA pcode", "Population group", "Category", "Category name", "Interviews", "Margin of error (%)",
                           "Households (all)", "% of households in accessible areas", "Households in accessible areas",
                           "Accessible-population basis", "Reason", "Stratum ID"], W)
ws = wb.create_sheet("Not assessed")
add_table(ws, na_out, ["State", "LGA", "LGA pcode", "Category", "Category name", "Population groups in the frame", "Households (design frame)", "Reason"], W)

os.makedirs(OUT_DIR, exist_ok=True)
name = "MSNA_N-WEC_2026_Round1_representativity_summary_2026-10-04%s.xlsx" % ("" if FINAL else "_DRAFT")
out = os.path.join(OUT_DIR, name)
wb.save(out)
vdir = os.path.join(OUT_DIR, "_csv_for_verification")
os.makedirs(vdir, exist_ok=True)
for fname, rows_, cols_ in [("strata.csv", strata_out, None), ("lga.csv", lga_out, None), ("state.csv", state_out, None), ("not_assessed.csv", [{k: v for k, v in r.items() if not k.startswith("_")} for r in na_out], None)]:
    with open(os.path.join(vdir, fname), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows_[0].keys()))
        w.writeheader(); w.writerows(rows_)
log("wrote %s" % os.path.relpath(out, WS))
with open(os.path.join(OUT_DIR, "BUILD_LOG.txt"), "w", encoding="utf-8", newline="\n") as f:
    f.write("\n".join(log_lines) + "\n")
