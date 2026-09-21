# ==============================================================================
# 2026-09-21 late evening. Every file in accessibility_reports_returned/ was
# found to carry a "[trash]/*.dat" + "customXml/*" set (13 entries), almost
# certainly SharePoint/OneDrive document-library metadata plus Excel's own
# auto-repair residue - not caused by any script here (checked: identical
# pattern on 11 files this session never touched). Verified before writing
# this:
#   - customXml/itemN.xml IS legitimately declared (Content_Types.xml
#     Override + workbook.xml.rels Relationship) - real SharePoint custom-
#     properties metadata (docProps/custom.xml's ContentTypeId/
#     MediaServiceImageTags are the same SharePoint infrastructure), not a
#     dangling reference.
#   - [trash]/*.dat is NOT referenced anywhere in the package (grepped every
#     other part) and its bytes are inert 0xFF+zero padding, not real data.
#   - Every real worksheet/data part is byte-identical to a clean baseline
#     (checked against git HEAD for the 7 files this session had touched).
# Removing both is zero-risk to content; done anyway since an unreferenced,
# non-standard "[trash]" part risks an Excel repair prompt for partners.
#
# Strips [trash]/* and customXml/* entries, and the 3 matching declarations
# each in [Content_Types].xml and xl/_rels/workbook.xml.rels (regex, exact
# literal match only - refuses if the expected pattern isn't found rather
# than guess). Everything else byte-for-byte unchanged, same compression.
# Backs up every touched file to accessibility_reports_returned/_archive/
# first, per the standing convention.
# Usage: python strip_trash_customxml_2026-09-21.py
# ==============================================================================
import os
import re
import shutil
import zipfile
from datetime import date

DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling\resampling\input\accessibility_reports_returned"
ARCHIVE = os.path.join(DIR, "_archive")
TODAY = str(date.today())

CT_ITEM_RE = re.compile(r'<Override PartName="/customXml/itemProps\d\.xml"[^>]*/>')
RELS_ITEM_RE = re.compile(r'<Relationship [^>]*Type="[^"]*/customXml"[^>]*/>')


def clean_one(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    junk = [n for n in names if n.startswith("[trash]") or n.startswith("customXml")]
    if not junk:
        return None
    ct = z.read("[Content_Types].xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    ct_new, n_ct = CT_ITEM_RE.subn("", ct)
    rels_new, n_rels = RELS_ITEM_RE.subn("", rels)
    n_custom_items = len([n for n in junk if re.match(r"^customXml/itemProps\d\.xml$", n)])
    if n_ct != n_custom_items or n_rels != n_custom_items:
        raise SystemExit(f"{path}: expected {n_custom_items} customXml declarations in each of "
                         f"Content_Types.xml ({n_ct} found) and workbook.xml.rels ({n_rels} found) - refusing, pattern mismatch.")

    tmp = path + ".cleaning.tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item in z.infolist():
            if item.filename in junk:
                continue
            data = z.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = ct_new.encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = rels_new.encode("utf-8")
            out.writestr(item, data)
    z.close()
    return tmp, len(junk)


def main():
    os.makedirs(ARCHIVE, exist_ok=True)
    cleaned = 0
    for fn in sorted(os.listdir(DIR)):
        if not fn.endswith("_accessibility_report.xlsx"):
            continue
        path = os.path.join(DIR, fn)
        result = clean_one(path)
        if result is None:
            print(f"  {fn}: already clean")
            continue
        tmp, n_junk = result
        bak = os.path.join(ARCHIVE, f"{fn[:-5]}_pre_trash_strip_{TODAY}.xlsx")
        shutil.copy2(path, bak)
        os.replace(tmp, path)
        # verify: opens, same visible row counts as before cleaning
        import openpyxl
        wb_before = openpyxl.load_workbook(bak, read_only=True)
        wb_after = openpyxl.load_workbook(path, read_only=True)
        rows_before = {ws: sum(1 for r in wb_before[ws].iter_rows(min_row=2, values_only=True) if r and any(v is not None for v in r)) for ws in wb_before.sheetnames}
        rows_after = {ws: sum(1 for r in wb_after[ws].iter_rows(min_row=2, values_only=True) if r and any(v is not None for v in r)) for ws in wb_after.sheetnames}
        if rows_before != rows_after:
            os.replace(bak, path)
            raise SystemExit(f"{fn}: row counts changed after cleaning ({rows_before} -> {rows_after}) - reverted, needs investigation.")
        print(f"  {fn}: stripped {n_junk} junk entries, verified {sum(rows_after.values())} rows unchanged across {len(rows_after)} sheets")
        cleaned += 1
    print(f"\n{cleaned} file(s) cleaned.")


if __name__ == "__main__":
    main()
