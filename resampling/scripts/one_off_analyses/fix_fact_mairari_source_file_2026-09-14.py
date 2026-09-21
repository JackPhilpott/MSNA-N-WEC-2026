"""
Fixes the ROOT CAUSE of a silent revert found 2026-09-14: FACT's actual
returned Ward Accessibility file still says "No" (dated 2026-04-09, the
stale "AOG attacks and no vehicle or access roads" report) for Guzamala/
Mairari - even though this project's log already carries a Jack-confirmed
override (request_id 6697, 2026-09-14) making it Accessible, for the
government-negotiated MSNA Light arrangement.

Because that override was logged under partner="FACT" (same key FACT's own
real reports use), 02_ingest_accessibility_reports.py's own job - "does the
returned file disagree with what's currently logged?" - saw FACT's
unedited file (still "No") disagree with the override (Yes) on some later
re-ingest today, and correctly (by its own logic) generated a NEW log row
(request_id 6732, partner=FACT, accessible=No, same stale AOG-attacks
reason) that silently re-superseded the override. Caught via the 2026-09-07
incident's own standing invariant tripping (assert_plausible("WORKING rows
in a currently-Inaccessible ward") = 156, all Guzamala MSNA Light rows).

The durable fix is to edit the SOURCE FILE itself, not just re-add another
log override - a log-only override will keep getting silently reverted by
every future re-ingest of FACT's unedited file, for as long as the file and
the log disagree. Editing the file makes them agree, so any future re-
ingest correctly finds "0 new/changed" instead of re-reverting.

Backs up the file first (matches this project's established convention).
"""
import shutil
from datetime import date
import openpyxl

PATH = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling\resampling\input\accessibility_reports_returned\FACT_accessibility_report.xlsx"
BACKUP = PATH.replace(".xlsx", "_PRE_mairari_source_fix_2026-09-14.xlsx.bak")

shutil.copy2(PATH, BACKUP)
print(f"Backed up to {BACKUP}")

wb = openpyxl.load_workbook(PATH)
ws = wb["Ward Accessibility"]
header = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(header)}

fixed = 0
for row_idx in range(2, ws.max_row + 1):
    lga = ws.cell(row=row_idx, column=idx["LGA"] + 1).value
    ward = ws.cell(row=row_idx, column=idx["Ward (GRID3)"] + 1).value
    if lga == "Guzamala" and ward == "Mairari":
        ws.cell(row=row_idx, column=idx["Accessible (Y/N)"] + 1).value = "Yes"
        ws.cell(row=row_idx, column=idx["Reason category"] + 1).value = "N/A - fully accessible"
        ws.cell(row=row_idx, column=idx["Reason notes"] + 1).value = (
            "Government-negotiated MSNA Light data collection arrangement for Guzamala (Mairari "
            "settlement) - confirmed directly by Jack Philpott, 2026-09-14. Supersedes the original "
            "2026-04-09 AOG-attacks/no-access-roads report, which predates and is unrelated to this "
            "negotiated access. See CLAUDE.md Update 2026-09-11 (MSNA Light) for full context."
        )
        ws.cell(row=row_idx, column=idx["Date reported"] + 1).value = date(2026, 9, 14)
        fixed += 1

wb.save(PATH)
print(f"Fixed {fixed} row(s) in FACT's returned file (Guzamala/Mairari -> Yes, N/A - fully accessible).")
