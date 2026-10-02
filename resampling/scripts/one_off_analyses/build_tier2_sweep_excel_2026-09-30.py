# ==============================================================================
# Rebuilds Tier2_exhaustion_full_sweep_2026-09-29.xlsx FRESH from the current,
# verified full_sweep_COMBINED_strategy_2026-09-30.csv - the old workbook
# (same filename) was the Phase-1-era two-separate-scenarios format Jack
# explicitly rejected as "genuinely ambiguous" when he asked for the combined
# single-strategy model; rebuilding fresh is cleaner than shoehorning new
# columns into a superseded structure (his own call to make either way).
#
# Adds one new derived column Jack asked for directly: increase over original
# target (both absolute and %), so he can filter/sort for the significant-
# increase cases himself without doing the arithmetic by hand.
# ==============================================================================
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling\resampling\output\full_sweep_2026-09-29"
CSV_PATH = OUT_DIR + r"\full_sweep_COMBINED_strategy_2026-09-30.csv"
XLSX_PATH = OUT_DIR + r"\Tier2_exhaustion_full_sweep_2026-09-29.xlsx"

with open(CSV_PATH, encoding="utf-8") as f:
    rows = list(csv.DictReader(f))

# Column order/labels for the sheet - reuses the CSV's own field names as the
# canonical source, relabelled for a partner/team-facing read.
COLUMNS = [
    ("state", "State"),
    ("lga", "LGA"),
    ("pop_type", "Pop type"),
    ("strata_id", "Strata ID"),
    ("partner", "Partner"),
    ("reaches_10pct", "Reaches 10% MoE? / status"),
    ("execution_status", "Execution status (30 Sep)"),
    ("current_moe_real_pct", "Current MoE % (real achieved only)"),
    ("projected_moe_final_pct", "Projected MoE % (combined strategy, final)"),
    ("original_target_samples", "Original target (design)"),
    ("current_revised_target_samples", "Current revised target (accessibility-adjusted)"),
    ("achieved_samples_so_far", "Achieved (real, field-collected)"),
    ("new_target_samples", "New target (samples)"),
    ("increase_over_original", "Increase over original target (samples)"),
    ("increase_over_original_pct", "Increase over original target (%)"),
    ("total_additional_real_interviews_combined", "Additional real interviews needed (draw + depth combined)"),
    ("real_new_clusters_deliverable", "Real new clusters deliverable (dry-run verified)"),
    ("existing_clusters_depth_additional_needed", "Existing-cluster depth top-up needed"),
    ("available_population_N_hh_accessible", "Available population (N_hh_accessible)"),
    ("reason_if_excluded", "Reason if excluded"),
    ("existing_clusters_depth_detail", "Depth top-up detail (per cluster)"),
    ("real_new_cluster_sizes", "Real new cluster sizes (if any)"),
]

PCT_COLS = {"current_moe_real_pct", "projected_moe_final_pct", "moe_after_real_draw_only_pct", "increase_over_original_pct"}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Combined sweep (30 Sep)"

header_fill = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap = Alignment(wrap_text=True, vertical="top")

for col_i, (key, label) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=col_i, value=label)
    c.fill = header_fill
    c.font = header_font
    c.alignment = wrap

significant_fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")  # >=50% increase, flagged for easy visual scan
SIGNIFICANT_PCT_THRESHOLD = 50.0

for r_i, r in enumerate(rows, start=2):
    orig = r.get("original_target_samples", "")
    new_t = r.get("new_target_samples", "")
    try:
        orig_n = float(orig)
        new_n = float(new_t)
        increase = new_n - orig_n
        increase_pct = (increase / orig_n * 100) if orig_n > 0 else None
    except (TypeError, ValueError):
        increase = "N/A"
        increase_pct = None
    r["increase_over_original"] = increase
    r["increase_over_original_pct"] = round(increase_pct, 1) if increase_pct is not None else "N/A"

    for col_i, (key, label) in enumerate(COLUMNS, start=1):
        val = r.get(key, "")
        if key in PCT_COLS and val not in ("", "N/A", None):
            try:
                val = float(val)
            except (TypeError, ValueError):
                pass
        cell = ws.cell(row=r_i, column=col_i, value=val)
        cell.alignment = Alignment(vertical="top", wrap_text=(key in ("existing_clusters_depth_detail", "execution_status", "reason_if_excluded")))
        if key in ("current_moe_real_pct", "projected_moe_final_pct") and isinstance(val, float):
            cell.number_format = "0.00"
    if isinstance(increase_pct, (int, float)) and increase_pct >= SIGNIFICANT_PCT_THRESHOLD:
        for col_i in range(1, len(COLUMNS) + 1):
            ws.cell(row=r_i, column=col_i).fill = significant_fill

widths = {
    "state": 10, "lga": 16, "pop_type": 9, "strata_id": 16, "partner": 12,
    "reaches_10pct": 30, "execution_status": 45, "current_moe_real_pct": 12,
    "projected_moe_final_pct": 14, "original_target_samples": 12, "current_revised_target_samples": 14,
    "achieved_samples_so_far": 11, "new_target_samples": 11, "increase_over_original": 12,
    "increase_over_original_pct": 12, "total_additional_real_interviews_combined": 14,
    "real_new_clusters_deliverable": 12, "existing_clusters_depth_additional_needed": 12,
    "available_population_N_hh_accessible": 14, "reason_if_excluded": 35,
    "existing_clusters_depth_detail": 45, "real_new_cluster_sizes": 20,
}
for col_i, (key, label) in enumerate(COLUMNS, start=1):
    ws.column_dimensions[get_column_letter(col_i)].width = widths.get(key, 14)

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows)+1}"

# Small README sheet for context, matching this project's usual workbook convention.
readme = wb.create_sheet("README", 0)
readme_lines = [
    "Tier 2 exhaustion - full national sweep, combined strategy",
    "",
    "Rebuilt 2026-09-30 from full_sweep_COMBINED_strategy_2026-09-30.csv (the current, verified source).",
    "",
    "Scope: every stratum currently assumed non-Representative, with a real accessible population to draw on.",
    "For each: real Tier1+Tier2 draw dry-run outcome, then existing-cluster depth top-up on top (design-weighted,",
    "model B) - ONE combined answer per stratum, not separate draw-only/depth-only scenarios.",
    "",
    "'Increase over original target' (2 new columns, added on request) = new_target_samples - original_target_samples,",
    "and the same as a %. Rows with a >=50% increase over the original design target are highlighted for a quick scan -",
    "these are the cases most worth an operational/resourcing conversation with partners, not just a data question.",
    "",
    "'Execution status' shows what's actually been done as of 2026-09-30: 4 IDP strata (Charanchi, Obi, Wamako,",
    "Shinkafi) were executed for real that night - real interviews still need collecting, but the frame capacity/",
    "target is live. Binji IDP was reverted to normal per a real formula fix (not a manual override) - see project memory.",
    "The 7 Non-IDP strata (Gwoza, Ngala, Shani, Charanchi NI, Bunza, Jega, Kalgo) were approved but blocked that night",
    "on a real mechanism gap (no tested Non-IDP existing-cluster top-up path existed) - in progress separately.",
    "",
    "'No available population' rows are accessibility-capped, NOT population-capped - the population exists but the",
    "assigned clusters sit in currently-inaccessible wards. Binji is the one genuine population-estimate edge case,",
    "now resolved (see 'Execution status' for that row).",
]
for i, line in enumerate(readme_lines, start=1):
    c = readme.cell(row=i, column=1, value=line)
    if i == 1:
        c.font = Font(bold=True, size=14)
readme.column_dimensions["A"].width = 130

wb.save(XLSX_PATH)
print(f"Wrote {XLSX_PATH}")
print(f"{len(rows)} rows, {len(COLUMNS)} columns.")
sig = sum(1 for r in rows if isinstance(r.get('increase_over_original_pct'), (int,float)) and r['increase_over_original_pct'] >= SIGNIFICANT_PCT_THRESHOLD)
print(f"{sig} rows flagged as >=50% increase over original target.")
