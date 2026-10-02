"""Round 1 representativity PROTOTYPE - 2 Oct 2026 (Coordinator, overnight, Jack-approved scope).

Applies Jack's 1 Oct dual-round decisions (memory: project_dual_round_representativity_design_2026-10-01)
to the FROZEN Round 1 snapshot (28,046 submissions, pre-recovery):
  - Round 1 gate = the same MoE formula with deff = 1 (no cluster design effect), uniform for every
    stratum, no certainty-PSU treatment (B4, B5). Z = 1.6449, p = 0.5, FPC as effective n, floored at 0.
  - Representative = MoE <= 10%; Indicative - meets reporting threshold = >= 20 achieved;
    otherwise Dropped (B6). Excluded/inaccessible strata are Dropped regardless of formula (B7).
  - Aggregates (D10-D13, Q2, Q3): Strata -> LGA -> State, each with its OWN MoE:
        MoE = Z * sqrt( sum_h W_h^2 * Var_h ),  W_h = accessible households share,
        Var_h = p(1-p) / ndeff_h  (deff = 1).
    ALL Dropped strata are left out of every aggregate ("as if they're not in the design at all");
    each aggregate reports the share of its population it actually speaks for.
    Combined IDP + Non-IDP and pop-type-split views (ToR assumption: WorldPop Non-IDP excludes IDPs).

Reuses 05_build_accessibility_impact_workbook.py's own loaders and formulas by importing it
(its main() is NOT run, so it writes nothing). Only the two submission inputs are re-pointed at the
frozen snapshot. Writes ONLY to resampling/output/round1_representativity_prototype_2026-10-02/.

The strict column is for reference: 05's own real-interview MoE (certainty-aware, else Kish
unequal-cluster), i.e. the "strict real-today" basis reported to Jack on 30 Sep.
"""
import csv
import importlib.util
import math
import os
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026"
SAMPLING_DIR = PROJECT_DIR + r"\1_sampling"
# Usage: python build_round1_representativity_prototype_2026-10-02.py <snapshot_dir> [out_dir]
# <snapshot_dir> holds the real_submissions.csv + CONFIRMED_DELETIONS_OVERLAY.csv to use.
SNAP = sys.argv[1] if len(sys.argv) > 1 else SAMPLING_DIR + r"\resampling\output\dual_round_prototype_snapshot_2026-10-01_pre_recovery"
OUT = sys.argv[2] if len(sys.argv) > 2 else SAMPLING_DIR + r"\resampling\output\round1_representativity_prototype_2026-10-02"
Z, P = 1.6448536269514722, 0.5
TARGET = 10.0
FLOOR = 20

spec = importlib.util.spec_from_file_location("acc05", SAMPLING_DIR + r"\resampling\scripts\05_build_accessibility_impact_workbook.py")
m05 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m05)
m05.REAL_SUBMISSIONS_CSV = SNAP + r"\real_submissions.csv"
m05.CONFIRMED_DELETIONS_OVERLAY_CSV = SNAP + r"\CONFIRMED_DELETIONS_OVERLAY.csv"


def ndeff(n, N):
    return n * (N - 1) / (N - n)


def var_simple(n, N):
    """deff = 1 variance of a proportion at p = 0.5, FPC-adjusted; 0 when the stratum is fully enumerated."""
    if n <= 0:
        return None
    if N <= n:
        return 0.0
    return P * (1 - P) / ndeff(n, N)


def moe_from_var(v):
    return None if v is None else 100 * Z * math.sqrt(v)


def main():
    os.makedirs(OUT, exist_ok=True)
    ward_status = m05.load_ward_status_lookup()
    provenance = m05.load_ward_provenance_lookup()
    lga_fractions = m05.load_lga_area_pop_fractions()
    households = m05.classify_households(ward_status)
    cluster_status = m05.load_cluster_status()
    ach_by_strata, ach_by_cluster = m05.load_real_achieved()
    cluster_rows = m05.build_cluster_level(households, provenance, cluster_status, ach_by_cluster)
    clusters_by_strata = defaultdict(list)
    for c in cluster_rows:
        clusters_by_strata[c["strata_id"]].append(c)

    # Universe = the main workbook's own 327 strata: 305 covered + the excluded ones that were once
    # active (05: exclusion_reason set and not partner_coverage_declined). The other ~245 rows of the
    # strata file are LGAs no partner was ever assigned to - out of scope, never in any aggregate.
    all_strata = m05.load_csv(m05.STRATA_CSV)

    def is_covered(s):
        return s.get("coverage_status") == "covered" and s.get("exclusion_reason") in ("none", "", None)

    def is_excluded_in_scope(s):
        return (s.get("coverage_status") != "covered" and s.get("exclusion_reason") not in ("none", "", None)
                and "partner_coverage_declined" not in (s.get("exclusion_reason") or ""))

    strata = [s for s in all_strata if is_covered(s) or is_excluded_in_scope(s)]
    out_of_scope_with_data = [(s["strata_id"], s["adm2_name"], s.get("exclusion_reason"), ach_by_strata.get(s["strata_id"], 0))
                              for s in all_strata if s not in strata and ach_by_strata.get(s["strata_id"], 0) > 0]
    print(f"universe: {len(strata)} strata ({sum(1 for s in strata if is_covered(s))} covered + "
          f"{sum(1 for s in strata if not is_covered(s))} excluded) of {len(all_strata)} rows")
    print("OUT-OF-SCOPE strata that nonetheless have Round 1 interviews:", out_of_scope_with_data or "none")
    known = {s["strata_id"] for s in all_strata}
    orphan = {k: v for k, v in ach_by_strata.items() if k not in known}
    print("interviews matched to a strata_id absent from the strata file:", orphan or "none")
    out_rows = []
    for s in strata:
        sid, pop_type = s["strata_id"], s["pop_type"]
        covered = is_covered(s)
        frac = lga_fractions.get((s["adm2_pcode"], "Non-IDP" if pop_type == "non_idp" else "IDP"), {"pct_pop_accessible_gis": 0})
        pct_acc = frac["pct_pop_accessible_gis"]
        N_hh = float(s["N_hh"] or 0)
        N_acc = N_hh * pct_acc / 100
        n = ach_by_strata.get(sid, 0)
        cl = clusters_by_strata.get(sid, [])
        n_by_clusters = sum(ach_by_cluster.get(c["cluster_id"], 0) for c in cl)

        v = var_simple(n, N_acc) if (covered and N_acc > 0) else None
        moe_simple = moe_from_var(v)

        moe_strict = None
        if covered and N_acc > 0 and n > 0:
            real_cl = [dict(c, n_primary_ceiling_contribution=ach_by_cluster.get(c["cluster_id"], 0)) for c in cl]
            if sum(c["n_primary_ceiling_contribution"] for c in real_cl) > 0:
                moe_strict = m05.realized_moe_certainty_aware(real_cl, N_acc, int(s["m_used"]), pop_type, ICC=0.06)
                if moe_strict is None:
                    moe_strict = m05.realized_moe_unequal(sum(c["n_primary_ceiling_contribution"] for c in real_cl), N_acc,
                                                          [c["n_primary_ceiling_contribution"] for c in real_cl], ICC=0.06)

        def label(moe):
            if not covered:
                return "Dropped", f"Dropped - LGA/strata excluded ({s.get('exclusion_reason')})"
            if N_acc <= 0:
                return "Dropped", "Dropped - no accessible population"
            if moe is not None and moe <= TARGET:
                return "Representative", "Representative (MoE <= 10%)"
            if n >= FLOOR:
                return "Indicative - meets reporting threshold", "Indicative - meets reporting threshold (>=20 samples achieved)"
            return "Dropped", "Dropped - <20 samples achieved"

        r1, r1_reason = label(moe_simple)
        st, _ = label(moe_strict)
        flags = []
        if r1 == "Representative" and n < FLOOR:
            flags.append(f"Representative on {n} interviews (<20) - small accessible population")
        if covered and N_acc > 0 and n >= N_acc:
            flags.append("achieved >= accessible population estimate: MoE floored at 0")
        if n != n_by_clusters:
            flags.append(f"stratum total {n} != cluster sum {n_by_clusters}")
        out_rows.append({
            "State": s["adm1_name"], "LGA": s["adm2_name"], "Pop type": "IDP" if pop_type == "idp" else "Non-IDP",
            "strata_id": sid, "adm2_pcode": s["adm2_pcode"], "Sampling method": s.get("sampling_method", ""),
            "Covered in design": "Yes" if covered else "No",
            "Households (N_hh)": round(N_hh), "% population accessible": round(pct_acc, 1),
            "Accessible households": round(N_acc, 1), "Achieved (Round 1)": n,
            "MoE Round 1 (deff=1) %": None if moe_simple is None else round(moe_simple, 2),
            "Round 1 label": r1, "Round 1 reason": r1_reason,
            "Strict MoE (reference) %": None if moe_strict is None else round(moe_strict, 2),
            "Strict label (reference)": st,
            "Label changes vs strict": "Yes" if r1 != st else "",
            "Flags": "; ".join(flags),
            "_var": v, "_N_acc": N_acc, "_N_hh": N_hh,
        })

    # ---- aggregates ----
    def aggregate(rows, level_keys, scope):
        groups = defaultdict(list)
        for r in rows:
            if scope == "combined" or r["Pop type"] == scope:
                groups[tuple(r[k] for k in level_keys)].append(r)
        agg = []
        for key, rs in sorted(groups.items()):
            inc = [r for r in rs if r["Round 1 label"] != "Dropped"]
            N_all = sum(r["_N_hh"] for r in rs)
            N_acc_all = sum(r["_N_acc"] for r in rs)
            N_inc = sum(r["_N_acc"] for r in inc)
            n_inc = sum(r["Achieved (Round 1)"] for r in inc)
            if inc and N_inc > 0:
                var = sum((r["_N_acc"] / N_inc) ** 2 * r["_var"] for r in inc)
                moe = moe_from_var(var)
                lab = "Representative" if moe <= TARGET else "Indicative - meets reporting threshold"
            else:
                moe, lab = None, "Dropped - no included strata"
            row = dict(zip(level_keys, key))
            row.update({
                "Scope": {"combined": "IDP + Non-IDP", "IDP": "IDP only", "Non-IDP": "Non-IDP only"}[scope],
                "Strata in unit": len(rs), "Strata included": len(inc),
                "Strata dropped": ", ".join(f"{r['LGA'] if 'LGA' not in level_keys else ''}{' ' if 'LGA' not in level_keys else ''}{r['Pop type']}".strip() for r in rs if r["Round 1 label"] == "Dropped"),
                "Achieved (included strata)": n_inc,
                "Households (all strata)": round(N_all), "Accessible households (all strata)": round(N_acc_all),
                "Households represented (included, accessible)": round(N_inc),
                "% of all households represented": round(100 * N_inc / N_all, 1) if N_all else None,
                "% of accessible households represented": round(100 * N_inc / N_acc_all, 1) if N_acc_all else None,
                "MoE Round 1 (deff=1) %": None if moe is None else round(moe, 2),
                "Round 1 label": lab,
            })
            agg.append(row)
        return agg

    lga = aggregate(out_rows, ["State", "LGA"], "combined")
    state = aggregate(out_rows, ["State"], "combined") + aggregate(out_rows, ["State"], "IDP") + aggregate(out_rows, ["State"], "Non-IDP")
    state.sort(key=lambda r: (r["State"], r["Scope"]))

    # ---- write ----
    strata_cols = [k for k in out_rows[0] if not k.startswith("_")]
    strata_out = sorted(out_rows, key=lambda r: (r["State"], r["LGA"], r["Pop type"]))
    for name, rows, cols in (("round1_strata.csv", strata_out, strata_cols), ("round1_lga.csv", lga, list(lga[0])), ("round1_state.csv", state, list(state[0]))):
        with open(os.path.join(OUT, name), "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)

    def count(rows, key="Round 1 label"):
        c = defaultdict(int)
        for r in rows:
            c[r[key]] += 1
        return dict(c)

    s_r1, s_st = count(strata_out), count(strata_out, "Strict label (reference)")
    changed = [r for r in strata_out if r["Label changes vs strict"]]
    print("\nSTRATA (", len(strata_out), "): Round 1 simplified", s_r1, "| strict reference", s_st, "| labels that change:", len(changed))
    print("LGA combined (", len(lga), "):", count(lga))
    for sc in ("IDP + Non-IDP", "IDP only", "Non-IDP only"):
        print(f"STATE {sc}:", count([r for r in state if r["Scope"] == sc]))
    flagged = [r for r in strata_out if r["Flags"]]
    print("flagged strata:", len(flagged))
    for r in flagged[:12]:
        print("   ", r["strata_id"], r["LGA"], r["Pop type"], "|", r["Flags"])
    dropped_with_data = [r for r in strata_out if r["Round 1 label"] == "Dropped" and r["Achieved (Round 1)"] > 0]
    by_reason = defaultdict(lambda: [0, 0])
    for r in dropped_with_data:
        by_reason[r["Round 1 reason"]][0] += 1
        by_reason[r["Round 1 reason"]][1] += r["Achieved (Round 1)"]
    total_n = sum(r["Achieved (Round 1)"] for r in strata_out)
    print(f"Round 1 interviews in Dropped strata (NOT used in Round 1 estimates): {sum(v[1] for v in by_reason.values())} of {total_n} in {len(dropped_with_data)} strata")
    for reason, (k, n) in sorted(by_reason.items(), key=lambda kv: -kv[1][1]):
        print(f"    {reason}: {k} strata, {n} interviews")

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    notes = [
        ("Round 1 representativity - PROTOTYPE", True),
        ("Built 2 Oct 2026 overnight by Coordinator on the FROZEN Round 1 snapshot (28,046 submissions, pre-recovery). Not live; not partner-facing. Re-run after the recovery closeout with the post-recovery submissions.", False),
        ("Rules applied (Jack's decisions, 1 Oct)", True),
        ("Round 1 gate: same MoE formula as the main workbook with design effect = 1 (cluster structure ignored), applied uniformly to every stratum; certainty-site treatment not used. Z = 1.6449 (90% confidence), p = 0.5, finite-population correction on the accessible households; MoE = 0 when achieved >= the accessible-household estimate.", False),
        ("Labels: Representative = MoE <= 10%. Indicative - meets reporting threshold = MoE > 10% with >= 20 achieved. Dropped = LGA/stratum excluded, no accessible population, or < 20 achieved.", False),
        ("Achieved = completed interviews matched to a sampled point, minus confirmed/contested deletions (same definition as the main workbook). Accessible households = N_hh x GIS share of population accessible (same as the main workbook).", False),
        ("Aggregates: each LGA and State gets its own MoE: Z x sqrt(sum of W^2 x Var) over its included strata, W = share of accessible households, Var at design effect 1. ALL Dropped strata are left out of every aggregate, as if not in the design. Each aggregate shows the share of households it actually speaks for.", False),
        ("Combined IDP + Non-IDP aggregates rely on the ToR assumption that WorldPop Non-IDP estimates exclude IDPs (no double counting).", False),
        ("Strict columns (reference only): the main workbook's real-interview MoE (certainty-aware, else cluster-aware Kish) on the same frozen data - the 'strict real-today' basis. Round 2 keeps the strict formula.", False),
        ("Sheets", True),
        ("Strata: one row per stratum with both labels. LGA: IDP + Non-IDP combined (the pop-type split at LGA level is the Strata sheet). State: combined, IDP only, Non-IDP only. Flags: cases needing a human look (e.g. Representative on < 20 interviews because the accessible population is tiny).", False),
        ("Built by resampling/scripts/one_off_analyses/build_round1_representativity_prototype_2026-10-02.py", False),
    ]
    for i, (t, bold) in enumerate(notes, 1):
        c = readme.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=12 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 140

    head_fill = PatternFill("solid", fgColor="1B2A4A")
    for title, rows, cols in (("Strata", strata_out, strata_cols), ("LGA", lga, list(lga[0])), ("State", state, list(state[0]))):
        ws = wb.create_sheet(title)
        for j, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(45, len(col) + 2))
        for i, r in enumerate(rows, 2):
            for j, col in enumerate(cols, 1):
                ws.cell(row=i, column=j, value=r.get(col))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    path = os.path.join(OUT, "Round1_representativity_PROTOTYPE_2026-10-02.xlsx")
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    main()
