# ==============================================================================
# One-off supplemental report for FACT (2026-08-26): a SHORT version of the
# standard accessibility report template, containing ONLY the 189 wards
# missing from their return - not the full 1,067-ward file again, so they
# can close the gap without re-finding their way through what they already
# submitted (and without risk of overwriting good data if they resend the
# whole thing).
#
# Reuses the same ward-splitting/aggregation logic as
# 01_generate_accessibility_reports.py (duplicated, not imported, per this
# project's standalone-script convention), then filters to exactly the
# wards missing from the corrected, already-ingested
# FACT_accessibility_report.xlsx.
#
# Same template shape/dropdowns/styling as the standard generated report, so
# it's immediately familiar to fill in - just README text and scope adapted
# to explain this is a gap-fill supplement, not a fresh full report. No
# Cluster Accessibility sheet - not relevant to closing a ward-level gap.
# ==============================================================================
import csv
from collections import defaultdict

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
STAGE2_CSV = PROJECT_DIR + r"\output\data\data_collection\NGA_MSNA_2026_stage2_sampling_frame_v4_WORKING.csv"
RETURNED_FACT = PROJECT_DIR + r"\resampling\input\accessibility_reports_returned\FACT_accessibility_report.xlsx"
OUT_DIR = PROJECT_DIR + r"\resampling\input\accessibility_reports_generated"
PARTNER = "FACT"

ACCESSIBLE_OPTIONS = ["Yes", "No"]
REASON_OPTIONS = [
    "N/A - fully accessible", "Insecurity / conflict", "Physical access (terrain, flooding, roads)",
    "Population absent / relocated", "Building-footprint issue (no eligible structures)",
    "Access denied by authorities / community", "Other",
]
WARD_COLUMNS = [
    "State", "LGA", "Ward (GRID3)", "Ward (OCHA/COD)",
    "Ward spans multiple LGAs (Y/N)", "Other LGA(s) sharing this ward", "Note",
    "Non-IDP clusters", "IDP clusters", "Total target HHs (primary)",
    "Accessible (Y/N)", "Reason category", "Reason notes",
    "% of target achieved so far", "Date reported",
]
INPUT_COLUMNS = {"Accessible (Y/N)", "Reason category", "Reason notes", "% of target achieved so far", "Date reported"}
MULTI_LGA_WARD_NOTE = (
    "This ward's GRID3 polygon spans more than one LGA - you only need to report on the part of the ward that "
    "falls within your own LGA coverage. See the 'Other LGA(s) sharing this ward' column for where the rest "
    "of it sits."
)


def norm_pop_type(pt):
    return "Non-IDP" if pt == "non_idp" else "IDP"


def build_fact_ward_rows():
    """Same logic as 01's load_cluster_rows_by_partner()+build_ward_rows(),
    scoped to FACT only - full 1,067-ward set, before filtering to the gap."""
    with open(STAGE2_CSV, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    ward_to_lgas = defaultdict(set)
    for r in rows:
        ward_to_lgas[(r["adm1_name"], r["adm3_name"])].add(r["adm2_name"])
    ward_to_lgas = {k: sorted(v) for k, v in ward_to_lgas.items()}

    cluster_rows = defaultdict(list)
    for r in rows:
        cluster_rows[r["cluster_id"]].append(r)

    ward_records = []
    for cid, crows in cluster_rows.items():
        any_row = crows[0]
        partners = [p.strip() for p in any_row["partners_covering"].split(",") if p.strip() and p.strip() != "NA"]
        if PARTNER not in partners:
            continue
        pop_type_norm = norm_pop_type(any_row["pop_type"])

        by_ward = defaultdict(lambda: {"primary": 0, "ward_cod": ""})
        for r in crows:
            key = (r["adm1_name"], r["adm2_name"], r["adm3_name"])
            w = by_ward[key]
            if r["status"] == "primary":
                w["primary"] += 1
            ward_cod = r.get("admin3_cod_name")
            if ward_cod and ward_cod != "NA":
                w["ward_cod"] = ward_cod

        for (state, lga, ward), w in by_ward.items():
            ward_records.append({
                "State": state, "LGA": lga, "Ward (GRID3)": ward, "Ward (OCHA/COD)": w["ward_cod"],
                "Pop Type": pop_type_norm, "Target HHs (primary)": w["primary"],
            })

    agg = defaultdict(lambda: {"non_idp": 0, "idp": 0, "target_hh": 0, "ward_cod": ""})
    for r in ward_records:
        key = (r["State"], r["LGA"], r["Ward (GRID3)"])
        a = agg[key]
        if r["Pop Type"] == "Non-IDP":
            a["non_idp"] += 1
        else:
            a["idp"] += 1
        a["target_hh"] += r["Target HHs (primary)"]
        a["ward_cod"] = r["Ward (OCHA/COD)"]

    out = []
    for (state, lga, ward), a in agg.items():
        other_lgas = [l for l in ward_to_lgas.get((state, ward), [lga]) if l != lga]
        out.append({
            "State": state, "LGA": lga, "Ward (GRID3)": ward, "Ward (OCHA/COD)": a["ward_cod"],
            "Ward spans multiple LGAs (Y/N)": "Yes" if other_lgas else "No",
            "Other LGA(s) sharing this ward": "; ".join(other_lgas),
            "Note": MULTI_LGA_WARD_NOTE if other_lgas else "",
            "Non-IDP clusters": a["non_idp"], "IDP clusters": a["idp"], "Total target HHs (primary)": a["target_hh"],
        })
    out.sort(key=lambda r: (r["State"], r["LGA"], -r["Total target HHs (primary)"], r["Ward (GRID3)"]))
    return out


def already_submitted_keys():
    wb = openpyxl.load_workbook(RETURNED_FACT, data_only=True)
    ws = wb["Ward Accessibility"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        state, lga, ward = row[idx["State"]], row[idx["LGA"]], row[idx["Ward (GRID3)"]]
        if state and lga and ward:
            keys.add((state, lga, ward))
    return keys


def build_workbook(missing_rows):
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    readme.column_dimensions["A"].width = 30
    readme.column_dimensions["B"].width = 95

    r = 1
    readme.cell(row=r, column=1, value="FACT - NGA MSNA 2026 accessibility report - GAP-FILL SUPPLEMENT").font = \
        openpyxl.styles.Font(bold=True, size=14, color="1B2A4A")
    r += 2

    overview = (
        "Thank you for your accessibility report - this is NOT a new full report. It's a short supplement "
        f"containing ONLY the {len(missing_rows)} wards that were assigned to your coverage but weren't included "
        "in your original submission (out of 1,067 wards total assigned to FACT). Everything else you already "
        "reported has been processed and does not need to be resent.\n\nPlease fill in Accessible (Y/N) etc. for "
        "each row below, exactly as in the original template, and send this file back to us - no need to merge "
        "it with your original file, we'll combine them on our end."
    )
    cell = readme.cell(row=r, column=1, value=overview)
    cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    readme.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    readme.row_dimensions[r].height = 110
    r += 2
    readme.cell(row=r, column=1, value=f"Wards in this supplement: {len(missing_rows)}").font = \
        openpyxl.styles.Font(bold=True, italic=True)

    ws = wb.create_sheet("Ward Accessibility")
    ws.append(WARD_COLUMNS)
    header_fill_ref = openpyxl.styles.PatternFill("solid", fgColor="1B2A4A")
    header_fill_input = openpyxl.styles.PatternFill("solid", fgColor="2C5F8A")
    for c, col_name in enumerate(WARD_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = header_fill_input if col_name in INPUT_COLUMNS else header_fill_ref

    for row in missing_rows:
        ws.append([row.get(c, "") for c in WARD_COLUMNS])

    n_rows = len(missing_rows)
    accessible_col = WARD_COLUMNS.index("Accessible (Y/N)") + 1
    reason_col = WARD_COLUMNS.index("Reason category") + 1
    dv_access = DataValidation(type="list", formula1=f'"{",".join(ACCESSIBLE_OPTIONS)}"', allow_blank=True)
    dv_reason = DataValidation(type="list", formula1=f'"{",".join(REASON_OPTIONS)}"', allow_blank=True)
    ws.add_data_validation(dv_access)
    ws.add_data_validation(dv_reason)
    dv_access.add(f"{openpyxl.utils.get_column_letter(accessible_col)}2:{openpyxl.utils.get_column_letter(accessible_col)}{n_rows + 1}")
    dv_reason.add(f"{openpyxl.utils.get_column_letter(reason_col)}2:{openpyxl.utils.get_column_letter(reason_col)}{n_rows + 1}")
    ws.freeze_panes = "A2"
    for i, col_name in enumerate(WARD_COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(30, len(col_name) + 4))
    if n_rows:
        tbl = Table(displayName="FACTGapFillWards", ref=f"A1:{openpyxl.utils.get_column_letter(len(WARD_COLUMNS))}{n_rows + 1}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)

    return wb


def main():
    all_rows = build_fact_ward_rows()
    submitted = already_submitted_keys()
    missing_rows = [r for r in all_rows if (r["State"], r["LGA"], r["Ward (GRID3)"]) not in submitted]
    print(f"FACT total ward universe: {len(all_rows)}")
    print(f"Already submitted: {len(submitted & {(r['State'], r['LGA'], r['Ward (GRID3)']) for r in all_rows})}")
    print(f"Missing (this supplement): {len(missing_rows)}")

    wb = build_workbook(missing_rows)
    out_path = OUT_DIR + r"\FACT_accessibility_report_GAP_FILL_2026-08-26.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
