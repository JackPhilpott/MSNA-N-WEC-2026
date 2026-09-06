# ==============================================================================
# Pulls filled-in "<Partner>_accessibility_report.xlsx" copies from
# resampling/input/accessibility_reports_returned/ (drop returned copies
# there manually - pulled from the partner's SharePoint folder, saved from an
# email attachment, or transcribed by hand from a WhatsApp/verbal report into
# a copy of the template) and appends new/changed rows into the master
# resampling/output/resampling_requests_log.csv.
#
# Reads BOTH sheets (2026-08-20, v2 - see 01_generate_accessibility_reports.py
# for why): "Ward Accessibility" (primary - partners mostly report at ward
# level) and "Cluster Accessibility" (secondary - site-specific detail only).
# A ward-level row is logged with cluster_id/pop_type blank and report_level
# = "ward" - expanding it into the actual affected cluster IDs is deliberately
# NOT done here (that's 03_route_and_flag.py's job, once built), because that
# expansion must always be scoped to (this partner's own clusters in that
# specific LGA+Ward), never a bare ward-name lookup - see 01's header comment
# on why a free-floating ward-name join would risk pulling in a neighbouring
# LGA/partner's portion of a same-named ward.
#
# Design intent: the master log is APPEND-ONLY. This script is safe to rerun
# repeatedly (e.g. once a week, or whenever a partner sends an update) - it
# never edits or deletes an existing log row. For each (partner, cluster_id)
# or (partner, state, lga, ward) it only logs a NEW row when the returned
# file's content differs from the most recent non-superseded log row for that
# same key (comparing accessible/reason_category/reason_notes/
# pct_target_achieved) - an unchanged row on a re-ingested file is silently
# skipped, so you can drop the same partner's file in again after a small
# edit without creating duplicate log spam. A genuinely changed row gets a
# fresh request_id and its `supersedes_request_id` points at the row it
# replaces, so the log keeps every version of a partner's story over time,
# not just the latest.
#
# Rows where BOTH Accessible is blank AND Reason category is blank are
# skipped (untouched template rows - not an actual report).
# ==============================================================================
import csv
import glob
import os
from datetime import date

import openpyxl

PROJECT_DIR = r"c:\Users\JackPHILPOTT\ACTED\IMPACT NGA - 02. MSNA\4. Data\MSNA N-WEC 2026\1_sampling"
RETURNED_DIR = PROJECT_DIR + r"\resampling\input\accessibility_reports_returned"
LOG_CSV = PROJECT_DIR + r"\resampling\output\resampling_requests_log.csv"

LOG_FIELDS = [
    "request_id", "date_added", "partner", "report_level", "state", "lga", "ward_name", "cluster_id", "pop_type",
    "source_channel", "reported_by", "accessible", "reason_category", "reason_notes", "pct_target_achieved",
    "date_reported_by_partner", "status", "resolution_mechanism", "resolution_date",
    "resolution_notes", "supersedes_request_id",
]

# Compared to decide "is this actually a change from what's already logged".
COMPARE_FIELDS = ["accessible", "reason_category", "reason_notes", "pct_target_achieved"]

REPORTED_BY_COL = "Reported by (Partner / IMPACT-default)"
SOURCE_CHANNEL_COL = "Source channel"
# Fallback when a returned file leaves these blank (true for most rows - see
# 01_generate_accessibility_reports.py's README note telling partners to
# leave them for internal use): the row came back via the normal returned-
# template flow, so it's reasonable to assume Partner / the standard channel
# absent a more specific value actually entered on the sheet.
DEFAULT_REPORTED_BY = "Partner"
DEFAULT_SOURCE_CHANNEL = "Partner's own template"


def load_existing_log():
    if not os.path.exists(LOG_CSV):
        return []
    with open(LOG_CSV, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def latest_by_key(log_rows):
    """Most recent (highest request_id) row per (partner, report_level, key),
    ignoring which ones have since been superseded - we still want to diff
    against the newest known state either way."""
    latest = {}
    for r in log_rows:
        key = (r["partner"], r["report_level"], r["state"], r["lga"], r.get("ward_name", ""), r.get("cluster_id", ""))
        if key not in latest or int(r["request_id"]) > int(latest[key]["request_id"]):
            latest[key] = r
    return latest


def next_request_id(log_rows):
    if not log_rows:
        return 1
    return max(int(r["request_id"]) for r in log_rows) + 1


def read_sheet_rows(ws, level, partner):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    has_provenance_cols = REPORTED_BY_COL in idx and SOURCE_CHANNEL_COL in idx
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        accessible = row[idx["Accessible (Y/N)"]]
        reason = row[idx["Reason category"]]
        if not accessible and not reason:
            continue
        pct = row[idx["% of target achieved so far"]]
        rec = {
            "partner": partner, "report_level": level,
            "state": row[idx["State"]] or "", "lga": row[idx["LGA"]] or "",
            "ward_name": row[idx["Ward (GRID3)"]] or "" if "Ward (GRID3)" in idx else "",
            "cluster_id": row[idx["Cluster ID"]] or "" if "Cluster ID" in idx else "",
            "pop_type": row[idx["Pop Type"]] or "" if "Pop Type" in idx else "",
            "accessible": accessible or "",
            "reason_category": reason or "",
            "reason_notes": row[idx["Reason notes"]] or "",
            # explicit None-check, not `or ""` - a genuine 0 (fully-failed
            # ward/cluster, arguably the most important value a partner can
            # report) is falsy in Python and would otherwise silently
            # collapse to blank.
            "pct_target_achieved": "" if pct is None else pct,
            "date_reported_by_partner": str(row[idx["Date reported"]] or ""),
            # Read from the sheet rather than assumed - most rows leave these
            # blank (the columns are for OUR use, see 01_generate_accessibility_
            # reports.py's README note), in which case ingest() below falls back
            # to the same default this script always used before the columns
            # existed. A partner or coordinator who did fill them in (e.g.
            # transcribing a WhatsApp report into the sheet) overrides that
            # default, which is the whole point of the columns being editable.
            "reported_by": (row[idx[REPORTED_BY_COL]] or "") if has_provenance_cols else "",
            "source_channel": (row[idx[SOURCE_CHANNEL_COL]] or "") if has_provenance_cols else "",
        }
        out.append(rec)

    if out and not has_provenance_cols:
        print(f"  NOTE: {partner} ({level}) - returned file predates the provenance columns; "
              f"{len(out)} row(s) will get the default ('{DEFAULT_REPORTED_BY}' / '{DEFAULT_SOURCE_CHANNEL}').")
    return out


def read_returned_file(path, partner_from_filename):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    if "Ward Accessibility" in wb.sheetnames:
        rows += read_sheet_rows(wb["Ward Accessibility"], "ward", partner_from_filename)
    if "Cluster Accessibility" in wb.sheetnames:
        rows += read_sheet_rows(wb["Cluster Accessibility"], "cluster", partner_from_filename)
    return rows


def partner_name_from_path(path):
    base = os.path.basename(path)
    return base[: -len("_accessibility_report.xlsx")]


def ingest():
    log_rows = load_existing_log()
    latest = latest_by_key(log_rows)
    req_id = next_request_id(log_rows)
    new_rows = []

    files = sorted(glob.glob(os.path.join(RETURNED_DIR, "*_accessibility_report.xlsx")))
    if not files:
        print(f"No returned files found in {RETURNED_DIR}. Nothing to ingest.")
        return

    for path in files:
        partner = partner_name_from_path(path)
        try:
            reports = read_returned_file(path, partner)
        except (KeyError, PermissionError) as e:
            print(f"WARNING: could not read {path}: {e}")
            continue

        n_logged = 0
        for rep in reports:
            key = (partner, rep["report_level"], rep["state"], rep["lga"], rep["ward_name"], rep["cluster_id"])
            prior = latest.get(key)
            changed = prior is None or any(str(prior.get(f, "")) != str(rep.get(f, "")) for f in COMPARE_FIELDS)
            if not changed:
                continue

            log_row = {f: "" for f in LOG_FIELDS}
            log_row.update(rep)
            log_row["request_id"] = req_id
            log_row["date_added"] = date.today().isoformat()
            log_row["source_channel"] = rep["source_channel"] or DEFAULT_SOURCE_CHANNEL
            log_row["reported_by"] = rep["reported_by"] or DEFAULT_REPORTED_BY
            log_row["status"] = "new"
            log_row["supersedes_request_id"] = prior["request_id"] if prior else ""
            new_rows.append(log_row)
            latest[key] = log_row
            req_id += 1
            n_logged += 1

        print(f"{partner}: {len(reports)} reported rows in file ({n_logged} new/changed).")

    if not new_rows:
        print("\nNo new or changed rows to log.")
        return

    write_header = not os.path.exists(LOG_CSV) or os.path.getsize(LOG_CSV) == 0
    with open(LOG_CSV, "a", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=LOG_FIELDS)
        if write_header:
            writer.writeheader()
        for r in new_rows:
            writer.writerow(r)

    print(f"\nAppended {len(new_rows)} new/changed rows to {LOG_CSV}.")


if __name__ == "__main__":
    ingest()
